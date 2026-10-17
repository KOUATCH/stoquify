from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
W=ROOT/'work'
inv=json.loads((W/'inventory.json').read_text('utf-8'))
src={x['id']:x for x in inv}
def read(i):
    p=Path(src[i]['path'])
    return (W/(i+'.txt')).read_text('utf-8') if p.suffix=='.pdf' else p.read_text('utf-8-sig')
o=read('SRC01'); s=read('SRC08'); old=read('SRC05')
def between(txt,a,b): return txt.split(a,1)[1].split(b,1)[0]
def prefix(txt,p):
    return re.sub(r'(?<![\w-])((?:FBASE\d+|[BPFADIY]\d{2}))(?!\w)',lambda m:p+'-'+m[1],txt)
def standards(txt):
    txt=txt.replace('at least 74°C','75°C for 30 seconds').replace('reach 74°C','reach 75°C for 30 seconds')
    txt=txt.replace('at least 71°C','75°C for 30 seconds').replace('60°C or above','63°C or above')
    txt=txt.replace('at or below 5°C','at or below 4°C').replace('at 5°C or below','at 4°C or below')
    txt=txt.replace('≤5°C','≤4°C').replace('above 63°C','at least 63°C')
    return txt
def clean(txt):
    txt=standards(txt)
    txt=re.sub(r'\n{3,}','\n\n',txt)
    return txt.strip()
def get_sections(txt,pattern):
    matches=list(re.finditer(pattern,txt,re.M)); out=[]
    for m in matches:
        # Section ends at next same-or-higher heading, not merely next matching recipe.
        level=len(m[0])-len(m[0].lstrip('#'))
        tail=txt[m.end():]
        nxt=re.search(r'^#{1,'+str(level)+r'} ',tail,re.M)
        body=tail[:nxt.start()] if nxt else tail
        out.append((m[1],m[0].split(' ',1)[1],body))
    return out

cards_o=get_sections(o,r'^## ((?:B\d{2}|P\d{2}|A\d{2}|FBASE\d+|F\d{2}))[^\n]*')
# B36's source title has a narrative prefix.
m=re.search(r'^## Additional complete recipe B36: ([^\n]+)\n([\s\S]*?)(?=^## )',o,re.M)
cards_o.append(('B36','B36 '+m[1],m[2]))
cards_s=get_sections(s,r'^### ([BPFDIY]\d{2})[^\n]*')
assert len(cards_o)==75, len(cards_o)
assert len(cards_s)==28,len(cards_s)
parts=[]; coverage=[]
def add(t): parts.append(clean(t))
def sect(num,title,body): add(f'# {num}. {title}\n\n{body}')

add('''# Nkambe Bakery & Recipe Handbook

Bread • pastry • savoury food • frozen desserts • juice & yogurt preservation

Integrated working edition | 3 October 2026

Prepared from the accessible documents and bakery discussions dated or modified 4 September–3 October 2026.

Practical recipes, consistent methods and a measured route from pilot batch to daily production.

This is a development and operating handbook. Formulations, yields, equipment capabilities and commercial assumptions require the trials identified in each section. A recipe being fully written does not mean it has been physically tested or approved for sale.

The source and editorial record accompanies this handbook. It identifies the 27 discovered files, duplicate publications, retained variants, corrections, evidence limits and additional preservation references.

<!--TOC-->
''')
sect('1','How to use this handbook','''Begin with the shared controls, select a small opening menu, and print the relevant recipe and component cards for the station. Use one approved version per product. Enter actual batch results before increasing volume.

| Identifier | Meaning | How to use it |
|---|---|---|
| O-B01–O-B36 | October shared bakery and savoury components | Make these before assembling the O-P and O-A products. |
| O-P01–O-P20; O-A01–O-A04 | October finished products and four additions | Follow their component weights and service instructions. |
| O-FBASE1–7; O-F01–8 | October frozen bases and sale formats | Confirm machine suitability, treatment and cold chain first. |
| S-B/P/F/D/I/Y | September specialist recipes | Separate formulas, retained for their different product styles. S-F means fast food, whereas O-F means frozen. |
| T01 | Your cake-donut cracking trial | A controlled test based on your actual formula. |
| V series | Earlier recipes and development variants | Preserve ideas and formula differences; resolve the stated gaps before release. |

All principal formula quantities are grams, including milk, oil, eggs without shells, drained fish and prepared vegetables. The V-series preserves original litres, counts and ambiguous ingredients explicitly; these need weighing and specification before conversion. Only water may be approximated as 1 g/ml. A nominal egg count does not replace edible egg mass.

Baker’s percentages use flour as 100%; other percentages use the stated component basis. Do not mix them. Scale every ingredient by the same factor, then verify equipment load, mixing, heat penetration and cooling afresh. Larger batches are not validated by arithmetic alone.

Raw piece mass, cooked component mass and cooled sale weight are different measurements. Modelled loss is an assumption. Whole pieces = floor(usable batch mass ÷ actual portion mass). Record leftovers, rejects and samples. Air in ice cream increases volume, not ingredient mass.

October cost examples remain illustrative FCFA assumptions. September workbook prices and operating costs are largely unfilled; a missing price is unknown, never zero. The recipe IDs retain their original suffix so the source workbook remains traceable.

Use customer feedback to choose sweetness, size and price. Keep accessible descriptions alongside familiar product names. Clearly identify wheat, egg, milk, fish, groundnut, sesame and other supplier-declared allergens; vegetarian, dairy-free and allergen-free claims require ingredient and cross-contact checks.

Completeness is relative to the sources found: 27 files across the project, bakery chat-linked working folders and relevant Downloads, plus the cake-donut discussion. Google Drive and Pages searches produced no additional in-window bakery documents. Creation dates were not always available; modification dates and documented conversation context were used. No access to every private folder or external account is claimed.''')

controls=between(s,'## 5. Shared operator standards','## 6. Production cards')
controls=controls.replace('### H','## H')
controls=re.sub(r'FoodSafety.gov gives practical cooking benchmarks.*?\[S22\]\.', 'Use the common house cooking/reheating target of 75°C for 30 seconds for the meat, fish, poultry and cooked egg dishes in this handbook. Hold hot perishables at 63°C or above. These are operating targets informed by external guidance, not statements of Cameroonian law [N08].',controls,flags=re.S)
controls=controls.replace('current Yogourmet Original, 3 g per litre milk','Yogourmet Original, provisionally 3 g per litre milk in the source formulas')
controls=controls.replace('This is an accessible named trial choice,','The current manufacturer page specifies one sachet per litre and 42–44°C; verify the mass on the actual regional sachet before using the 6 g recipe entries. The manufacturer’s domestic consumption window is not transferred to commercial fortified or fruit yogurt. This is an accessible named trial choice,')
sect('2','Shared preparation, cooking and storage controls',controls+'''

## Commercial release and preservation

The study days in H05 are sampling times, not shelf lives. Neither a same-day policy nor a seven-day trial window makes an unvalidated process safe. Until the product-specific treatment, sanitation, storage and release procedure is established, use development batches for controlled evaluation only. Approved ready-to-use treated ingredients may simplify commissioning.

Do not use the dairy pot for raw meat. Use dedicated raw-food containers and tools, or a documented sanitation changeover validated for the equipment. A scraper is not automatically a high-shear blender, aerator, dough mixer or pasteurizer. Follow the equipment’s actual working capacity and operating instructions.

Cooling benchmarks apply to cooked perishables; yogurt has its own endpoint and cooling plan. When this handbook refers to S1 or F0, apply H01–H05 here and the frozen-control section. The common cold storage target is 4°C or below, with dairy aging at 2–4°C and frozen stock at −18°C or colder. The FDA reference cooling endpoint of 5°C is retained as a benchmark, followed by the house 4°C storage target.
''')

equipment=between(o,'# 2 Equipment facts and operating assumptions','# 3 The 20 core products')
sect('3','Equipment, capacity and the opening menu',equipment+prefix(between(o,'# 1 Executive recommendation','# 2 Equipment facts and operating assumptions'),'O')+'''

## Decisions carried forward from earlier documents

The older divider, planetary mixer, slicer and sheeter proposals remain conditional on measured bottlenecks. The latest inventory says a slicer is already available. Confirm its condition before purchasing another. A planetary mixer supports cake creaming and pastry work; a sheeter supports lamination; neither automatically has the best return. Request working loads, minimum loads, phase, voltage, rated and starting power, spare parts, manuals, warranty and a test with your actual product. Quote delivered cost and maintenance, not purchase price alone.

Older records describe the spiral as 120 L / 50 kg flour and 2.8/5.6 kW; the October plan gives 6.5 kW. Treat these as unresolved nameplate differences. Do not size wiring, inverter or batch mass from the conflicting claims. A VFD is not a universal solution for arbitrary three-phase machinery; have compatibility checked for the exact motor and supply.

At five baguettes per tray, a 16-tray rack holds 80, not 101. The earlier 18 kg flour, 11.88 kg combined water/ice, 360 g salt and 90 g yeast batch totals 30.33 kg: 101 × 300 g before losses, requiring more than one 80-piece rack. A 300 g piece with 15% bake loss weighs 255 g, not 250 g. Starting mixes every 45 minutes cannot continuously fill an oven cycling every 25 minutes; measure the whole line and schedule around its slowest stage.

Keep ice inside the total formula water. Choose chilled water and ice from measured dough temperature and friction, not a fixed ice percentage. Protect hot bread from road dust using clean ventilated food-grade transport; cool before sealing. Do not apply raw egg after baking. Direct combustion exhaust outdoors through a compliant installation, away from food and customers.

Offer sliced bread, household loaves and breakfast bundles as paid tests. Agege and tangzhong breads remain development opportunities. A soft crumb is a quality attribute, not proof of microbial shelf life. Avoid promises of free solar energy, guaranteed premium prices, instant market dominance or automatic profit from residual oven heat.''')

sect('4','Shared bakery and savoury components',prefix(between(o,'## Condensed milk formulation','## Complete base formulas'),'O')+'''

Use these component formulas with their linked finished-product cards. Costs are illustrative ingredient costs only. Each card’s production column is a planning scale, not an approved machine load. Portions and additional egg wash, coatings, frying oil and packaging appear on the finished-product cards.''')

def emit_o(ids):
    for ident,title,body in cards_o:
        if ident not in ids: continue
        text=prefix('## '+title+'\n\n'+body,'O')
        if ident=='B32': text=text.replace('Milk cake sponge batter','Milk cake sponge: batter to baked component')
        text=text.replace('**S1**','H01–H05').replace('under S1','under H01–H05').replace('Apply S1','Apply H01–H05')
        if ident=='P12': text=text.replace('visible 50 g chicken','a measured 50 g cooked-chicken input before baking')
        if ident=='P20': text=text.replace('B32 Milk cake sponge batter','B32 baked milk cake sponge')
        if ident=='F07': text+='\n\nPlanning clarification: 12 bars is the source’s conservative costing count. At 970 g usable mix and 70 g per bar, 13 whole bars are arithmetically possible with 60 g remaining; use measured mould capacity, losses and saleable count before changing costs.'
        text=re.sub(r'Pilot only small lots for sale within 7 days[^.]*\.', 'Set the commercial life only after the product-specific process and shelf-life assessment.',text)
        text=re.sub(r'\n(?:Release check|Batch release):[^\n]+','',text)
        add(text)
        coverage.append(('SRC01','O-'+ident,title,'Full formula/method retained; common controls harmonized'))
emit_o({f'B{i:02}' for i in range(1,37)})
sect('5','Bread, fried snacks and bakery meals','The finished cards use the shared components in chapter 4. Make the specified component amount, then assemble by weight. Alternatives in the S-series are separate product versions; do not blend their formulas accidentally.')
emit_o({f'P{i:02}' for i in range(1,21)}|{f'A{i:02}' for i in range(1,5)})
sect('6','Shawarma service and preparation ahead',prefix(between(o,'## Standard and small wrap assembly','# 5A Prepare ahead: burgers, meat, chicken and fish'),'O')+'\n\n'+prefix(between(o,'## Decision for this bakery','## Additional complete recipe B36:'),'O')+'\n\n'+prefix(between(o,'## Costing the preparation-ahead route','# 6 Frozen desserts and two-size popsicles'),'O'))

fcontrol=between(o,'## F0 Frozen dairy processing and machine control','## FBASE1 Gravity soft serve vanilla trial mix')
fcontrol=re.sub(r'For on-site dairy-mix processing,[\s\S]*?(?=\n\nAfter a validated)', 'For on-site dairy-mix processing, follow H02. Its batch reference is at least 69°C for 30 minutes throughout the mix in a qualified system. A continuous-process holding time cannot be transferred to an ordinary kettle. Use an approved ready mix if treatment and rapid chilling have not been demonstrated. [N04]',fcontrol)
fcontrol=re.sub(r'Pilot only small lots for sale within 7 days[^.]*\.', 'Do not assign a commercial life from a generic seven-day trial; validate the actual product, process and storage.',fcontrol)
sect('7','Frozen desserts and portion formats',fcontrol)
emit_o({f'FBASE{i}' for i in range(1,8)}|{f'F{i:02}' for i in range(1,9)})

sect('8','Specialist bread, pastry, meals and cultured dairy','''These 28 September formulas preserve meaningful differences: true-butter brioche, baguettes, cinnamon rolls, biscuits, cake-style beignets, vegetable samosas, cultured yogurt and other formats. They use their own yields and batch multiples. Generic preparation and costing instructions have been consolidated in H01–H05 and the commercial chapter; product-specific formulas, methods, defects, variants and packaging remain.

Every S-series card is a trial specification. The original fit classes mean: 1 = close fit with the stated equipment; 2 = operational adjustments/accessories; 3 = additional equipment or process qualification; 4 = unsuitable for the stated setup. These are judgments, not measured commissioning results. Recipe references [S01]–[S34] point to the inherited research register in the final chapter.''')
for ident,title,body in cards_s:
    # Keep component mass reconciliation: multicomponent assembly needs these yields.
    body=re.sub(r'#### Equipment, ingredients and preparation\n[\s\S]*?(?=#### Numbered production method)', '',body)
    body=re.sub(r'^Costing: workbook ID.*\n?', '',body,flags=re.M)
    body=body.replace(' For F02 this bread allocation is split after baking; for P05 slices are cut after the tray bake.','')
    body=re.sub(r'^Assumptions by mass:.*\n?', '',body,flags=re.M)
    body=body.replace(' Verify possible soy, milk or mustard in purchased fats, sauces, seasonings and coating; shared fryer is a cross-contact route.','')
    body=body.replace('###','##') # source subhead #### -> ###
    body=prefix(body,'S')
    if ident.startswith('Y'):
        body+='\n\nCulture check: the retained 6 g entry assumes two 3 g sachets for 2 L milk. Verify sachet mass and current regional supplier instructions before manufacture; a different culture is not substituted gram for gram.'
    add(prefix('## '+title,'S')+'\n\n'+body)
    coverage.append(('SRC08','S-'+ident,title,'Complete formula, yield, method and distinct variants retained; repeated shared text consolidated'))

sect('9','Fruit juice and yogurt preservation',(W/'preservation.md').read_text('utf-8'))
sect('10','Cake-donut cracking: your controlled trial', '''## T01 — Rolled and cut cake donuts

Source: your September discussion “Reduce cracks in cake donuts”. These are your formula quantities, not a new validated commercial recipe.

| Ingredient | Original batch g | 1 kg flour pilot g |
|---|---|---|
| Flour | 6,000 | 1,000 |
| Sugar | 1,000 | 166.67 |
| Butter | 1,000 | 166.67 |
| Baking powder | 80 | 13.33 |
| Egg, edible mass | 1,500 | 250 |
| Water | 1,500 | 250 |
| EDC, identity to confirm | 10 | 1.67 |
| Total | 11,090 | 1,848.34 rounded |

EDC was described as viscous, but its identity, concentration and supplier dose were not confirmed. Do not assume EDC-95 or increase its dose. Obtain the label before continued use; if it cannot be identified, conduct the development control without it and document the change.

1. Record the existing mixing sequence, flour type, dough temperature, rolling thickness, cutter size and piece mass. Hold these constant between comparable tests.
2. Mix briefly after flour addition. Rest covered about 10 minutes as a trial, use minimal bench flour and avoid repeated rolling of scraps.
3. Compare a control at 250 g water per 1 kg flour with 275 g only if the control is stiff or tears. Keep all other identified ingredients unchanged. If already soft, do not add more water simply to remove cracks.
4. Run any EDC omission comparison separately from the water change. One variable at a time makes the result interpretable.
5. Trial frying at 180–185°C with measured oil recovery and a small load. Let the first surface set before turning; determine cook time using actual thickness and a set center. Oil temperature alone is not proof of doneness.
6. Cool samples and compare cracks, greasy texture, interior moisture, handling breakage and eating quality. Record at least ten pieces when possible. Select the formulation from repeated results, not a single attractive piece.

No portion count, rolling thickness, frying duration or shelf life was established in the discussion. Set these during the trial before issuing a production card. Declare wheat, egg and milk plus any emulsifier-carrier allergens.''')
coverage.append(('CHAT01','T01','Rolled cake donuts','User quantities retained; unresolved EDC identity; controlled comparison'))
sect('11','Earlier formulas and development variants',(W/'variants.md').read_text('utf-8'))

sect('12','Costing, production and a staged launch',prefix(between(o,'# 7 Costing pricing and operating profitability','# 10 Sources outstanding information and next actions'),'O').replace('# 8 Two site production staffing and energy plan','## Two-site production, staffing and energy').replace('# 9 Launch validation and daily controls','## Launch validation and daily controls')+'''

## Additional commercial ideas retained from September

Test breakfast BHB (beignets, haricots and bouillie), school/office packs, pain omelette, sandwich preorders, afternoon brioche and seasonal fruit lines. Beans and porridge were concepts, not complete costed formulas in the gathered recipe cards; develop separate controlled recipes before adding them. Urban menu examples from Douala and Yaoundé show availability, not Nkambe demand. Use customer interviews, paid tastings and actual sell-through.

Use paper/ventilated protection for crisp bread and fully cooled sealed packs for soft sliced bread. Compare slice thickness, loaf weight, bag strength, closure, food-contact documentation, printed label space and supplier minimum orders before ordering packaging. The older 10,000–20,000 printed-bag proposal is a purchasing idea; prove the design and turnover with a small run first.

Track contribution per labour minute and bottleneck hour alongside margin per unit. Separate direct costs, overhead and working stock. Do not treat all produced units as sold. Measure food waste, refunds, cold-chain rejects, energy use and repeat purchases. Set replenishment from real time-of-day sales; do not produce every recipe daily.

The September workbook uses low/base/high stress cases: demand multipliers 0.6/1/1.3; ingredient cost 1.2/1/0.95; energy cost 1.3/1/0.95; energy availability 0.6/0.85/0.98; rejects 12%/5%/2%. These are management scenarios, not observed rates. Keep them separate from the October price illustration.

Its capacity example assumes 16 trays, 5 loaves/tray, 30-minute complete oven cycle, 32 working proof trays, 60-minute proof, 16.513 kg usable mixer dough, 330.26 g raw pieces, 30-minute mixer occupancy, 100 units/hour dividing and 100 cooling/packing. The resulting limiting rate is 100 units/hour before availability; a 6-hour shift at 85% availability gives 510 units of illustrative capacity, not demand. Confirm every input through a time study.

Read the workbook’s Formulas, Yields and Batch plan sheets using the S-series suffixes. The source workbook was inspected, including its formulas and unfilled input cells; it was not recalculated or represented as a completed financial forecast. Use the handbook’s cost method to populate verified quotations and actual operating inputs.
''')

sect('13','Quick-reference records and troubleshooting','''## Daily batch record

| Field | Entry |
|---|---|
| Date / station / operator / recipe version | ____________________ |
| Batch ID / ingredient lots / actual weights | ____________________ |
| Mixer or treatment load / equipment ID | ____________________ |
| Dough temperature or milk/juice pH and °Brix | ____________________ |
| Proof / cook / treatment readings and times | ____________________ |
| Cooling start, intermediate and final readings | ____________________ |
| Raw, cooked and cooled weights / count / leftovers | ____________________ |
| Packaging lot / allergen label / storage | ____________________ |
| Dispatch and receipt temperatures / quantity | ____________________ |
| Release decision / deviation / responsible person | ____________________ |

## Common defects

| Observation | Check first | Next controlled adjustment |
|---|---|---|
| Dense bread | Flour strength, yeast, dough development, proof | Correct one measured cause; avoid extra yeast without diagnosis. |
| Greasy brioche | Butter condition, dough/proofer temperature | Chill and incorporate progressively; do not hide with flour. |
| Soggy pies or leaking rolls | Filling water, cooling, sealing, pack condensation | Reduce free liquid, weigh filling and verify cooling/ventilation. |
| Greasy fried snacks | Oil recovery, load, piece thickness, proof | Reduce load and confirm temperature before changing formula. |
| Icy or soft frozen bars | Sugar/solids, freezing rate, temperature cycling | Measure Brix and storage history; reformulate only in a new trial. |
| Yogurt whey separation | Milk solids, heat treatment, acidification, shear | Review logs; stabilizer is not a substitute for process control. |
| Juice separation or browning | Fruit lot, oxygen, formulation, package | Trial smaller sealed portions; distinguish quality from safety. |
| Gas, swelling, mould, unusual odour | Spoilage or lost process control | Isolate and withhold; do not taste or reprocess for sale. |

## Working glossary

Baker’s percentage: ingredient mass relative to flour. °Brix: refractive soluble-solids estimate, not a safety test. DDT: desired dough temperature. Overrun: percentage volume increase from incorporated air. Pasteurization: validated pathogen-reduction treatment, not sterilization. Proof: final dough fermentation before baking. Syneresis: whey separation. Usable yield: measured output available for the next step. Contribution: revenue minus the defined variable costs; it is not net profit.
''')

refs=between(s,'## 11.', '') if False else ''
source_register=read('SRC11')
source_register=source_register[source_register.find('### S01'):]
source_register=re.sub(r'^### (S\d+)',r'## \1',source_register,flags=re.M)
sect('14','Sources and edition notes','''Source documents are identified by SRC numbers in the separate editorial record. Recipes are inherited trial formulations, not laboratory-tested results. October source links embedded in retained text were supplied with that edition; September references below retain their original access date of 13 September 2026. The preservation references N01–N09 were checked for this edition on 3 October 2026. An inherited citation does not certify every adapted quantity or local market claim.

'''+(W/'new_references.md').read_text('utf-8')+'\n\n'+source_register)

manuscript='\n\n'.join(parts)+'\n'
# Repair references to sections in original standalone works.
manuscript=manuscript.replace('section 5A','chapter 6').replace('section 4’s assumed composition','the condensed-milk assumptions in chapter 4')
manuscript=manuscript.replace('components in section 5','components in chapter 4')
manuscript=manuscript.replace('workflow in section 5','workflow in chapter 6')
manuscript=manuscript.replace('section 7','chapter 12').replace('section 4','chapter 4').replace('section 6','chapter 7')
manuscript=re.sub(r'\bS1\b','H01–H05',manuscript)
manuscript=manuscript.replace('75°C for 30 seconds for 30 seconds','75°C for 30 seconds')
manuscript=manuscript.replace('80°C/25-second route','qualified dairy treatment route')
(ROOT/'Nkambe_Bakery_Handbook.md').write_text(manuscript,'utf-8')
(W/'recipe_coverage.json').write_text(json.dumps(coverage,ensure_ascii=False,indent=2),'utf-8')

# Numeric checks on retained primary formulas: all explicitly labelled total rows.
checks=[]
for family,txt in [('O',o),('S',s)]:
    for table in re.findall(r'(?:^\|[^\n]+\n)+',txt,re.M):
        rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in table.splitlines()]
        if len(rows)<3: continue
        head=rows[0]
        if 'Pilot g' not in head: continue
        pi=head.index('Pilot g'); vals=[]; total=None
        for r in rows[2:]:
            if len(r)<=pi:continue
            try:v=float(r[pi].replace(',','').replace('**',''))
            except:continue
            if 'TOTAL' in r[0].upper(): total=v
            else:vals.append(v)
        if total is not None:checks.append({'family':family,'first':rows[2][0],'sum':round(sum(vals),3),'stated':total,'pass':abs(sum(vals)-total)<.1})
(W/'arithmetic_checks.json').write_text(json.dumps(checks,indent=2),'utf-8')
assert all(x['pass'] for x in checks), [x for x in checks if not x['pass']]

# Word publishing: native headings, bookmarks, clickable links, repeating table headers.
def field(p,code):
    r=OxmlElement('w:r'); b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin');r.append(b);p._p.append(r)
    r=OxmlElement('w:r'); t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=code;r.append(t);p._p.append(r)
    r=OxmlElement('w:r'); e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');r.append(e);p._p.append(r)
def inline(p,txt):
    for part in re.split(r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*)',txt):
        m=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if m:
            rel=p.part.relate_to(m[2],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'315A69');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
        else:
            rr=p.add_run(part.replace('**','').replace('`',''));rr.bold=part.startswith('**')
def make_doc(md,path):
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    for st in d.styles:
        for border in st.element.findall('.//'+qn('w:pBdr')): border.getparent().remove(border)
    sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
    sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
    normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.03
    for name,size in [('Title',30),('Heading 1',21),('Heading 2',15),('Heading 3',11.5)]:
        st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=True;st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(7)
    d.styles['Heading 1'].paragraph_format.page_break_before=True
    hp=sec.header.paragraphs[0];hp.text='NKAMBE  /  BAKERY & RECIPE HANDBOOK';hp.style='Caption';hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(75,75,75)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;fp.add_run('Working edition • 3 October 2026     |     ');field(fp,' PAGE ')
    lines=md.splitlines();i=0;first=True;bid=1
    while i<len(lines):
        l=lines[i].strip()
        if not l: i+=1;continue
        if l=='<!--TOC-->':
            d.add_page_break();d.add_heading('Contents',level=1);field(d.add_paragraph(),' TOC \\o "1-2" \\h \\z \\u ');i+=1;continue
        if l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in cells): rows.append(cells)
                i+=1
            n=max(map(len,rows));t=d.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            widths=([3.55]+[(7.1-3.55)/(n-1)]*(n-1)) if n in [2,3,4] and any('Pilot' in x for x in rows[0]) else [7.1/n]*n
            if n==2:widths=[2.3,4.8]
            for j,col in enumerate(t.columns):col.width=Inches(widths[j])
            for ri,row in enumerate(rows):
                cells=t.add_row().cells
                for j,c in enumerate(cells):
                    c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);inline(p,row[j] if j<len(row) else '')
                    for rr in p.runs:rr.font.size=Pt(9 if n>4 else 10);rr.bold=ri==0 or rr.bold
                    pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');mar.append(el)
                    pr.append(mar)
                    if ri==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8E4DB');pr.append(sh)
                trpr=t.rows[-1]._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
                if ri==0:
                    rep=OxmlElement('w:tblHeader');trpr.append(rep)
                    for c in cells:c.paragraphs[0].paragraph_format.keep_with_next=True
            d.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        hm=re.match(r'^(#{1,6}) (.*)',l)
        if hm:
            level=min(len(hm[1]),3)
            if first:p=d.add_paragraph(hm[2],'Title');first=False;p.paragraph_format.space_before=Pt(60)
            else:p=d.add_heading(hm[2],level=level)
            bm=OxmlElement('w:bookmarkStart');bm.set(qn('w:id'),str(bid));bm.set(qn('w:name'),'section_'+str(bid));p._p.insert(0,bm);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bid));p._p.append(end);bid+=1
        elif re.match(r'^\[S\d+\]:',l):
            p=d.add_paragraph(l);p.style='Caption'
        elif l in ['---','***']:pass
        else:
            p=d.add_paragraph();inline(p,l)
            if re.match(r'^\d+\. ',l):p.paragraph_format.left_indent=Inches(.15);p.paragraph_format.first_line_indent=Inches(-.15)
        i+=1
    cp=d.core_properties;cp.title='Nkambe Bakery & Recipe Handbook';cp.subject='Integrated production recipes and preservation';cp.author='Prepared for the bakery owner';cp.keywords='bakery, recipes, yogurt, juice, preservation, Nkambe'
    settings=d.settings.element;upd=OxmlElement('w:updateFields');upd.set(qn('w:val'),'true');settings.append(upd)
    d.save(path)
make_doc(manuscript,ROOT/'Nkambe_Bakery_Handbook.docx')
print(json.dumps({'primary_cards':len(coverage),'formula_total_checks':len(checks),'words':len(manuscript.split()),'docx':str(ROOT/'Nkambe_Bakery_Handbook.docx')},indent=2))
