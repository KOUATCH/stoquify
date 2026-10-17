from pathlib import Path
import openpyxl, json, hashlib

root=Path(__file__).resolve().parent
src=Path(r'E:\retail management systems\Bakery recipes   melto 2026.xlsx')
w=openpyxl.load_workbook(src,data_only=False)
v=openpyxl.load_workbook(src,data_only=True)
cost=w[' COST rRAW MATERIAL']
prices=[]
for r in range(4,41):
    name,price,weight=[cost.cell(r,c).value for c in (2,3,4)]
    if isinstance(name,str) and isinstance(price,(int,float)) and isinstance(weight,(int,float)):
        prices.append(dict(row=r,name=name,price=price,weight=weight))
data={'file':src.name,'sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'prices':prices,'sheets':[]}
out=['PRICES',json.dumps(prices)]
for s in w:
    if s.title.strip()=='COST rRAW MATERIAL': continue
    cells={c.coordinate:{'value':c.value,'cached':v[s.title][c.coordinate].value} for row in s for c in row if c.value is not None}
    data['sheets'].append({'name':s.title,'cells':cells})
    out.append('\n### '+s.title)
    for r in range(1,s.max_row+1):
        parts=[]
        for c in s[r]:
            if c.value is not None:
                val=str(c.value)
                if c.data_type=='f': val+=' [cached='+str(v[s.title][c.coordinate].value)+']'
                parts.append(c.coordinate+':'+val)
        out.append(' | '.join(parts))
(root/'source_extract.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
(root/'source_cells.txt').write_text('\n'.join(out),encoding='utf-8')
print(json.dumps({'prices':prices,'sheets':[(s.title,s.max_row) for s in w],'sha256':data['sha256']}))
