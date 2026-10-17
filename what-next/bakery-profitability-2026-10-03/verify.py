from pathlib import Path
import json,math,hashlib,shutil,re
import openpyxl
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image,ImageDraw
root=Path(__file__).parent; repo=root.parents[1]; d=json.loads((root/'analysis.json').read_text(encoding='utf-8'))
source=Path(r'E:\retail management systems')/d['source']
assert hashlib.sha256(source.read_bytes()).hexdigest()==d['sha256'],'Source changed'
source_w=openpyxl.load_workbook(source,data_only=True)
for rec in d['recipes']:
    s=source_w[rec['name']]
    for r in range(5,rec['last']+1):
        name=s.cell(r,3).value
        if isinstance(name,str) and name.strip() and name.strip().lower()!='totals':
            assert any(x['ingredient']==name for x in rec['lines']),(rec['name'],r,name)
assert len(d['recipes'])==len(source_w.sheetnames)-1==41
assert sum(r['missing']>0 for r in d['recipes'])==6
book=repo/'outputs/bakery-profitability-2026-10-03/Bakery_Melto_Profitability_2026-10-03.xlsx'
w=openpyxl.load_workbook(book,data_only=True); wf=openpyxl.load_workbook(book,data_only=False)
for i,f in enumerate(d['formats'],7):
    s=w['Sale formats']; expected={'H':f['units'],'I':f['units'],'J':f['revenue'],'K':f['unitCost'],'L':f['surplus'],'M':f['margin']}
    for col,v in expected.items():assert math.isclose(s[f'{col}{i}'].value,v,abs_tol=1e-6),(i,col,s[f'{col}{i}'].value,v)
    assert s[f'Q{i}'].value=='n.a.'
for s in w:
    for row in s:
        for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
stage=repo/'docs/bakery/_profitability-staging';name='Bakery_Melto_Profitability_2026-10-03'
assert (stage/(name+'.md')).read_bytes()==(repo/'docs/bakery'/(name+'.md')).read_bytes()
pdf=stage/(name+'.pdf');reader=PdfReader(pdf);texts=[p.extract_text() for p in reader.pages]
joined='\n'.join(texts)
for n,r in enumerate(d['recipes'],1):assert f'{n:02d}. {r["name"]}' in joined,(n,r['name'])
assert '−402.72' in joined or '-402.72' in joined
assert len(reader.pages)>10
doc=pdfium.PdfDocument(str(pdf));thumbs=[]
review=root/'review';review.mkdir(exist_ok=True)
for i in range(len(doc)):
    img=doc[i].render(scale=.65).to_pil().convert('RGB')
    canvas=Image.new('RGB',(img.width,img.height+25),'white');canvas.paste(img,(0,25));ImageDraw.Draw(canvas).text((8,5),f'Page {i+1}',fill='black');thumbs.append(canvas)
for start in range(0,len(thumbs),12):
    subset=thumbs[start:start+12];cw=max(x.width for x in subset);ch=max(x.height for x in subset)
    contact=Image.new('RGB',(cw*4,ch*math.ceil(len(subset)/4)), '#D0D0D0')
    for j,img in enumerate(subset):contact.paste(img,((j%4)*cw,(j//4)*ch))
    contact.save(review/f'contact-{start//12+1}.jpg',quality=92)
for i in [0,1,3,len(doc)-1]:doc[i].render(scale=1.4).to_pil().save(review/f'page-{i+1}.png')
print(json.dumps({'sourceIntact':True,'recipeCoverage':41,'sellingFormats':len(d['formats']),'missingIngredientRecipes':6,'allExportedOfferValuesVerified':True,'pdfPages':len(reader.pages),'freezePanes':{s.title:s.freeze_panes for s in wf}},indent=2))
(root/'verification.json').write_text(json.dumps({'sourceIntact':True,'recipeCoverage':41,'sellingFormats':len(d['formats']),'allExportedOfferValuesVerified':True,'pdfPages':len(reader.pages)},indent=2))
