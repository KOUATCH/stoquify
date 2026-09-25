"""Build the proposed register operating guide; no application or approval writes."""
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'approval-register'
QA = Path(__file__).resolve().parent
STEM = 'STOQUIFY_APPROVAL_REGISTER_OPERATING_GUIDE_2026-09-03'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(.492)

def sty(name, size, color='111111', bold=False, before=0, after=6, spacing=1.10):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name, s.font.size, s.font.bold = 'Calibri', Pt(size), bold
    s.font.italic = False
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.line_spacing = spacing
    s.paragraph_format.widow_control = True
    for border in list(s.element.findall('.//' + qn('w:pBdr'))):
        border.getparent().remove(border)
    for tracking in list(s.element.findall('.//' + qn('w:spacing'))):
        if tracking.getparent().tag == qn('w:rPr'):
            tracking.getparent().remove(tracking)
    return s

sty('Normal', 11)
sty('Title', 24, '17365D', True, after=6)
sty('Subtitle', 12, '555555', after=10)
sty('Heading 1', 16, '2E74B5', True, 16, 8)
sty('Heading 2', 13, '2E74B5', True, 12, 6)
sty('Heading 3', 12, '1F4D78', True, 8, 4)
sty('Header', 9, '555555', after=0)
sty('Footer', 9, '555555', after=0)
# Named operating-guide overrides: table cells, metadata, references, fillable fields.
sty('Guide cell', 10, after=3, spacing=1.05)
sty('Guide table header', 10, '17365D', True, after=3, spacing=1.05)
sty('Guide metadata', 10, '555555', after=3)
sty('Guide note', 9.5, '555555', after=6)
sty('Guide field', 11, after=12)
sty('Table citation', 9, '555555', before=4, after=4)
sty('List Number', 11, after=8, spacing=1.167)
for name in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[name].paragraph_format.keep_with_next = True

# A real numbering definition with explicit standard_business_brief geometry.
numbering = doc.part.numbering_part.element
abstract_id = max([int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))] or [0]) + 1
num_id = max([int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))] or [0]) + 1
abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(abstract_id))
level = OxmlElement('w:lvl'); level.set(qn('w:ilvl'), '0')
for tag, val in [('start', '1'), ('numFmt', 'decimal'), ('lvlText', '%1.'), ('lvlJc', 'left')]:
    el = OxmlElement('w:' + tag); el.set(qn('w:val'), val); level.append(el)
pp = OxmlElement('w:pPr')
tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), '720'); tabs.append(tab); pp.append(tabs)
indent = OxmlElement('w:ind'); indent.set(qn('w:left'), '720'); indent.set(qn('w:hanging'), '360'); pp.append(indent)
spacing = OxmlElement('w:spacing'); spacing.set(qn('w:before'), '0'); spacing.set(qn('w:after'), '160'); spacing.set(qn('w:line'), '280'); spacing.set(qn('w:lineRule'), 'auto'); pp.append(spacing)
level.append(pp); abstract.append(level); numbering.append(abstract)
num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
aid = OxmlElement('w:abstractNumId'); aid.set(qn('w:val'), str(abstract_id)); num.append(aid); numbering.append(num)

section.header.paragraphs[0].text = 'STOQUIFY  /  APPROVAL REGISTER  /  PROPOSED WORKFLOW'
footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Operating guide v1  |  03 Sep 2026  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
md = ['# Stoquify Approval Register — operating guide', '', '**Proposed workflow; not implemented or adopted. Prepared 3 September 2026.**', '']

def p(text, style='Normal'):
    doc.add_paragraph(text, style=style)
    md.extend([('## ' if style == 'Title' else '') + text, ''])

def heading(text):
    doc.add_heading(text, 2)
    md.extend(['### ' + text, ''])

def page(title, subtitle):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    p(title, 'Title'); p(subtitle, 'Subtitle')

def step(text):
    para = doc.add_paragraph(text, style='List Number')
    props = para._p.get_or_add_pPr(); np = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); np.append(il)
    ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(num_id)); np.append(ni); props.append(np)
    md.extend(['1. ' + text, ''])

def form(label, note=''):
    para = doc.add_paragraph(style='Guide field')
    para.add_run(label + ': ').bold = True
    para.add_run('________________________________________________')
    md.extend(['**' + label + ':** [Fill in]', ''])
    if note: p(note, 'Guide note')

def table(headers, rows, widths):
    assert sum(widths) == 9360
    t = doc.add_table(rows=1 + len(rows), cols=len(headers)); t.autofit = False
    pr = t._tbl.tblPr
    for name, attrs in [('tblW', {'w': '9360', 'type': 'dxa'}), ('tblInd', {'w': '120', 'type': 'dxa'}), ('tblLayout', {'type': 'fixed'})]:
        el = pr.find(qn('w:' + name))
        if el is None: el = OxmlElement('w:' + name); pr.append(el)
        for k, v in attrs.items(): el.set(qn('w:' + k), v)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + side)
        for k, v in {'val': 'single', 'sz': '4', 'color': 'B8C2CC'}.items(): el.set(qn('w:' + k), v)
        borders.append(el)
    pr.append(borders)
    grid = t._tbl.tblGrid
    for el in list(grid): grid.remove(el)
    for width in widths:
        el = OxmlElement('w:gridCol'); el.set(qn('w:w'), str(width)); grid.append(el)
    for i, values in enumerate([headers, *rows]):
        rowpr = t.rows[i]._tr.get_or_add_trPr(); rowpr.append(OxmlElement('w:cantSplit'))
        if i == 0: rowpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(values):
            cell = t.cell(i, j); cell.width = Inches(widths[j] / 1440)
            cp = cell._tc.get_or_add_tcPr()
            cw = cp.find(qn('w:tcW')); cw.set(qn('w:w'), str(widths[j])); cw.set(qn('w:type'), 'dxa')
            margins = OxmlElement('w:tcMar')
            for k, v in {'top': 80, 'bottom': 80, 'start': 120, 'end': 120}.items():
                el = OxmlElement('w:' + k); el.set(qn('w:w'), str(v)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            cp.append(margins)
            if i == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'F2F4F7'); cp.append(fill)
            cell.paragraphs[0].style = 'Guide table header' if i == 0 else 'Guide cell'
            cell.paragraphs[0].text = value
    md.append('| ' + ' | '.join(headers) + ' |')
    md.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')
    for row in rows: md.append('| ' + ' | '.join(x.replace('|', '/') for x in row) + ' |')
    md.append('')

# Page 1
page('Stoquify Approval Register', 'How to establish it, use it, and obtain every required reference.')
p('Status: proposed design — screens and automatic outputs are not yet implemented.', 'Guide metadata')
p('Purpose: one internal register for X1 now and other Stoquify approval workflows later.', 'Guide metadata')
p('You should only need to provide the real decision, the correct material, and the authority behind it. Once established, the register creates the references that your current form asks you to find.')
heading('Your normal approval process')
step('Open Governance → Approvals and choose the request assigned to you.')
step('Read the exact document version, scope, conditions and expiry. Choose Approve, Reject, or Request revision.')
step('Confirm your identity and the displayed decision. Stoquify checks your appointment and records the actual event.')
step('Open the resulting record. Use Copy references or Copy X1 fields to reuse the generated values.')
heading('What the system supplies')
table(['Output', 'What it means'], [
    ('Approval record ID', 'Unique reference to the decision actually recorded. A draft request ID does not mean approved.'),
    ('Decision timestamp', 'Server-recorded time of the native confirmation, shown with timezone.'),
    ('Verification reference', 'The event proving which account confirmed this exact decision and under which authority.'),
    ('SOURCE RECORD URI', 'A stable protected link to the record, approved version and verification evidence.'),
    ('File hash', 'A check of the exact stored file bytes. It does not identify the signer.'),
], [2750, 6610])
p('This is an authenticated internal approval process. It requires adoption by the actual owner and verification of people and authority. It does not require buying a commercial signing platform.', 'Guide note')

# Page 2
page('Establish the register once', 'The first record does not need an earlier approval-system ID.')
p('The administrator opens a setup draft. You identify the owner and the authority to establish this process; Stoquify assigns the draft reference. You can upload an existing authority source without first giving it a register URI.')
table(['Information needed', 'Who supplies it and where to get it'], [
    ('Owning organization or project', 'You: actual business/project name and whether this starts with Stoquify development only. Use ownership records, not an assumed fixture tenant.'),
    ('Owner identity and account', 'You: real name, role and existing account/contact. The administrator resolves the internal user ID and verifies the account binding.'),
    ('Basis of authority', 'You: ownership evidence, executive mandate or appointment. For your independently owned project, an owner declaration can start setup with supporting ownership evidence.'),
    ('Setup verifier', 'You name a suitable person who can check identity and authority against an independently established source. They record their check and its result.'),
    ('Process and appointments', 'Owner: adopt the in-app confirmation method and appoint real people to scoped product, security/data and release responsibilities.'),
    ('Validity, retention and recovery', 'Owner/control personnel: decide expiry limits, permitted role combinations, retention policy, archive owner and revocation/recovery responsibility.'),
    ('Address and private storage', 'Administrator: obtain the real HTTPS origin and configured private evidence storage from hosting/infrastructure settings. Do not provide passwords here.'),
], [2850, 6510])
heading('What makes setup effective?')
p('The owner confirms the establishment charter after identity and authority checks. Appointees accept their appointments. The administrator activates the register only after those records and storage/access checks are complete. Stoquify retains the establishment record, its sources and verification history.')
p('If a source or suitable verifier is missing, save the draft and identify the missing item. A role label or another account for the same person does not supply independent verification. Module development and document preparation can still proceed.', 'Guide note')

# Page 3
page('Choose people and verify decisions', 'A technical login and business authority answer different questions.')
table(['Responsibility', 'Needed from you / source'], [
    ('Register owner', 'Actual person entitled to adopt the process and appoint others. Supply authority source and scope from ownership or organizational records.'),
    ('Administrator / controller', 'Person who manages accounts, evidence custody and register operation. An administrator does not receive business approval power automatically.'),
    ('Product authority', 'Real person who owns product priorities and performance-budget decisions. Appointment states scope and validity.'),
    ('Security / data authority', 'Person or people authorized over security, credential use and fixture-data custody. Obtain names from those actual responsibilities.'),
    ('Independent technical reviewer', 'Person permitted to assess evidence independently of implementation/capture. Appointment and review history must show that independence.'),
], [3000, 6360])
heading('How many people do I need?')
p('There is no blanket five-person requirement. One person may cover compatible scopes when the adopted policy permits it. The proposed default separates product and security decisions and retains independent technical review. Existing domain rules still apply. If you are the sole developer, identify the missing reviewer instead of filling that role with a second account or a seeded persona.')
heading('How native confirmation is verified')
p('Stoquify displays the exact decision and document version, asks for fresh identity confirmation, checks the appointment, and records the event. Your approval ID, confirmation reference and timestamp are produced only after commitment. Replacing the file or changing the scope requires a new decision.')
p('The native receipt is evidence of an authenticated internal decision. It is not represented as a third-party signing certificate. The organization adopts this verification method during setup.', 'Guide note')
heading('When someone leaves or loses authority')
p('The owner or authorized delegate revokes the appointment, giving the record and reason. Stoquify preserves past history, records the change and prevents ineligible future use. A replacement receives a new appointment; nobody edits the old signer into a new person.')

# Page 4
page('Case: X1 product approval', 'Approve the proposed performance budgets and exact scope.')
p('Start from the prepared X1 product form and its review bundle. The register submits an immutable version of both. The person actually appointed to product decisions reviews and confirms the request.')
table(['What is needed', 'Your action and source'], [
    ('Decision', 'Choose Approve, Reject or Request revision after reviewing the form. This is your judgment; it cannot be taken from the database.'),
    ('Budget and scope', 'Review the seven proposed signals in the existing product form. The technical team supplies current contracts and the frozen source bundle.'),
    ('Reason or changes', 'If rejecting or revising, state the issue and desired changes. Revised material must be submitted again before approval.'),
    ('Authority', 'Choose your valid product appointment in the register. Obtain a new appointment from the owner if none covers this scope.'),
    ('Validity / conditions', 'Review the proposed 90-day validity in the existing form and any earlier invalidation on change or revocation. Changes require a revised request.'),
    ('Identity confirmation', 'Use your own verified account and confirm the displayed decision. Do not substitute a seeded approver candidate.'),
], [2750, 6610])
heading('The four values you receive')
p('Stoquify supplies the product decision, product approval reference, verified decision timestamp and signature-verification reference. The Copy X1 fields action carries these into the existing intake. Each source column can use the permanent record link when it exposes the corresponding value and evidence.')
heading('Where to find the prepared material')
p('In docs/X1: X1_PRODUCT_PERFORMANCE_APPROVAL_FOR_SIGNATURE_2026-09-03.docx and X1_PRODUCT_APPROVAL_REVIEW_BUNDLE_2026-09-03.zip.', 'Guide note')
p('The form is currently unsigned. The technical team must freeze the exact current files when submitting. Existing measurements do not establish that the proposed budgets pass: the September 2 investigation found timing overhead and observations needing clarification.', 'Guide note')
p('Product approval authorizes assessment against the agreed budgets. It does not itself authorize credentials, evidence collection or release.')

# Page 5
page('Case: X1 security and evidence', 'Authorize the exact credentials, tenants and collection window.')
p('The security/data request can cover several scopes in one record only when the signer has authority for each and the decision explicitly covers each. Otherwise collect the separate decisions and link them in a verified package.')
table(['Information needed', 'Where to obtain it / what to confirm'], [
    ('Environment and target', 'Environment administrator: real non-production origin, database and approved test scope. Current local candidate: stoquify_replay_next_20260830.'),
    ('Fixture tenants and data', 'Fixture/data owner: authorize tenant rds_org_001 and the rds_org_002 isolation probe if appropriate. Verify data classification and allowed access; these candidates are not approval.'),
    ('Credential reference', 'Secrets/environment administrator: managed secret name or workload-identity reference and its allowed operations. Supply the reference, never the password or token.'),
    ('Capture account / persona', 'Administrator: actual authorized evidence identity and permissions. Record any mapping to test personas separately from human approvers.'),
    ('Permitted activity and window', 'Security/data authority: approve specific read/capture operations, start/end time, date filters and any limits. Do not infer production or write permission.'),
    ('Redaction and retention', 'Data/control owner: identify what may appear in screenshots/logs, where artifacts may be stored, who may read them and how long to retain them.'),
    ('Decision and confirmation', 'Each appointed authority reviews their scope and confirms Approve, Reject or Request revision using their own account.'),
], [2750, 6610])
heading('The references you receive')
p('Evidence decision, credential approval reference, seeded-data approval reference, verified approval timestamp and signature-verification reference. One complete decision can supply both approval references. Split decisions must retain their individual signatures and times; the package cannot invent a combined signature.')
p('A permanent source link may be reused across fields when it exposes the relevant data and verification. Read-only credentials still require the authorization relevant to this evidence capture.', 'Guide note')

# Page 6
page('Finish X1 and fill the old form', 'The record provides approval evidence; technical checks establish readiness.')
table(['Existing intake field', 'Copy from the register'], [
    ('Product decision', 'Actual selected product decision.'),
    ('Product approval reference', 'Issued product decision record ID.'),
    ('Product approval timestamp', 'Verified product decision time with timezone.'),
    ('Product signature reference', 'Product confirmation or verified signature-event reference.'),
    ('Evidence decision', 'Actual security/data decision.'),
    ('Credential approval reference', 'Record explicitly covering credential use.'),
    ('Seeded-data approval reference', 'Record explicitly covering approved fixture data.'),
    ('Evidence approval timestamp', 'Verified time of the complete authorization; retain constituent times when split.'),
    ('Evidence signature reference', 'Verification evidence covering all required security/data scopes.'),
], [3700, 5660])
p('For SOURCE RECORD URI, choose Copy record link. The reviewer must have access. Do not use an expiring download link or the unsigned form reference as an approval ID.', 'Table citation')
p('For “Approval register / controlled system”, copy the active register name and owning organization from Register details. The other registration fields come from the decision record above.', 'Guide note')
heading('What the technical team does next')
p('Run fresh authenticated evidence after the latest required approval. Collect the 42 route observations across seven routes, two languages and three viewport sizes, with current build, fixture and measurement versions. Assess the other required performance signals and resolve zero-row observations against filters, permissions and actual fixture data.')
heading('What the independent reviewer needs')
p('Current approved scope, test results, evidence artifacts and hashes, timing definitions, fixture reconciliation, and an account/appointment independent of implementation and capture. Obtain technical material from the evidence run and CI; obtain authority from the register. Record the actual pass/fail assessment and reasons.')
p('X1 advances only when its applicable approval and technical conditions pass. Approval of review delegation is not a completed technical review. Production release has its own authorization.', 'Guide note')

# Page 7
page('Other cases the register will serve', 'Use a template that names the decision and its required authority.')
table(['Case', 'Information needed from you', 'Where to get it'], [
    ('Policy or procedure', 'Exact revision, purpose, affected scope, effective/expiry dates and decision.', 'Document owner and previous controlled version; appointing owner supplies authority.'),
    ('Change / production release', 'Target, release/build, change plan, rollback, test evidence, window and responsible people.', 'CI/repository and release/infrastructure owners. Execution retains its separate gate.'),
    ('Risk exception', 'Specific waivable requirement, reason, owner, mitigation and short expiry.', 'Risk owner and the policy defining what may be waived. Hard controls remain enforced.'),
    ('Appointment / renewal', 'Real person, account, scope, limits, dates and acceptance.', 'Owner or authorized delegator; previous appointment and identity records.'),
    ('Revoke / supersede', 'Existing record, reason, scope affected and replacement if any.', 'Record page and responsible authority. History is retained; changed content needs a new decision.'),
    ('Paper / external signature', 'Final signed file, signer identity/authority, claimed signing time and verification evidence.', 'Original signer and signing provider/audit trail, or the adopted manual verification process.'),
    ('Future purchasing / payroll / accounting links', 'Existing transaction/approval reference and exact domain scope.', 'The relevant module. Existing business approvals and financial controls remain authoritative.'),
    ('Optional agent delegation', 'Bounded actions, authority source, validity and distinct machine identities.', 'Owner/security/identity administrators. Existing X1 POL-004 is required for that automation path.'),
], [2100, 3500, 3760])
heading('If the signature is on paper')
p('Upload the final signed scan as unverified. Under the adopted process, the signer can log in and confirm the exact uploaded version, or a designated verifier can check the original and identity through an independently established channel and retain the evidence. Store the claimed signing time separately from upload and verification times. Missing proof remains unverified.')
p('Signing a scan again inside Stoquify creates a current confirmation event; it must not pretend to prove an earlier date. A hash confirms unchanged bytes, not who signed them. External signatures require authentic provider verification, not just a pictured signature.', 'Guide note')

# Page 8
page('Your first setup worksheet', 'Fill what you know. Write “not yet available” for anything missing.')
p('This worksheet is preparation, not an approval or an appointment. You do not need to supply record IDs, timestamps, hashes or SOURCE RECORD URIs here.', 'Guide note')
form('Organization or personally owned project', 'Source: actual ownership/business records. State whether this initially covers Stoquify development only.')
form('Your real name, role and account/contact', 'Source: your own details and account profile. The administrator can retrieve the internal user ID.')
form('Your authority basis / existing source', 'Source: ownership evidence, executive mandate or appointment. If you own the project without an existing appointment document, state that fact.')
form('Person who can verify initial setup', 'Name, role and contact of someone able to check identity and authority against an independently established source.')
form('Product decision maker', 'Use a real person from your team or an appointed reviewer, not a seeded candidate.')
form('Security / data decision maker(s)', 'Identify people responsible for security, credential use and fixture data; compatible scopes may be combined where allowed.')
form('Independent technical reviewer', 'Identify someone independent of implementation/capture, or state that the role is not yet filled.')
form('Internal confirmation method acceptable?', 'State accept / prefer external signing / need to decide. The establishment record will capture the actual adoption later.')
p('Before activation, the owner and administrator will also complete validity, retention, revocation, real hosting address and private storage settings. The START_HERE worksheet lists those sources. Do not put passwords or sensitive identity documents into this file or chat.', 'Guide note')

doc.core_properties.title = 'Stoquify Approval Register — Operating Guide'
doc.core_properties.subject = 'Proposed reusable internal register, first-time setup, X1 and per-case inputs'
doc.core_properties.author = 'Stoquify'
doc.core_properties.comments = 'standard_business_brief; memo_masthead without rule; explicit Guide cell/metadata/note/field overrides; no operational approval generated.'
doc.save(OUT / (STEM + '.docx'))
(OUT / (STEM + '.md')).write_text('\n'.join(md), encoding='utf-8')

# Structural verification before rendering.
assert section.page_width.twips == 12240 and section.page_height.twips == 15840
assert all(x.twips == 1440 for x in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin))
for t in doc.tables:
    widths = [int(c.get(qn('w:w'))) for c in t._tbl.tblGrid]
    assert sum(widths) == 9360
    assert t._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w')) == '9360'
    assert t._tbl.tblPr.find(qn('w:tblInd')).get(qn('w:w')) == '120'
    for row in t.rows:
        assert [int(c._tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w'))) for c in row.cells] == widths
        for c in row.cells:
            margins = c._tc.get_or_add_tcPr().find(qn('w:tcMar'))
            assert {x.tag.split('}')[1]: int(x.get(qn('w:w'))) for x in margins} == {'top':80, 'bottom':80, 'start':120, 'end':120}
for name, before, after in [('Heading 1',16,8),('Heading 2',12,6),('Heading 3',8,4)]:
    s = doc.styles[name]
    assert s.paragraph_format.space_before.pt == before and s.paragraph_format.space_after.pt == after
receipt = {'status': 'PROPOSED_NOT_IMPLEMENTED', 'docx': STEM + '.docx', 'plannedPages':8, 'tables':len(doc.tables), 'structuralChecks':'passed', 'visualQA':'pending', 'sha256':hashlib.sha256((OUT / (STEM + '.docx')).read_bytes()).hexdigest()}
(QA / 'document-qa.json').write_text(json.dumps(receipt, indent=2) + '\n', encoding='utf-8')
print(json.dumps(receipt, indent=2))
