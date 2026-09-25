from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "X1" / "X1_FINANCE_ACCOUNTING_APPROVAL_INTAKE_PREFILLED_2026-09-01.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5A6773"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF7D6"
PALE_RED = "FDECEC"
WHITE = "FFFFFF"
BLACK = "111111"
BORDER = "B8C2CC"


def set_run_font(run, size=None, bold=None, color=BLACK, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        set_row_cant_split(row)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_cell_text(cell, size=9.5, bold=False, color=BLACK):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_table_header(table, labels):
    row = table.rows[0]
    set_repeat_table_header(row)
    for cell, label in zip(row.cells, labels):
        cell.text = label
        shade_cell(cell, LIGHT_BLUE)
        format_cell_text(cell, size=9, bold=True, color=NAVY)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_callout(doc, title, body, fill=LIGHT_BLUE, title_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, size=10.5, bold=True, color=title_color)
    body_paragraph = cell.add_paragraph()
    body_paragraph.paragraph_format.space_after = Pt(0)
    body_paragraph.paragraph_format.line_spacing = 1.15
    body_run = body_paragraph.add_run(body)
    set_run_font(body_run, size=10, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_response_box(doc, prompt, response_hint, extra_rows=None):
    rows = [("VALUE TO ENTER", response_hint)]
    rows.extend(extra_rows or [])
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2300, 7060])
    for index, (label, value) in enumerate(rows):
        left, right = table.rows[index].cells
        left.text = label
        right.text = value
        shade_cell(left, LIGHT_BLUE)
        shade_cell(right, PALE_YELLOW)
        format_cell_text(left, size=9, bold=True, color=NAVY)
        format_cell_text(right, size=10, color=BLACK)
        right.paragraphs[0].paragraph_format.space_after = Pt(8)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(3)
    note.paragraph_format.space_after = Pt(6)
    run = note.add_run(prompt)
    set_run_font(run, size=8.5, italic=True, color=MUTED)


def add_field_section(doc, number, field):
    heading = add_heading(doc, f"{number}. {field['name']}", level=2)
    # Each controlled field starts on its own page so its explanation and
    # response box stay together when Word paginates the form.
    heading.paragraph_format.page_break_before = True

    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1950, 7410])
    entries = [
        ("Plain English", field["plain"]),
        ("Why a human", field["why"]),
        ("Who supplies it", field["who"]),
        ("Where to get it", field["where"]),
        ("Required format", field["format"]),
        ("Valid example", field["example"]),
        ("Do not enter", field["reject"]),
    ]
    for row, (label, value) in zip(table.rows, entries):
        row.cells[0].text = label
        row.cells[1].text = value
        shade_cell(row.cells[0], LIGHT_GRAY)
        format_cell_text(row.cells[0], size=9, bold=True, color=NAVY)
        format_cell_text(row.cells[1], size=9.5, color=BLACK)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_response_box(
        doc,
        field["prompt"],
        field["response"],
        [
            ("SOURCE RECORD URI", "____________________________________________________________"),
            ("COMPLETED BY / DATE", "Name: __________________________    Date: __________________"),
        ],
    )


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_paragraph.add_run("STOQUIFY | X1 FINANCE AND ACCOUNTING")
    set_run_font(header_run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run("DRAFT - NOT EFFECTIVE   |   Page ")
    set_run_font(footer_run, size=9, bold=True, color=MUTED)
    add_page_field(footer_paragraph)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "X1 Finance and Accounting Approval Intake"
    doc.core_properties.subject = "Prefilled form for nine human-issued X1 approval fields"
    doc.core_properties.author = "Stoquify Engineering and Controls"
    doc.core_properties.keywords = "Stoquify, X1, finance, approval, evidence, intake"

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(4)
    kicker_run = kicker.add_run("CONTROLLED COMPLETION FORM")
    set_run_font(kicker_run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run("X1 Finance and Accounting Approval Intake")
    set_run_font(title_run, size=25, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run(
        "System-prefilled context and human completion fields for product budgets and authenticated evidence authorization"
    )
    set_run_font(subtitle_run, size=12, color=MUTED)

    metadata = doc.add_table(rows=5, cols=2)
    metadata.style = "Table Grid"
    set_table_geometry(metadata, [2700, 6660])
    metadata_rows = [
        ("Document status", "DRAFT - NOT EFFECTIVE"),
        ("Prepared date", "2026-09-01"),
        ("Target", "X1 finance and accounting"),
        ("Authorized fixture tenant", "rds_org_001"),
        ("Human fields remaining", "9"),
    ]
    for row, (label, value) in zip(metadata.rows, metadata_rows):
        row.cells[0].text = label
        row.cells[1].text = value
        shade_cell(row.cells[0], LIGHT_BLUE)
        format_cell_text(row.cells[0], size=9.5, bold=True, color=NAVY)
        format_cell_text(row.cells[1], size=10, color=BLACK)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    add_callout(
        doc,
        "Important control rule",
        "Do not invent, backdate, or copy approval values from the database. Complete the nine fields only from an authenticated approval or signature record. Evidence capture must be rerun after the latest valid approval timestamp.",
        fill=PALE_RED,
        title_color="9B1C1C",
    )

    add_heading(doc, "How to use this form", level=1)
    steps = [
        "Review the prefilled system context and confirm that the tenant and named candidates are correct.",
        "Have the appointed product approver complete fields 1 through 4.",
        "Have the appointed security or control owner complete fields 5 through 9.",
        "Attach direct record links or signature references. Do not attach passwords, tokens, cookies, or private keys.",
        "Return the completed values for automated manifest binding and a fresh post-approval evidence run.",
    ]
    for step in steps:
        paragraph = doc.add_paragraph(step, style="List Number")
        paragraph.paragraph_format.keep_together = True

    add_heading(doc, "System-prefilled context", level=1)
    summary = doc.add_table(rows=1, cols=2)
    summary.style = "Table Grid"
    set_table_geometry(summary, [3000, 6360])
    add_table_header(summary, ["Information", "System-derived value"])
    summary_rows = [
        ("Database", "stoquify_replay_next_20260830"),
        ("Authorized fixture tenant", "rds_org_001 - Stoquify"),
        ("Foreign isolation probe", "rds_org_002"),
        ("Product-approver candidate", "Kouatchoua Mark - rds_org_001_compliance_user_product_approver"),
        ("Evidence-authorizer candidate", "Yonga Springfield - rds_org_001_compliance_user_security_owner"),
        ("Evidence persona", "Carine Manga - rds_org_001_compliance_user_external_accountant"),
        ("Contract manifest SHA-256", "cdb16e90a74ea942e204bc8191445041a98102850619363096786768e777c9e4"),
        ("Readiness report SHA-256", "312e483e946e22c5078b100649f19a903b5be60ba41ce894db393f4f5cf3f38a"),
        ("Local authenticated baseline", "LOCAL-X1-20260901144339 - 42 artifacts - local only"),
        ("Baseline manifest SHA-256", "4d540d077479567c22ed6a485f029b6ab1cffb5e31250c7870c7dc047e69c256"),
        ("Authoritative blocker", "Traceable product and evidence approvals"),
    ]
    for label, value in summary_rows:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = value
        format_cell_text(cells[0], size=9, bold=True, color=NAVY)
        format_cell_text(cells[1], size=9, color=BLACK)
    set_table_geometry(summary, [3000, 6360])

    add_heading(doc, "What database information proves - and does not prove", level=1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(
        "The database proves that candidate users, roles, permissions, tenant membership, and technical artifacts exist. It does not prove that a person accepted a decision, that the person had formal appointment authority, or that an approval system issued a signed record."
    )
    set_run_font(run, size=11, color=BLACK)

    fields = [
        {
            "name": "Product decision",
            "plain": "The product approver's explicit answer on whether the seven proposed X1 engineering budgets may govern the release evidence run.",
            "why": "A role assignment cannot express consent. Silence, account existence, and prior activity are not approval.",
            "who": "The formally appointed product approver. Current database candidate: Kouatchoua Mark.",
            "where": "Open the product-budget approval request in the recognized approval register and select one decision.",
            "format": "Exactly one of: APPROVE, REJECT, or REVISE.",
            "example": "APPROVE",
            "reject": "Yes, OK, looks good, silence, inferred approval, or a decision entered by automation.",
            "response": "[ ] APPROVE      [ ] REJECT      [ ] REVISE",
            "prompt": "Select exactly one option. If REVISE is selected, describe required changes in the source approval record.",
        },
        {
            "name": "Product approval reference",
            "plain": "The immutable identifier of the approval-system record containing the product decision.",
            "why": "Stoquify can identify the candidate approver but cannot issue an external governance record ID.",
            "who": "The authoritative approval register generates it automatically when the request is created.",
            "where": "The approval page, envelope details, ticket details, or controlled approval-register record.",
            "format": "A stable immutable ID, not a title or editable filename.",
            "example": "APR-X1-PRODUCT-2026-000123 (format example only)",
            "reject": "Draft filename, email subject, database user ID, local simulation ID, or a made-up sequence.",
            "response": "____________________________________________________________",
            "prompt": "Paste the provider-issued immutable record ID and direct record URI.",
        },
        {
            "name": "Product approval timestamp",
            "plain": "The exact signed time at which the product approver submitted the recorded decision.",
            "why": "The time must come from the approval event, not the document creation time or the database clock.",
            "who": "The authoritative approval register records it automatically.",
            "where": "Approval event history, signature certificate, envelope completion record, or immutable audit log.",
            "format": "ISO-8601 with timezone, preferably UTC: YYYY-MM-DDThh:mm:ssZ.",
            "example": "2026-09-02T09:15:30Z (format example only)",
            "reject": "Date without time zone, typed estimate, file modified date, backdated value, or future timestamp.",
            "response": "YYYY-MM-DDThh:mm:ssZ: ______________________________________",
            "prompt": "Copy the timestamp exactly as recorded by the approval system. Do not convert it manually unless the original offset is retained.",
        },
        {
            "name": "Product signature reference",
            "plain": "The verifiable identifier that proves which authenticated approver performed the product decision.",
            "why": "A name or role in the database does not prove that the person personally signed or approved.",
            "who": "The approval or electronic-signature system issues it.",
            "where": "Signature certificate, envelope audit trail, approval event detail, or governance-register signature record.",
            "format": "Provider-issued event, envelope, signature, or certificate ID.",
            "example": "approval-event:7f2c... or envelope:2a91... (format examples only)",
            "reject": "Typed name, pasted image of a signature, database role ID, unsigned Word document, or email screenshot alone.",
            "response": "____________________________________________________________",
            "prompt": "Enter the reference and attach or link the provider's verification record.",
        },
        {
            "name": "Evidence decision",
            "plain": "The security or control owner's explicit decision on whether the named non-production credential, tenant, routes, locales, and viewports may be used for evidence capture.",
            "why": "Technical access does not equal permission to collect and retain evidence.",
            "who": "The formally appointed security or control owner. Current database candidate: Yonga Springfield.",
            "where": "The evidence-authorization request in the recognized approval register.",
            "format": "Exactly one of: APPROVE, REJECT, or REVISE.",
            "example": "APPROVE",
            "reject": "Assumed consent, prior login success, database access, local simulation approval, or an automated decision.",
            "response": "[ ] APPROVE      [ ] REJECT      [ ] REVISE",
            "prompt": "Select exactly one option for the defined rds_org_001 evidence scope.",
        },
        {
            "name": "Credential approval reference",
            "plain": "The immutable record proving that Carine Manga's external-accountant fixture credential may be used for the specified authenticated evidence run.",
            "why": "A working credential and storage state prove authentication, not authorization to use that identity as evidence.",
            "who": "The approval register issues the ID after the security or control owner authorizes credential use.",
            "where": "Credential-use approval, non-production access authorization, or evidence-capture approval record.",
            "format": "Immutable approval record ID with a direct URI.",
            "example": "APR-X1-CREDENTIAL-2026-000124 (format example only)",
            "reject": "Password, token, cookie, storage-state contents, user ID, local preparation ID, or secret-manager value.",
            "response": "____________________________________________________________",
            "prompt": "Enter only the approval reference. Never place credentials or browser cookies in this document.",
        },
        {
            "name": "Seeded-data approval reference",
            "plain": "The immutable record authorizing the specific disposable database and tenants as the evidence fixture scope.",
            "why": "The existence of seeded data does not prove that it is approved for screenshots, measurements, or retained evidence.",
            "who": "The security, privacy, data owner, or delegated control owner according to your governance policy.",
            "where": "Seeded-data authorization, test-data approval, privacy review, or evidence-scope approval record.",
            "format": "Immutable approval record ID naming stoquify_replay_next_20260830, rds_org_001, and rds_org_002.",
            "example": "APR-X1-SEED-DATA-2026-000125 (format example only)",
            "reject": "Database name by itself, seed script output, local readiness report, or an approval for a different tenant.",
            "response": "____________________________________________________________",
            "prompt": "Confirm that the linked record names the exact database, primary tenant, and isolation-probe tenant.",
        },
        {
            "name": "Evidence approval timestamp",
            "plain": "The exact signed time at which the evidence authorization became effective.",
            "why": "Release evidence must be captured after both product and evidence approvals. The timestamp establishes that chronology.",
            "who": "The authoritative approval register records it automatically.",
            "where": "Evidence approval event history, signature certificate, or immutable audit log.",
            "format": "ISO-8601 with timezone, preferably UTC: YYYY-MM-DDThh:mm:ssZ.",
            "example": "2026-09-02T09:20:45Z (format example only)",
            "reject": "Estimated time, document modified date, date without timezone, backdated value, or timestamp copied from product approval.",
            "response": "YYYY-MM-DDThh:mm:ssZ: ______________________________________",
            "prompt": "The new evidence run must begin after the later of this timestamp and the product approval timestamp.",
        },
        {
            "name": "Evidence signature reference",
            "plain": "The verifiable identifier proving which authenticated security or control owner authorized the evidence scope.",
            "why": "Database identity and role membership do not prove a personal approval act.",
            "who": "The approval or electronic-signature system issues it.",
            "where": "Signature certificate, envelope audit trail, approval event detail, or governance-register signature record.",
            "format": "Provider-issued event, envelope, signature, or certificate ID.",
            "example": "approval-event:91bd... or envelope:83c4... (format examples only)",
            "reject": "Typed name, signature image, database role ID, unsigned document, or an unverifiable email statement.",
            "response": "____________________________________________________________",
            "prompt": "Enter the reference and link the verification record. Do not edit a signed artifact after hashing.",
        },
    ]

    for number, field in enumerate(fields, start=1):
        add_field_section(doc, number, field)

    add_heading(doc, "Final completion checklist", level=1)
    checklist = [
        "[ ] Exactly one product decision is selected.",
        "[ ] Product approval ID, timestamp, and signature reference resolve to the same record.",
        "[ ] Exactly one evidence decision is selected.",
        "[ ] Credential and seeded-data references name the exact approved scope.",
        "[ ] Evidence approval timestamp and signature reference resolve to the same record.",
        "[ ] Every timestamp includes a timezone and is copied from the approval system.",
        "[ ] No password, token, cookie, private key, or storage-state content is included.",
        "[ ] No value was inferred from a Stoquify role, local simulation, or database record.",
        "[ ] Direct record URIs are accessible to the independent reviewer.",
    ]
    for item in checklist:
        paragraph = doc.add_paragraph(item)
        paragraph.paragraph_format.left_indent = Inches(0.187)
        paragraph.paragraph_format.first_line_indent = Inches(-0.187)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.runs[0]
        set_run_font(run, size=10.5, color=BLACK)

    add_callout(
        doc,
        "Chronology after completion",
        "Automation will determine the later of the product approval timestamp and evidence approval timestamp, generate a new run ID, and capture a fresh 42-artifact matrix after that point. Existing local artifacts remain development evidence and their timestamps must not be changed.",
        fill=PALE_YELLOW,
        title_color="7A5A00",
    )

    add_heading(doc, "Submission details", level=1)
    submission = doc.add_table(rows=5, cols=2)
    submission.style = "Table Grid"
    set_table_geometry(submission, [2600, 6760])
    submission_rows = [
        ("Completed package URI", "____________________________________________________________"),
        ("Independent reviewer", "____________________________________________________________"),
        ("Review date", "____________________________________________________________"),
        ("Overall decision", "[ ] ACCEPT FOR MANIFEST BINDING    [ ] RETURN FOR REVISION"),
        ("Reviewer notes", "____________________________________________________________\n____________________________________________________________"),
    ]
    for row, (label, value) in zip(submission.rows, submission_rows):
        row.cells[0].text = label
        row.cells[1].text = value
        shade_cell(row.cells[0], LIGHT_BLUE)
        shade_cell(row.cells[1], PALE_YELLOW)
        format_cell_text(row.cells[0], size=9.5, bold=True, color=NAVY)
        format_cell_text(row.cells[1], size=10, color=BLACK)

    final_note = doc.add_paragraph()
    final_note.paragraph_format.space_before = Pt(10)
    final_note.paragraph_format.space_after = Pt(0)
    final_run = final_note.add_run(
        "This completed Word form is an intake package. It becomes authoritative only when the referenced approval and signature records are valid under the organization's governance policy."
    )
    set_run_font(final_run, size=9, italic=True, color=MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    build_document()
