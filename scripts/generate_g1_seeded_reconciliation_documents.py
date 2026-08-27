from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE_DIR = ROOT / "docs" / "blockers-and-gates" / "hris-payroll-compliance-prefill"
SEED_DIR = ROOT / ".seed-artifacts"
SOURCE_DOCX = EVIDENCE_DIR / "COMPLIANCE_AUTHORIZATION_G1_EVIDENCE_RECONCILED_WORKING_COPY_2026-08-19.docx"
SOURCE_PDF = EVIDENCE_DIR / "COMPLIANCE_AUTHORIZATION_G1_EVIDENCE_RECONCILED_WORKING_COPY_2026-08-19.pdf"
SEED_JSON = SEED_DIR / "seed-login-credentials.json"
OUTPUT_DOCX = EVIDENCE_DIR / "COMPLIANCE_AUTHORIZATION_G1_EVIDENCE_SEEDED_RECONCILIATION_WORKING_COPY_2026-08-22.docx"
OUTPUT_PDF = EVIDENCE_DIR / "COMPLIANCE_AUTHORIZATION_G1_EVIDENCE_SEEDED_RECONCILIATION_WORKING_COPY_2026-08-22.pdf"
OUTPUT_MANIFEST = EVIDENCE_DIR / "COMPLIANCE_AUTHORIZATION_G1_EVIDENCE_SEEDED_RECONCILIATION_MANIFEST_2026-08-22.json"
CREDENTIAL_DOCX = SEED_DIR / "SEEDED_DEVELOPMENT_LOGIN_CREDENTIALS_2026-08-22.docx"
APPENDIX_PDF = SEED_DIR / ".g1-seeded-reconciliation-appendix-2026-08-22.pdf"


CANONICAL_ROLES = [
    {"role": "Product owner", "candidate": "Arielle Yongwa", "aliases": ["product owner"]},
    {"role": "Financial controller", "candidate": "Tchami Jennifer", "aliases": ["financial controller"]},
    {"role": "Payments owner", "candidate": "Yonga Junie", "aliases": ["payments owner", "payment owner"]},
    {"role": "Retail operations owner", "candidate": "Tamen Max", "aliases": ["retail operations owner", "retail ops owner"]},
    {"role": "POS architect", "candidate": "tchakoumiLorrain", "aliases": ["pos architect"]},
    {"role": "Security owner", "candidate": "Yonga Springfield", "aliases": ["security owner"]},
    {"role": "Treasury owner", "candidate": "Tamen Stanick", "aliases": ["treasury owner"]},
    {"role": "Risk owner", "candidate": "Tamen Martial", "aliases": ["risk owner"]},
    {"role": "QA owner", "candidate": "Sonkeng Steve", "aliases": ["qa owner", "quality assurance owner"]},
    {"role": "Support owner", "candidate": "Etoo Naomie", "aliases": ["support owner"]},
    {
        "role": "Qualified Cameroon country-pack reviewer",
        "candidate": "Kouatchoua mMark",
        "aliases": ["qualified cameroon country pack reviewer", "qualified cm reviewer", "country pack reviewer"],
    },
    {"role": "SRE owner", "candidate": "Ronald Djakou", "aliases": ["sre owner", "site reliability owner"]},
    {
        "role": "Order-to-cash product owner",
        "candidate": "Tamen Marceline",
        "aliases": ["order to cash product owner", "o2c product owner", "order-to-cash product owner"],
    },
    {
        "role": "Qualified accounting reviewer",
        "candidate": "Tchana Nikita",
        "aliases": ["qualified accounting reviewer", "accounting reviewer"],
    },
    {"role": "Inventory controller", "candidate": "Tchana Rose", "aliases": ["inventory controller"]},
    {"role": "Fulfillment owner", "candidate": "Yonga Lysette", "aliases": ["fulfillment owner", "fulfilment owner"]},
    {"role": "Accounting owner", "candidate": "Yongwa Eli", "aliases": ["accounting owner"]},
]


ADDITIONAL_CANDIDATES = [
    {"role": "Migration operator", "candidate": "Sango Malo", "aliases": ["migration operator"]},
    {"role": "Migration checker", "candidate": "Maximilliano Bonga", "aliases": ["migration checker"]},
    {"role": "Maker", "candidate": "Tamen Marceline", "aliases": ["maker"]},
    {
        "role": "Checker",
        "candidate": "Yonga Claude / Yonga Springfield (source conflict)",
        "aliases": ["checker"],
    },
]


DECISION_SHEETS = [
    ("D-01", ["Product owner", "Financial controller", "Payments owner"]),
    ("D-02", ["Retail operations owner", "POS architect", "Security owner"]),
    ("D-03", ["Payments owner", "Treasury owner", "Security owner"]),
    ("D-04", ["Product owner", "Risk owner", "Retail operations owner"]),
    ("D-05", ["Retail operations owner", "QA owner", "Support owner"]),
    ("D-06", ["Financial controller", "Retail operations owner", "Risk owner"]),
    ("D-07", ["Product owner", "Financial controller", "Qualified Cameroon country-pack reviewer"]),
    ("D-08", ["SRE owner", "Product owner", "Support owner"]),
    ("D-09", ["Financial controller", "Order-to-cash product owner", "Qualified accounting reviewer"]),
    ("D-10", ["Inventory controller", "Fulfillment owner", "Accounting owner"]),
    ("D-11", ["Financial controller", "Treasury owner", "Retail operations owner"]),
]


WARNING = (
    "DEVELOPMENT WORKING COPY ONLY — NOT AUTHORITY, APPOINTMENT, QUALIFICATION, "
    "SEGREGATION-OF-DUTIES, SIGNATURE, APPROVAL, PRODUCTION ACCESS, OR RELEASE EVIDENCE."
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def cell_text(value: Any) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, list):
        return ", ".join(str(item) for item in value)
    return str(value)


def extract_credentials(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return payload
    if not isinstance(payload, dict):
        raise ValueError("Unsupported seed credential JSON structure")
    for key in ("credentials", "users", "logins", "records"):
        value = payload.get(key)
        if isinstance(value, list):
            return value
    for value in payload.values():
        if isinstance(value, list) and value and isinstance(value[0], dict) and "email" in value[0]:
            return value
    raise ValueError("No credential list was found in the seed credential JSON")


def org_id(record: dict[str, Any]) -> str:
    return str(record.get("organizationId") or record.get("organization_id") or record.get("orgId") or "UNASSIGNED")


def org_name(record: dict[str, Any]) -> str:
    return str(record.get("organizationName") or record.get("organization_name") or record.get("orgName") or org_id(record))


def role_haystack(record: dict[str, Any]) -> str:
    codes = record.get("roleCodes") or record.get("role_codes") or []
    return normalized(" ".join([str(record.get("role") or ""), cell_text(codes)]))


def find_role_record(records: Iterable[dict[str, Any]], aliases: list[str]) -> dict[str, Any] | None:
    prepared = [(record, role_haystack(record)) for record in records]
    for alias in aliases:
        needle = normalized(alias)
        exact = [record for record, haystack in prepared if haystack == needle]
        if exact:
            return exact[0]
    for alias in aliases:
        needle = normalized(alias)
        contained = [record for record, haystack in prepared if needle in haystack]
        if contained:
            return contained[0]
    return None


def candidate_for_role(role: str) -> str:
    for item in CANONICAL_ROLES:
        if item["role"] == role:
            return item["candidate"]
    return "—"


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def write_docx_table(document: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if "Table Grid" in [style.name for style in document.styles]:
        table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        header.cells[index].text = label
        shade_cell(header.cells[index], "1F4E78")
        for run in header.cells[index].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(7)
        header.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            row.cells[index].text = cell_text(value)
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in row.cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7)
        if widths:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_warning_box(document: Document, text: str = WARNING) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if "Table Grid" in [style.name for style in document.styles]:
        table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = text
    shade_cell(cell, "FFF2CC")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.bold = True
        run.font.color.rgb = RGBColor(156, 87, 0)
        run.font.size = Pt(10)
    document.add_paragraph()


def add_title(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    if level == 0:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(31, 78, 120)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(31, 78, 120)
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(5)
    else:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(49, 95, 125)
        paragraph.paragraph_format.space_before = Pt(7)
        paragraph.paragraph_format.space_after = Pt(3)


def append_seeded_reconciliation(
    document: Document,
    credentials: list[dict[str, Any]],
    organizations: list[str],
    source_hashes: dict[str, str],
) -> dict[str, Any]:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    add_title(document, "PART VII — SEEDED DEVELOPMENT PERSONA RECONCILIATION", level=0)
    subtitle = document.add_paragraph("Evidence continuation dated 22 August 2026")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_warning_box(document)
    document.add_paragraph(
        "This continuation preserves every heading, rubric, table, candidate name, and unresolved evidence field in the "
        "19 August source working copy. It adds a complete reconciliation of the current seeded development personas. "
        "A seeded account proves only that a test login exists; it does not prove employment, appointment, qualification, "
        "independence, conflict clearance, fresh authentication, approval, or production authorization."
    )

    add_title(document, "VII.1 Source and handling provenance")
    provenance_rows = [
        ["Source compliance DOCX", SOURCE_DOCX.name, source_hashes["source_docx"], "Preserved verbatim as Parts I–VI"],
        ["Source compliance PDF", SOURCE_PDF.name, source_hashes["source_pdf"], "Preserved verbatim in combined PDF"],
        ["Seed login inventory", SEED_JSON.name, source_hashes["seed_json"], "Passwords excluded from tracked evidence"],
        ["Authority disposition", "G0/G1", "OPEN / UNRESOLVED", "Human authority evidence remains required"],
    ]
    write_docx_table(document, ["Artifact", "Filename", "SHA-256 / state", "Treatment"], provenance_rows)

    by_org: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in credentials:
        by_org[org_id(record)].append(record)

    add_title(document, "VII.2 Canonical 17-role reconciliation")
    canonical_rows: list[list[Any]] = []
    canonical_matches = 0
    for item in CANONICAL_ROLES:
        row: list[Any] = [item["role"], item["candidate"]]
        matches_for_role = 0
        for organization in organizations:
            match = find_role_record(by_org[organization], item["aliases"])
            if match:
                matches_for_role += 1
                row.extend([match.get("name"), match.get("email")])
            else:
                row.extend(["MISSING", "MISSING"])
        if matches_for_role == len(organizations):
            canonical_matches += 1
        row.extend([
            "Seed login only — no authority credit",
            "Verified identity; HR assignment; G1 appointment; SoD/COI; qualification where applicable; controlled approval",
        ])
        canonical_rows.append(row)
    headers = ["Canonical G1 role", "19-Aug candidate"]
    for organization in organizations:
        label = org_name(by_org[organization][0]) if by_org[organization] else organization
        headers.extend([f"{label} — seeded name", f"{label} — seeded email"])
    headers.extend(["Current evidence value", "Evidence still required"])
    write_docx_table(document, headers, canonical_rows)

    add_title(document, "VII.3 Additional named operational candidates and source conflicts")
    additional_rows: list[list[Any]] = []
    additional_matches = 0
    for item in ADDITIONAL_CANDIDATES:
        row = [item["role"], item["candidate"]]
        role_matches = 0
        for organization in organizations:
            match = find_role_record(by_org[organization], item["aliases"])
            if match:
                role_matches += 1
                row.extend([match.get("name"), match.get("email")])
            else:
                row.extend(["MISSING", "MISSING"])
        if role_matches == len(organizations):
            additional_matches += 1
        row.extend([
            "UNRESOLVED",
            "Reconcile the Yonga Claude/Yonga Springfield checker conflict; verify identity and formally assign maker/checker duties",
        ])
        additional_rows.append(row)
    additional_headers = ["Operational role", "19-Aug candidate / conflict"]
    for organization in organizations:
        label = org_name(by_org[organization][0]) if by_org[organization] else organization
        additional_headers.extend([f"{label} — seeded name", f"{label} — seeded email"])
    additional_headers.extend(["Authority state", "Required resolution"])
    write_docx_table(document, additional_headers, additional_rows)
    document.add_paragraph(
        "Name-normalization review: the source spellings ‘tchakoumiLorrain’ and ‘Kouatchoua mMark’ are retained exactly. "
        "Any normalized display form (for example, spacing or capitalization) must be confirmed against identity evidence "
        "before an authoritative register is changed."
    )

    add_title(document, "VII.4 Decision-sheet prefill overlays (D-01 through D-11)")
    document.add_paragraph(
        "The following tables prefill candidate and development-login references only. Signature, decision, timestamp, "
        "fresh-auth, appointment, qualification, and segregation-of-duties fields deliberately remain unresolved."
    )
    decision_rows_count = 0
    for sheet_id, roles in DECISION_SHEETS:
        add_title(document, f"{sheet_id} — candidate and seed-account crosswalk", level=2)
        rows: list[list[Any]] = []
        for role in roles:
            role_item = next(item for item in CANONICAL_ROLES if item["role"] == role)
            row = [role, role_item["candidate"]]
            for organization in organizations:
                match = find_role_record(by_org[organization], role_item["aliases"])
                row.extend([match.get("name") if match else "MISSING", match.get("email") if match else "MISSING"])
            row.extend([
                "UNRESOLVED — no signature/approval credit",
                "Appointment + SoD/COI + fresh-auth controlled approval evidence",
            ])
            rows.append(row)
            decision_rows_count += 1
        decision_headers = ["Required role", "19-Aug candidate"]
        for organization in organizations:
            label = org_name(by_org[organization][0]) if by_org[organization] else organization
            decision_headers.extend([f"{label} — seeded name", f"{label} — seeded email"])
        decision_headers.extend(["Approval state", "Completion evidence"])
        write_docx_table(document, decision_headers, rows)

    add_title(document, "VII.5 Complete seeded development identity directory")
    document.add_paragraph(
        "All seeded identities are listed for reconciliation. Passwords are intentionally absent. The separate local "
        "credential document is stored under .seed-artifacts and must never be committed or treated as production access evidence."
    )
    for organization in organizations:
        records = sorted(by_org[organization], key=lambda item: (normalized(item.get("role")), normalized(item.get("name"))))
        label = org_name(records[0]) if records else organization
        add_title(document, f"Organization: {label} ({organization})", level=2)
        rows = [
            [
                record.get("name"),
                record.get("email"),
                record.get("role"),
                record.get("roleCodes") or record.get("role_codes"),
                record.get("locale"),
                record.get("notes") or record.get("note"),
                "DEVELOPMENT ONLY",
            ]
            for record in records
        ]
        write_docx_table(
            document,
            ["Seeded name", "Email", "Role", "Role code(s)", "Locale", "Seed note / provenance", "Authority value"],
            rows,
        )

    add_title(document, "VII.6 Gate disposition and completion controls")
    gate_rows = [
        ["G0 human-input handoff", "OPEN", "Real authorized owners, identity evidence, HR assignment, appointment, SoD/COI and qualification evidence", "Human evidence owner + independent reviewer"],
        ["G1 canonical role authorization", "OPEN", "17/17 formal appointments and role-appropriate authority evidence", "Compliance authority"],
        ["G1 obligation approvals", "OPEN", "33/33 controlled decisions with server-derived actor, tenant and fresh-auth evidence", "Named approvers under valid appointments"],
        ["Decision sheets D-01–D-11", "OPEN", "Complete decisions, signatures/fresh-auth evidence, timestamps and immutable audit/outbox records", "Assigned decision owners"],
        ["Segregation of duties", "OPEN", "Maker/checker and other incompatible-role analysis; conflict resolution; independent approval", "Risk/compliance reviewer"],
        ["Qualified reviewer evidence", "OPEN", "Current professional qualification, jurisdiction/scope, validity dates and independent verification", "Statutory expert reviewer"],
        ["Production authorization", "NOT AUTHORIZED", "Production credentials/provider evidence, managed-database/migration proof and operational release signoff", "Final release authority"],
    ]
    write_docx_table(document, ["Gate / control", "State", "Objective completion evidence", "Required owner"], gate_rows)

    add_title(document, "VII.7 Controlled human completion record")
    completion_rows = [
        ["Prepared by", "", "", "Seeded reconciliation preparation only"],
        ["Independent evidence reviewer", "", "", "Verify source hashes, identity links and non-authority treatment"],
        ["Compliance authority", "", "", "Approve only after every unresolved field has objective evidence"],
        ["Final release authority", "", "", "Production remains unauthorized until all upstream gates close"],
    ]
    write_docx_table(document, ["Capacity", "Verified name", "Date/time + fresh-auth evidence ID", "Decision / scope"], completion_rows)

    document.add_paragraph(
        "Disposition: TECHNICAL SEED RECONCILIATION COMPLETE; HUMAN AUTHORITY AND PRODUCTION RELEASE GATES REMAIN OPEN."
    ).runs[0].bold = True

    return {
        "canonical_roles": len(CANONICAL_ROLES),
        "canonical_roles_matched_in_all_organizations": canonical_matches,
        "additional_roles": len(ADDITIONAL_CANDIDATES),
        "additional_roles_matched_in_all_organizations": additional_matches,
        "decision_sheets": len(DECISION_SHEETS),
        "decision_role_rows": decision_rows_count,
    }


def set_doc_metadata(document: Document, title: str, subject: str) -> None:
    props = document.core_properties
    props.title = title
    props.subject = subject
    props.author = "Stoquify evidence automation — development working copy"
    props.keywords = "Stoquify, G0, G1, compliance, seeded personas, non-authoritative"
    props.comments = WARNING
    props.modified = datetime.now(timezone.utc)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def create_credential_document(credentials: list[dict[str, Any]], organizations: list[str], seed_hash: str) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    add_title(document, "SEEDED DEVELOPMENT LOGIN CREDENTIALS", level=0)
    subtitle = document.add_paragraph("Generated 22 August 2026 from the local seed credential artifact")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_warning_box(
        document,
        "SENSITIVE LOCAL DEVELOPMENT CREDENTIALS — DO NOT COMMIT, SHARE, OR USE IN PRODUCTION. "
        "THESE ACCOUNTS CONFER NO REAL-WORLD AUTHORITY, APPOINTMENT, QUALIFICATION, OR APPROVAL STATUS.",
    )
    document.add_paragraph(f"Source: {SEED_JSON.name}\nSource SHA-256: {seed_hash}\nCredential rows: {len(credentials)}")
    document.add_paragraph(
        "Handling: keep this file only in the Git-ignored .seed-artifacts directory. Rotate/reseed immediately if it leaves "
        "the intended local development environment. Do not attach this file to compliance, audit, release, or support tickets."
    )

    by_org: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in credentials:
        by_org[org_id(record)].append(record)
    for organization in organizations:
        records = sorted(by_org[organization], key=lambda item: (normalized(item.get("role")), normalized(item.get("name"))))
        label = org_name(records[0]) if records else organization
        add_title(document, f"Organization: {label} ({organization})")
        rows = [
            [
                record.get("name"),
                record.get("email"),
                record.get("password"),
                record.get("role"),
                record.get("roleCodes") or record.get("role_codes"),
                record.get("locale"),
                record.get("notes") or record.get("note"),
            ]
            for record in records
        ]
        write_docx_table(
            document,
            ["Name", "Email", "Password", "Role", "Role code(s)", "Locale", "Seed note / intended scenario"],
            rows,
        )
    for section in document.sections:
        section.header.paragraphs[0].text = "SENSITIVE — LOCAL DEVELOPMENT ONLY — .seed-artifacts"
        section.header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in section.header.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(192, 0, 0)
        add_page_number(section.footer.paragraphs[0])
    set_doc_metadata(document, "Seeded Development Login Credentials", "Sensitive local development-only credential inventory")
    document.save(CREDENTIAL_DOCX)


def pdf_paragraph(value: Any, style: ParagraphStyle) -> Paragraph:
    text = cell_text(value)
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(text, style)


def pdf_table(
    rows: list[list[Any]],
    headers: list[str],
    body_style: ParagraphStyle,
    header_style: ParagraphStyle,
    widths: list[float] | None = None,
) -> Table:
    data = [[pdf_paragraph(item, header_style) for item in headers]]
    data.extend([[pdf_paragraph(item, body_style) for item in row] for row in rows])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9EADBA")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
            ]
        )
    )
    return table


def build_appendix_pdf(credentials: list[dict[str, Any]], organizations: list[str], source_hashes: dict[str, str]) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=20,
        textColor=colors.HexColor("#1F4E78"), alignment=TA_CENTER, spaceAfter=9,
    )
    h1 = ParagraphStyle(
        "H1Custom", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=11, leading=13,
        textColor=colors.HexColor("#1F4E78"), spaceBefore=7, spaceAfter=5,
    )
    h2 = ParagraphStyle(
        "H2Custom", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=9, leading=11,
        textColor=colors.HexColor("#315F7D"), spaceBefore=5, spaceAfter=4,
    )
    body = ParagraphStyle("BodySmall", parent=styles["BodyText"], fontSize=5.5, leading=6.8, alignment=TA_LEFT)
    header = ParagraphStyle("HeaderSmall", parent=body, fontName="Helvetica-Bold", textColor=colors.white, fontSize=5.2, leading=6.2)
    warning = ParagraphStyle(
        "Warning", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8, leading=10,
        textColor=colors.HexColor("#9C5700"), backColor=colors.HexColor("#FFF2CC"),
        borderColor=colors.HexColor("#C9A227"), borderWidth=0.6, borderPadding=6, alignment=TA_CENTER, spaceAfter=8,
    )
    normal = ParagraphStyle("NormalCompact", parent=styles["BodyText"], fontSize=7, leading=9, spaceAfter=5)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 6)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(12 * mm, 7 * mm, "Stoquify G1 seeded reconciliation — development working copy")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Appendix page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(APPENDIX_PDF),
        pagesize=landscape(A4),
        leftMargin=9 * mm,
        rightMargin=9 * mm,
        topMargin=10 * mm,
        bottomMargin=11 * mm,
        title="G1 Seeded Development Persona Reconciliation",
        author="Stoquify evidence automation — development working copy",
    )
    story: list[Any] = [
        Paragraph("PART VII — SEEDED DEVELOPMENT PERSONA RECONCILIATION", title_style),
        Paragraph("Evidence continuation dated 22 August 2026", ParagraphStyle("Sub", parent=normal, alignment=TA_CENTER)),
        Paragraph(WARNING, warning),
        Paragraph(
            "This continuation preserves the 19 August source PDF and adds a complete development-persona reconciliation. "
            "Seeded accounts do not prove authority, appointment, qualification, SoD/COI clearance, approval, or production access.",
            normal,
        ),
        Paragraph("VII.1 Source and handling provenance", h1),
    ]
    provenance_rows = [
        ["Source compliance DOCX", SOURCE_DOCX.name, source_hashes["source_docx"], "Preserved in combined output"],
        ["Source compliance PDF", SOURCE_PDF.name, source_hashes["source_pdf"], "Preserved verbatim as prior pages"],
        ["Seed login inventory", SEED_JSON.name, source_hashes["seed_json"], "Passwords excluded"],
        ["Authority disposition", "G0/G1", "OPEN / UNRESOLVED", "Human evidence required"],
    ]
    story.extend([pdf_table(provenance_rows, ["Artifact", "Filename", "SHA-256 / state", "Treatment"], body, header, [31*mm, 61*mm, 85*mm, 81*mm]), Spacer(1, 5)])

    by_org: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in credentials:
        by_org[org_id(record)].append(record)
    story.append(Paragraph("VII.2 Canonical 17-role reconciliation", h1))
    canonical_rows: list[list[Any]] = []
    for item in CANONICAL_ROLES:
        row = [item["role"], item["candidate"]]
        for organization in organizations:
            match = find_role_record(by_org[organization], item["aliases"])
            row.extend([match.get("name") if match else "MISSING", match.get("email") if match else "MISSING"])
        row.extend(["Seed login only", "Identity + HR + appointment + SoD/COI + controlled approval"])
        canonical_rows.append(row)
    headers = ["Canonical role", "19-Aug candidate"]
    for organization in organizations:
        headers.extend([f"{organization} name", f"{organization} email"])
    headers.extend(["Evidence value", "Still required"])
    story.extend([pdf_table(canonical_rows, headers, body, header), Spacer(1, 5)])

    story.append(Paragraph("VII.3 Additional candidates and source conflicts", h1))
    additional_rows: list[list[Any]] = []
    for item in ADDITIONAL_CANDIDATES:
        row = [item["role"], item["candidate"]]
        for organization in organizations:
            match = find_role_record(by_org[organization], item["aliases"])
            row.extend([match.get("name") if match else "MISSING", match.get("email") if match else "MISSING"])
        row.extend(["UNRESOLVED", "Verify identity, assignment, SoD and source-name conflict"])
        additional_rows.append(row)
    additional_headers = ["Role", "Candidate / conflict"]
    for organization in organizations:
        additional_headers.extend([f"{organization} name", f"{organization} email"])
    additional_headers.extend(["Authority", "Resolution"])
    story.extend([pdf_table(additional_rows, additional_headers, body, header), Spacer(1, 5)])

    story.append(PageBreak())
    story.append(Paragraph("VII.4 Decision-sheet prefill overlays", h1))
    story.append(Paragraph("All decision, signature, fresh-auth, appointment and approval fields remain unresolved.", normal))
    for sheet_id, roles in DECISION_SHEETS:
        story.append(Paragraph(f"{sheet_id} — candidate and seed-account crosswalk", h2))
        rows = []
        for role in roles:
            item = next(candidate for candidate in CANONICAL_ROLES if candidate["role"] == role)
            row = [role, item["candidate"]]
            for organization in organizations:
                match = find_role_record(by_org[organization], item["aliases"])
                row.extend([match.get("name") if match else "MISSING", match.get("email") if match else "MISSING"])
            row.extend(["UNRESOLVED", "Appointment + SoD/COI + fresh-auth approval"])
            rows.append(row)
        sheet_headers = ["Required role", "19-Aug candidate"]
        for organization in organizations:
            sheet_headers.extend([f"{organization} name", f"{organization} email"])
        sheet_headers.extend(["Approval", "Completion evidence"])
        story.extend([pdf_table(rows, sheet_headers, body, header), Spacer(1, 4)])

    story.append(PageBreak())
    story.append(Paragraph("VII.5 Complete seeded development identity directory", h1))
    story.append(Paragraph("Passwords are intentionally excluded. The sensitive credential DOCX remains under .seed-artifacts.", normal))
    for organization in organizations:
        records = sorted(by_org[organization], key=lambda item: (normalized(item.get("role")), normalized(item.get("name"))))
        label = org_name(records[0]) if records else organization
        story.append(Paragraph(f"Organization: {label} ({organization})", h2))
        rows = [
            [record.get("name"), record.get("email"), record.get("role"), record.get("roleCodes") or record.get("role_codes"), record.get("locale"), record.get("notes") or record.get("note"), "DEVELOPMENT ONLY"]
            for record in records
        ]
        story.extend([pdf_table(rows, ["Name", "Email", "Role", "Role code(s)", "Locale", "Seed note", "Authority value"], body, header), Spacer(1, 5)])

    story.append(Paragraph("VII.6 Gate disposition", h1))
    gate_rows = [
        ["G0 human-input handoff", "OPEN", "Verified identity, HR assignment, appointment, SoD/COI and qualification evidence"],
        ["G1 canonical role authorization", "OPEN", "17/17 formal appointments"],
        ["G1 obligation approvals", "OPEN", "33/33 controlled decisions with server-derived evidence"],
        ["D-01–D-11", "OPEN", "Decisions, fresh-auth evidence, timestamps, immutable audit/outbox"],
        ["Production", "NOT AUTHORIZED", "All external and operational release evidence"],
    ]
    story.extend([pdf_table(gate_rows, ["Gate", "State", "Completion evidence"], body, header, [55*mm, 35*mm, 168*mm]), Spacer(1, 6)])
    story.append(Paragraph("TECHNICAL SEED RECONCILIATION COMPLETE; HUMAN AUTHORITY AND PRODUCTION RELEASE GATES REMAIN OPEN.", warning))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def merge_pdf() -> None:
    writer = PdfWriter()
    for source in (SOURCE_PDF, APPENDIX_PDF):
        reader = PdfReader(str(source))
        for page in reader.pages:
            writer.add_page(page)
    writer.add_metadata(
        {
            "/Title": "Compliance Authorization G1 Evidence — Seeded Reconciliation Working Copy",
            "/Author": "Stoquify evidence automation — development working copy",
            "/Subject": WARNING,
        }
    )
    with OUTPUT_PDF.open("wb") as handle:
        writer.write(handle)


def docx_xml_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )


def validate_outputs(credentials: list[dict[str, Any]], source_doc: Document) -> dict[str, Any]:
    output_doc = Document(OUTPUT_DOCX)
    credential_doc = Document(CREDENTIAL_DOCX)
    output_pdf = PdfReader(str(OUTPUT_PDF))
    source_pdf = PdfReader(str(SOURCE_PDF))
    tracked_xml = docx_xml_text(OUTPUT_DOCX)
    credential_xml = docx_xml_text(CREDENTIAL_DOCX)
    tracked_pdf_bytes = OUTPUT_PDF.read_bytes()

    passwords = [str(record.get("password") or "") for record in credentials]
    if any(not password for password in passwords):
        raise ValueError("At least one seed credential has no password")
    for password in set(passwords):
        if password in tracked_xml or password.encode("utf-8") in tracked_pdf_bytes:
            raise ValueError("A seed password leaked into a tracked compliance output")
    for record in credentials:
        email = str(record.get("email") or "")
        if not email or email not in tracked_xml or email not in credential_xml:
            raise ValueError("A seeded email is missing from one or more required documents")
    for password in set(passwords):
        expected = sum(1 for item in passwords if item == password)
        if credential_xml.count(password) < expected:
            raise ValueError("The sensitive credential document is not exhaustive")
    if len(output_doc.tables) <= len(source_doc.tables):
        raise ValueError("The compliance working copy did not retain source tables and append reconciliation tables")
    if len(output_pdf.pages) <= len(source_pdf.pages):
        raise ValueError("The combined PDF did not retain source pages and append reconciliation pages")
    return {
        "source_docx_tables": len(source_doc.tables),
        "output_docx_tables": len(output_doc.tables),
        "credential_docx_tables": len(credential_doc.tables),
        "source_pdf_pages": len(source_pdf.pages),
        "output_pdf_pages": len(output_pdf.pages),
        "credential_rows": len(credentials),
        "distinct_passwords": len(set(passwords)),
        "tracked_password_leaks": 0,
    }


def main() -> None:
    for path in (SOURCE_DOCX, SOURCE_PDF, SEED_JSON):
        if not path.exists():
            raise FileNotFoundError(path)
    SEED_DIR.mkdir(exist_ok=True)
    payload = json.loads(SEED_JSON.read_text(encoding="utf-8-sig"))
    credentials = extract_credentials(payload)
    organizations = sorted({org_id(record) for record in credentials})
    source_hashes = {
        "source_docx": sha256(SOURCE_DOCX),
        "source_pdf": sha256(SOURCE_PDF),
        "seed_json": sha256(SEED_JSON),
    }

    source_doc = Document(SOURCE_DOCX)
    output_doc = Document(SOURCE_DOCX)
    reconciliation = append_seeded_reconciliation(output_doc, credentials, organizations, source_hashes)
    set_doc_metadata(
        output_doc,
        "Compliance Authorization G1 Evidence — Seeded Reconciliation Working Copy",
        WARNING,
    )
    output_doc.save(OUTPUT_DOCX)

    create_credential_document(credentials, organizations, source_hashes["seed_json"])
    build_appendix_pdf(credentials, organizations, source_hashes)
    merge_pdf()
    validation = validate_outputs(credentials, source_doc)

    manifest = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "classification": "DEVELOPMENT_WORKING_COPY_NON_AUTHORITATIVE",
        "warning": WARNING,
        "sources": {
            "complianceDocx": {"path": str(SOURCE_DOCX.relative_to(ROOT)), "sha256": source_hashes["source_docx"]},
            "compliancePdf": {"path": str(SOURCE_PDF.relative_to(ROOT)), "sha256": source_hashes["source_pdf"]},
            "seedCredentials": {"path": str(SEED_JSON.relative_to(ROOT)), "sha256": source_hashes["seed_json"], "sensitive": True},
        },
        "outputs": {
            "complianceDocx": {"path": str(OUTPUT_DOCX.relative_to(ROOT)), "sha256": sha256(OUTPUT_DOCX), "containsPasswords": False},
            "compliancePdf": {"path": str(OUTPUT_PDF.relative_to(ROOT)), "sha256": sha256(OUTPUT_PDF), "containsPasswords": False},
            "credentialDocx": {"path": str(CREDENTIAL_DOCX.relative_to(ROOT)), "sha256": sha256(CREDENTIAL_DOCX), "sensitive": True, "gitIgnoredLocation": True},
        },
        "coverage": {
            **reconciliation,
            "g1_obligations_preserved_from_source": 33,
            "seed_credentials": len(credentials),
            "organizations": len(organizations),
            "distinct_roles": len({normalized(record.get("role")) for record in credentials}),
        },
        "validation": validation,
        "gateDisposition": {
            "technicalSeedReconciliation": "COMPLETE",
            "g0HumanInputHandoff": "OPEN",
            "g1AuthorityAuthorization": "OPEN",
            "productionAuthorization": "NOT_AUTHORIZED",
        },
    }
    OUTPUT_MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    APPENDIX_PDF.unlink(missing_ok=True)
    print(
        json.dumps(
            {
                "status": "generated",
                "complianceDocx": str(OUTPUT_DOCX.relative_to(ROOT)),
                "compliancePdf": str(OUTPUT_PDF.relative_to(ROOT)),
                "credentialDocx": str(CREDENTIAL_DOCX.relative_to(ROOT)),
                "manifest": str(OUTPUT_MANIFEST.relative_to(ROOT)),
                "coverage": manifest["coverage"],
                "validation": validation,
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
