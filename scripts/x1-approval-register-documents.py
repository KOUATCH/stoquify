"""Create and extract the local register's plain Word forms. Never sign them."""
from pathlib import Path
import argparse
import json
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs/X1/approval-register'
LABELS = {
    'decision': 'Decision APPROVE REJECT or REVISE',
    'approverName': 'Actual approver full name',
    'approverUserId': 'Approver account mapping',
    'authorityReference': 'Authority appointment or adoption reference',
    'signedAt': 'Actual signing time with timezone',
    'expiresAt': 'Approval expiry time with timezone',
    'reason': 'Reason if rejecting or revising',
    'controllerCertificateSha256': 'Document controller certificate fingerprint',
    'retentionDays': 'Retention period in calendar days',
    'backupLocation': 'Separate backup location',
    'reviewerName': 'Actual document controller name',
    'identityVerification': 'How I checked the actual signer identity',
    'authorityVerification': 'How I checked the signer authority',
    'signatureVerification': 'How I checked the signature and exact decision scope',
    'signingTimeVerification': 'How I checked the actual signing time',
}
POLICIES = {
    'adoption': [
        'By approving and signing, the authorized owner adopts this local register for the specified non-production X1 scope. The owner recognizes a manual signature process: the appointed document controller checks original signatures, identity, authority, decision, scope and time, then seals a verification receipt with the certificate identified in this form.',
        'The owner confirms the person-to-account mappings below and appoints the product, security and independent-review authorities for this scope. These database candidates become appointments only through the actual authorized decision and retained authority evidence.',
        'Automation may allocate draft IDs, calculate hashes, preserve records, verify controller seals, synchronize approved decisions and prepare evidence bindings. Automation may not choose an approval decision, impersonate a signer, override rejection or revocation, or authorize production release. This is the bounded local delegation corresponding to POL-004.',
        'The owner accepts the stated retention period and backup location. Records and audit checkpoints must be backed up separately. No automatic deletion is enabled; any legal hold or required retention takes priority. The controller private key stays in the Windows certificate store. A repository administrator can alter local files, so external checkpoint copies and independent review remain important.',
    ],
    'product': [
        'This records the product authority decision for X1-finance-accounting-performance-budgets, budget contract version 1, against the exact contract and review bundle identified below (POL-001). Approval accepts the evaluation targets. It does not assert that measured performance already meets them.',
        'Proposed targets: bounded local sort/filter response under 100 ms on p75 test hardware; server loading feedback under 150 ms and useful results under 1,000 ms at p75; cached-shell page navigation under 750 ms at p75; immediate bulk-job receipt with progress and partial results; no visible column-layout flash; bounded rendered rows; export duration recorded by result tier and format.',
        'The 900-day duration in the earlier unsigned form is only a proposal. Enter the actual chosen expiry here, within the adopted authority window. Material changes to the approved contract or scope require a new decision.',
    ],
    'evidence': [
        'This is the security/control owner decision for credential use and fixture data in the exact scope below (POL-002). The credential-use and seeded-data clauses are separately addressable under this record ID, so one signed decision can supply both references.',
        'Credential-use clause: authorize the named external-accountant fixture persona for read-only authenticated EN/FR desktop, tablet and mobile evidence. Store passwords and browser sessions outside this register. Generate or renew browser storage state through the existing approved local preparation workflow. This approval does not itself reset passwords, change roles or revoke sessions.',
        'Seeded-data clause: authorize the named database and fixture tenant, with the distinct foreign tenant used only for read-only isolation probes. No production access, shared database reset, fixture reseed, posting, reversal, deletion, payment execution or permission change is authorized by this record.',
        'Capture only the seven routes from the bound X1 contract. Preserve tenant isolation, server-side authorization, protected-action restrictions and redaction. Capture after the actual product and evidence approvals. Keep evidence under the adopted retention and access controls; approval does not promote older local screenshots.',
    ],
    'independent': [
        'This appoints the mapped independent reviewer for X1 technical evidence review (POL-003). The reviewer checks the real approval records, current contract hashes, evidence chronology, authenticated EN/FR coverage and required measurements before any promotion decision.',
        'This delegation is not a review of evidence that has not yet been captured. A separate dated review result must identify the actual run, reviewer, disposition and evidence after capture. The independent reviewer account must differ from the product and security approver accounts.',
        'The manual workflow uses an enrolled document controller rather than invented service-principal identities. Production release and autonomous business decisions remain outside the local delegation.',
    ],
}

def load(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))

def document(title, reference):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    for name, size in [('Normal', 10), ('Title', 21), ('Heading 1', 14), ('Heading 2', 11)]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.05
        borders = style.element.find('.//' + qn('w:pBdr'))
        if borders is not None:
            borders.getparent().remove(borders)
    section.header.paragraphs[0].text = 'STOQUIFY  |  LOCAL X1 APPROVAL REGISTER'
    footer = section.footer.paragraphs[0]
    footer.add_run(reference + '  |  Page ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.add_paragraph(title, 'Title')
    doc.core_properties.author = 'Stoquify'
    doc.core_properties.title = title
    return doc

def table(doc, pairs, fillable=False):
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Inches(2.7), Inches(4.2)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + side)
        for k, v in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items(): b.set(qn('w:' + k), v)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for label, value in pairs:
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for index, text in enumerate((label, value or '[FILL]')):
            cell = row.cells[index]
            cell.width = Inches(2.7 if index == 0 else 4.2)
            cell.text = str(text)
            cell.vertical_alignment = 1
            cell.paragraphs[0].paragraph_format.space_after = Pt(5)
            cell.paragraphs[0].paragraph_format.space_before = Pt(5)
            if index == 0: cell.paragraphs[0].runs[0].bold = True
            if fillable and index == 1:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'FFF2CC'); cell._tc.get_or_add_tcPr().append(shade)
    return t

def create():
    settings = load(BASE / 'register.json')
    made = []
    for kind, record_id in settings['records'].items():
        folder = BASE / 'records' / record_id
        record = load(folder / 'approval-record.json')
        if not (folder / 'approval-form.docx').exists():
            doc = document(record['title'], record_id)
            doc.add_paragraph('Read the complete scope, fill the yellow fields on page 2, then sign. The issued ID reserves this record; approval begins only after a real decision and completed verification.')
            doc.add_paragraph('Issued record ID: ' + record_id)
            doc.add_paragraph('Register: ' + record['registerName'])
            for text in POLICIES[kind]: doc.add_paragraph(text)
            doc.add_heading('Exact fixture and contract scope', 1)
            s = record['scope']
            doc.add_paragraph(f"Database {s['database']}; tenant {s['organizationId']}; foreign read-only probe {s['foreignProbeOrganizationId']}; origin {s['origin']}. Persona {s['personaRole']} / {s['personaUserId']}. Languages EN and FR; desktop, tablet and mobile.")
            doc.add_paragraph('Contract manifest: ' + record['contract']['path'])
            p = doc.add_paragraph('SHA-256: ' + record['contract']['sha256'])
            for run in p.runs: run.font.size = Pt(8)
            p = doc.add_paragraph('Review bundle SHA-256: ' + record['contract']['bundleSha256'])
            for run in p.runs: run.font.size = Pt(8)
            if kind == 'adoption':
                doc.add_heading('Authority appointments subject to this decision', 2)
                for key, identity in settings['identityAssignments'].items():
                    doc.add_paragraph(f"{identity['roleCode']}: {identity['name']} ({identity['userId']})")
            doc.add_page_break()
            doc.add_heading('Decision and signature', 1)
            doc.add_paragraph('Use the actual signing time, with timezone. If signing paper, enter it at the signing event and copy it accurately into the Word form before controller review. Never alter the signed PDF to fill later registration details.')
            values = load(folder / 'approver-input.json')
            keys = list(values) if kind == 'adoption' else [key for key in values if key not in ('controllerCertificateSha256', 'retentionDays', 'backupLocation')]
            table(doc, [(LABELS[key], values[key]) for key in keys], True)
            doc.add_paragraph('Approver signature: __________________________________________')
            doc.add_paragraph('By signing an APPROVE decision, I confirm my authority and the exact scope in this form. REJECT or REVISE prevents approval synchronization. An expiry change, different signer mapping, or changed contract requires appropriate review before signing.')
            doc.add_paragraph('Retain the actual signed PDF as signed-approval.pdf in this record folder. The controller records later verification details in the separate controller-review form.')
            if kind == 'adoption': doc.add_paragraph('Also supply authority-evidence.pdf: the actual appointment, ownership or delegation evidence the controller will check. A seeded role assignment alone is insufficient.')
            doc.save(folder / 'approval-form.docx'); made.append(str(folder / 'approval-form.docx'))
        if not (folder / 'controller-review.docx').exists():
            doc = document('Document controller verification', record_id)
            doc.add_paragraph('Complete this separate review after checking the signed approval. Record concrete observations and references, not just YES. The software checks your cryptographic attestation; it cannot inspect a handwritten signature for you.')
            doc.add_paragraph('Approval record: ' + record_id)
            table(doc, [(LABELS[key], None) for key in ('reviewerName', 'identityVerification', 'authorityVerification', 'signatureVerification', 'signingTimeVerification')], True)
            doc.add_heading('Before sealing', 1)
            doc.add_paragraph('Compare the signed PDF with the filled Word form, including decision, actual signer, account mapping, contract hashes, scope, dates and any rejection reason. Check the original signature and retained authority evidence. For adoption, also compare the controller certificate fingerprint with the enrollment output.')
            doc.add_paragraph('Run the register manager to import forms, prepare the receipt and seal it. The Windows controller must personally confirm the exact record ID. The certificate key stays in the Windows certificate store; no password or private key belongs in these forms.')
            doc.add_paragraph('After sealing, do not edit either Word form, its imported data, the signed PDF or the verification receipt. Preserve and revoke incorrect records instead of rewriting their history.')
            doc.save(folder / 'controller-review.docx'); made.append(str(folder / 'controller-review.docx'))
    lines = ['# Forms to complete', '', 'Start by enrolling the document controller through Open-Register.cmd, then complete adoption before the other decisions. Fill yellow cells in Word and keep the field labels unchanged.', '', '| Packet | Approver form | Controller review |', '| --- | --- | --- |']
    for kind, record_id in settings['records'].items():
        lines.append(f'| {kind.title()} | [Open approval form](records/{record_id}/approval-form.docx) | [Open controller review](records/{record_id}/controller-review.docx) |')
    lines += ['', 'Each actual signed PDF goes into its corresponding record folder with the exact name `signed-approval.pdf`. Only adoption also requires the real `authority-evidence.pdf`.', '', 'The software fills record IDs, scope and contract hashes. The appointed people supply actual decisions, signing times, expiry, authority evidence and verification observations. No approval is inferred from a prepared form.']
    index = BASE / 'FORMS.md'
    if not index.exists(): index.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(json.dumps({'createdForms': made, 'signaturesCreated': False}, indent=2))

def fields(path):
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read('word/document.xml'))
    reverse = {value: key for key, value in LABELS.items()}
    result = {}
    for row in xml.findall('.//w:tr', ns):
        cells = row.findall('w:tc', ns)
        if len(cells) != 2: continue
        text = [''.join(node.text or '' for node in cell.findall('.//w:t', ns)).strip() for cell in cells]
        key = reverse.get(text[0])
        if key:
            if key in result: raise ValueError('Duplicate form field: ' + key)
            result[key] = None if not text[1] or text[1].startswith('[FILL]') else text[1]
    return result

def extract():
    settings = load(BASE / 'register.json')
    imported = []
    for kind, record_id in settings['records'].items():
        folder = BASE / 'records' / record_id
        if load(folder / 'signature-verification.json').get('signature'): continue
        for docname, jsonname in [('approval-form.docx', 'approver-input.json'), ('controller-review.docx', 'verification-input.json')]:
            value = fields(folder / docname)
            previous = load(folder / jsonname)
            expected = set(previous)
            if kind != 'adoption' and docname == 'approval-form.docx': expected -= {'controllerCertificateSha256', 'retentionDays', 'backupLocation'}
            if set(value) != expected: raise ValueError(f'Preserve all field labels in {record_id}/{docname}.')
            previous.update(value)
            (folder / jsonname).write_text(json.dumps(previous, indent=2) + '\n', encoding='utf-8')
        imported.append(record_id)
    print(json.dumps({'imported': imported, 'approvalGranted': False}))

def prefill_controller():
    settings = load(BASE / 'register.json')
    enrolled = load(BASE / 'controllers/controller-public.json')
    for kind, record_id in settings['records'].items():
        folder = BASE / 'records' / record_id
        if load(folder / 'signature-verification.json').get('signature'): continue
        targets = [('controller-review.docx', LABELS['reviewerName'], enrolled['name'])]
        if kind == 'adoption': targets.append(('approval-form.docx', LABELS['controllerCertificateSha256'], enrolled['certificateSha256']))
        for filename, label, value in targets:
            doc = Document(folder / filename)
            changed = False
            for t in doc.tables:
                for row in t.rows:
                    if row.cells[0].text == label and row.cells[1].text == '[FILL]':
                        row.cells[1].text = value; changed = True
            if changed: doc.save(folder / filename)
    print('Controller name and fingerprint inserted into unsealed forms.')

def qa():
    import sys
    from pdf2image import convert_from_path
    settings = load(BASE / 'register.json')
    poppler = next((Path(sys.executable).parents[1] / 'native/poppler').rglob('pdftoppm.exe')).parent
    out = ROOT / '.codex-tmp/x1-register-previews'
    out.mkdir(parents=True, exist_ok=True)
    images = []
    for kind, record_id in settings['records'].items():
        for stem in ('approval-form', 'controller-review'):
            pdf = BASE / 'records' / record_id / (stem + '-unsigned.pdf')
            for number, page in enumerate(convert_from_path(str(pdf), dpi=100, poppler_path=str(poppler)), 1):
                image = out / f'{kind}-{stem}-{number}.png'
                page.save(image)
                images.append(str(image))
    print(json.dumps({'pages': len(images), 'images': images}, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['create', 'extract', 'prefill-controller', 'qa'])
    action = parser.parse_args().action
    {'create': create, 'extract': extract, 'prefill-controller': prefill_controller, 'qa': qa}[action]()
