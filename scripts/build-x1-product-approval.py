"""Create an unsigned product decision form and a frozen review bundle."""
from pathlib import Path
import hashlib
import json
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'X1'
OUT.mkdir(parents=True, exist_ok=True)
STEM = 'X1_PRODUCT_PERFORMANCE_APPROVAL_FOR_SIGNATURE_2026-09-05'
BUNDLE = OUT / 'X1_PRODUCT_APPROVAL_REVIEW_BUNDLE_2026-09-05.zip'
DOC_REF = 'X1-PRODUCT-FORM-20260905-v2'
CONTRACT_MANIFEST = OUT / 'X1_CONTRACT_ARTIFACT_MANIFEST_2026-09-05.json'
RECEIPT = OUT / 'X1_PRODUCT_APPROVAL_DRAFT_RECEIPT_2026-09-05.json'
LOCAL_RUN_ID = 'LOCAL-X1-20260905093056'
LOCAL_MANIFEST_PATH = f'what-next/evidence/finance-x1/local-simulation/{LOCAL_RUN_ID}/manifest.local-simulation.json'
SUMMARY_PATH = ROOT / 'what-next/evidence/finance-x1/X1_LOCAL_EVIDENCE_AND_APPROVAL_HANDOFF_2026-09-05.md'
for output in (BUNDLE, CONTRACT_MANIFEST, RECEIPT, OUT / (STEM + '.docx'), SUMMARY_PATH):
    if output.exists():
        raise FileExistsError(f'Preserve the existing approval artifact; choose a new version: {output}')
SOURCE_PATHS = [
    'components/DataTableComponents/table-system-contract.ts',
    'components/DataTableComponents/finance-accounting-contracts.ts',
    'components/DataTableComponents/next-phase-contracts.ts',
    'components/DataTableComponents/finance-x1-evidence-binding.ts',
    'tests/e2e/finance-x1-route-observation.ts',
    'tests/e2e/finance-x1-authenticated-evidence.spec.ts',
    'tests/e2e/finance-x1-local-authenticated-evidence.spec.ts',
    'tests/e2e/finance-x1-local-evidence-aggregate.setup.ts',
    'components/finance/FinanceCommandCenterDashboard.tsx',
    'components/ui/scroll-area.tsx',
    'components/dashboard/history/TransactionHistoryWorkbenchShell.tsx',
    'components/DataTableComponents/__tests__/FinanceX1RouteObservation.test.ts',
    'what-next/development-progression-agent/finance-investigation-2026-09-02/INVESTIGATION.md',
    'what-next/evidence/finance-x1/FINANCE_X1_APPROVAL_AND_EVIDENCE_HANDOFF_2026-08-31.md',
    LOCAL_MANIFEST_PATH,
]
SOURCE_BYTES = {p: (ROOT / p).read_bytes() for p in SOURCE_PATHS}
local_manifest = json.loads(SOURCE_BYTES[LOCAL_MANIFEST_PATH])
assert local_manifest['runId'] == LOCAL_RUN_ID
assert local_manifest['authenticated'] is True
assert local_manifest['authoritative'] is False
assert local_manifest['approvedForReleaseEvidence'] is False
assert local_manifest['productionBudgetsApproved'] is False
assert local_manifest['artifactCount'] == len(local_manifest['artifacts']) == 42
assert local_manifest['routeCount'] == 7
local_artifacts = local_manifest['artifacts']
expected_checks = ('routeLoaded', 'authorizationVerified', 'tenantIsolationVerified',
                   'redactionBoundaryVerified', 'layoutUsable', 'accessibilityReviewed')
assert all(all(a['checks'].get(check) is True for check in expected_checks) for a in local_artifacts)
coverage = []
for viewport in ('desktop', 'tablet', 'mobile'):
    for locale in ('en', 'fr'):
        matching = [a for a in local_artifacts if a['viewport'] == viewport and a['locale'] == locale]
        assert len(matching) == len({a['route'] for a in matching}) == 7
        timings = [a['measurements']['usefulResultMs'] for a in matching]
        coverage.append({'viewport': viewport, 'locale': locale, 'routeCount': len(matching),
                         'usefulResultMsMin': min(timings), 'usefulResultMsMax': max(timings),
                         'maxRenderedRows': max(a['measurements']['renderedRowCount'] for a in matching)})
local_summary = {
    'runId': LOCAL_RUN_ID, 'classification': local_manifest['classification'],
    'capturedAt': local_manifest['capturedAt'], 'artifactCount': len(local_artifacts),
    'approvedForReleaseEvidence': False, 'productionBudgetsApproved': False,
    'coverage': coverage,
    'measurementLimits': 'Development-server route observations; not repeated p75 interaction, cached-navigation, dataset-scale, or export-duration acceptance measurements.',
}
scope = {
    'schemaVersion': 1, 'preparedOn': '2026-09-05', 'documentReference': DOC_REF,
    'approvalStatus': 'UNSIGNED_NO_DECISION', 'approvalRecordId': None,
    'budgetContractVersion': 1, 'scope': 'X1-finance-accounting-performance-budgets',
    'environment': 'non-production', 'fixtureTenant': 'rds_org_001',
    'files': [{'path': p, 'sha256': hashlib.sha256(b).hexdigest()} for p, b in SOURCE_BYTES.items()],
    'localEvidence': local_summary,
    'notes': [
        'Frozen review copy; source files retain their own historical status. This bundle creates no approval.',
        'The legacy 2026-08-31 contract manifest has a stale evidence-binding file hash; this review bundle records current bytes.',
        'Budget thresholds remain proposed and unverified. The route matrix alone does not measure every interaction or dataset tier.',
        'The form reference is not an approval-system record ID. Actual approval and signature references are supplied after execution.',
        'The September 3 partially completed form is preserved. This version includes the September 5 accessibility and mobile-observation fixes and needs a new signed decision.',
        'The proposed 900-day validity is carried forward from the unsigned September 3 form for confirmation or revision by the approver.',
    ],
}
with zipfile.ZipFile(BUNDLE, 'x', zipfile.ZIP_DEFLATED) as z:
    z.writestr('REVIEW_SCOPE.json', json.dumps(scope, indent=2) + '\n')
    for p, b in SOURCE_BYTES.items():
        z.writestr(p, b)
BUNDLE_HASH = hashlib.sha256(BUNDLE.read_bytes()).hexdigest()
contract_manifest = {
    'schemaVersion': 1, 'artifactId': 'stoquify-x1-table-finance-contract-bundle-v3',
    'artifactUri': 'repository://stoquify/' + CONTRACT_MANIFEST.relative_to(ROOT).as_posix(),
    'scope': 'X1 finance/accounting budgets, table contracts, accessibility fixes, evidence harness, and local capture results',
    'hashAlgorithm': 'SHA-256', 'source': BUNDLE.name,
    'reviewBundleUri': 'repository://stoquify/' + BUNDLE.relative_to(ROOT).as_posix(),
    'reviewBundleSha256': BUNDLE_HASH,
    'approvalStatus': 'pending-authenticated-human-approval', 'approvalSystemRecordId': None,
    'files': [{'uri': 'repository://stoquify/' + item['path'], **item} for item in scope['files']],
    'localEvidence': local_summary,
    'supersedesForNewApprovalRequests': 'repository://stoquify/docs/X1/X1_CONTRACT_ARTIFACT_MANIFEST_2026-09-03.json',
    'notes': scope['notes'],
}
with CONTRACT_MANIFEST.open('x', encoding='utf-8') as output:
    output.write(json.dumps(contract_manifest, indent=2) + '\n')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

def style(name, size, color='111111', bold=False, before=0, after=6, spacing=1.10):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name, s.font.size, s.font.bold = 'Calibri', Pt(size), bold
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.line_spacing = spacing
    s.paragraph_format.widow_control = True
    return s

style('Normal', 11)
style('Title', 24, '17365D', True, after=6)
style('Subtitle', 12, '555555', after=10)
style('Heading 1', 16, '2E74B5', True, 16, 8)
style('Heading 2', 13, '2E74B5', True, 12, 6)
style('Heading 3', 12, '1F4D78', True, 8, 4)
# Named form overrides: compact repeated fields, table cells, and technical hash.
style('Form metadata', 10, '555555', after=3)
style('Form field', 11, after=12)
style('Form note', 9.5, '555555', after=6)
style('Budget cell', 10, after=3, spacing=1.05)
style('Budget header', 10, '17365D', True, after=3, spacing=1.05)
style('Table citation', 9, '555555', before=4, after=4)
style('Technical hash', 8.5, '333333', after=6).font.name = 'Consolas'
style('Header', 9, '555555', after=0)
style('Footer', 9, '555555', after=0)
for n in ('Heading 1', 'Heading 2', 'Heading 3'):
    doc.styles[n].paragraph_format.keep_with_next = True
sec.header.paragraphs[0].text = 'STOQUIFY  /  X1 PRODUCT DECISION'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run(DOC_REF + '  |  Page ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def p(text, sty=None):
    return doc.add_paragraph(text, style=sty)

def labeled(label, text='', sty='Form field'):
    para = doc.add_paragraph(style=sty)
    para.add_run(label + ': ').bold = True
    para.add_run(text or '________________________________________________')
    return para

def table(headers, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.autofit = False
    pr = t._tbl.tblPr
    for name, attrs in [('tblW', {'w': str(sum(widths)), 'type': 'dxa'}), ('tblInd', {'w': '120', 'type': 'dxa'}), ('tblLayout', {'type': 'fixed'})]:
        el = pr.find(qn('w:' + name))
        if el is None: el = OxmlElement('w:' + name); pr.append(el)
        for k, v in attrs.items(): el.set(qn('w:' + k), v)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + side)
        for k, v in {'val':'single', 'sz':'4', 'color':'B8C2CC'}.items(): el.set(qn('w:' + k), v)
        borders.append(el)
    pr.append(borders)
    grid = t._tbl.tblGrid
    for el in list(grid): grid.remove(el)
    for width in widths:
        el = OxmlElement('w:gridCol'); el.set(qn('w:w'), str(width)); grid.append(el)
    for i, values in enumerate([headers, *rows]):
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(values):
            cell = t.cell(i, j)
            cell.width = Inches(widths[j] / 1440)
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn('w:tcW')); tcw.set(qn('w:w'), str(widths[j])); tcw.set(qn('w:type'), 'dxa')
            margins = OxmlElement('w:tcMar')
            for k, v in {'top':80, 'bottom':80, 'start':120, 'end':120}.items():
                el = OxmlElement('w:' + k); el.set(qn('w:w'), str(v)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            if i == 0:
                el = OxmlElement('w:shd'); el.set(qn('w:fill'), 'F2F4F7'); tcpr.append(el)
            cell.paragraphs[0].style = doc.styles['Budget header' if i == 0 else 'Budget cell']
            cell.paragraphs[0].text = value
    return t

# Page 1: the complete decision subject, with no preselected decision.
p('X1 product performance approval', 'Title')
p('Review the proposed budgets, then complete and sign page 2.', 'Subtitle')
labeled('Form reference', DOC_REF + '  |  Prepared 05 September 2026', 'Form metadata')
labeled('Status', 'Prepared for signature; no decision has been recorded.', 'Form metadata')
labeled('Scope', 'X1-finance-accounting-performance-budgets  |  Budget contract v1', 'Form metadata')
p('This form records the product authority\'s decision on the targets below for Stoquify\'s X1 finance/accounting workflow. It becomes an approval only through the authorized decision and verified signed record described on pages 2 and 3.')
doc.add_heading('Proposed engineering budgets', 1)
table(['Signal', 'Target proposed for approval'], [
    ('Local sort / filter', 'Interaction response under 100 ms on p75 test hardware, using a bounded client dataset.'),
    ('Server filter / sort', 'Loading indication under 150 ms; useful result under 1,000 ms at p75.'),
    ('Page navigation', 'Useful rows under 750 ms at p75 after a cached shell.'),
    ('Bulk action acknowledgement', 'Immediate job receipt, with progress and partial-result handling.'),
    ('Preference / view hydration', 'No visible column-layout flash after the initial shell.'),
    ('Browser row DOM', 'Bounded by pagination or a proven virtualization strategy.'),
    ('Export duration', 'Measured and retained by dataset tier and export format.'),
], [2800, 6560])
p('Source: current next-phase-contracts.ts and the X1 approval/evidence handoff, preserved in the review bundle identified on page 3.', 'Table citation')
doc.add_heading('Measurement and scope', 2)
p('The September 5 local run produced all 42 EN/FR screenshots after accessibility and mobile-observation fixes. Its 43 checks passed. This demonstrates local route coverage; performance-budget acceptance still requires the specified measurements. Route timings exclude audit and screenshot overhead and remain development-server observations.')
p('Scope: non-production; database stoquify_replay_next_20260830; tenant rds_org_001; isolation probe rds_org_002. Seven finance/accounting routes are covered in English and French on desktop, tablet, and mobile. Product approval does not itself authorize credential use, data capture, or release.', 'Form note')

# Page 2: practical decision and signature, all actual acts left blank.
doc.add_page_break()
p('Decision and signature', 'Title')
p('Complete this page as the appointed product authority.', 'Subtitle')
labeled('Form reference', DOC_REF, 'Form metadata')
p('Select exactly one decision:', 'Form field')
p('[  ] APPROVE as written     [  ] REJECT     [  ] REVISE', 'Form field')
p('By selecting APPROVE and signing below, I approve budget contract version 1 and the scope stated in this form, including the exact review bundle on page 3. I confirm that I hold the product decision authority identified below. Approval authorizes evaluation against these budgets; it does not establish that the implementation meets them.')
p('Proposed validity: 900 calendar days from the actual signed decision, carried forward from the unsigned September 3 form. Confirm or revise before signing. Approval ends earlier on revocation or a material change to the contract, routes, fixture scope, targets, or measurement method.', 'Form note')
labeled('Full legal / professional name')
labeled('Organization and title')
labeled('Product authority appointment / delegation reference')
labeled('Approving user ID / identity mapping')
p('The approving identity must be verified for the product_approver role when binding the manifest. A seeded persona name or role assignment alone does not establish your authority.', 'Form note')
labeled('Actual signed decision timestamp', '____________________________  (ISO-8601 with timezone)')
p('Copy the authenticated signing event time. For an accepted manual-signature process, the register must retain verification of the actual signing time.', 'Form note')
labeled('Signature', '________________________________________________')
p('Use the accepted electronic-signature workflow, or a manual signature with the verification required by your organization\'s controlled approval register.', 'Form note')
labeled('Reason for REJECT / REVISE, if applicable')
p('____________________________________________________________________', 'Form field')
p('For REVISE, issue revised controlled text for a new decision. No option selected, an unverified signer, or an unsigned form leaves the product approval pending.', 'Form note')

# Page 3: prevent circular or invented record references; exact intake mapping.
doc.add_page_break()
p('Register the signed record', 'Title')
p('For the document controller or approver after signing.', 'Subtitle')
p('A particular commercial approval platform is not prescribed here. Your organization must recognize the register and its signature-verification process. A saved file or shared-folder link alone does not establish authenticated approval.')
labeled('Approval register / controlled system')
labeled('Issued approval record ID')
labeled('Signature verification / event / certificate reference')
labeled('Permanent signed-record URI')
p('The URI is the reviewer-accessible link to the completed signed record and its verification evidence. Retain the approved version, its signing history, and any later revocation. Store the final signed PDF and compute its hash after signing; editing it afterwards changes that hash.', 'Form note')
doc.add_heading('Copy into the existing X1 intake', 2)
table(['Intake field', 'Value to copy'], [
    ('1. Product decision', 'The selected decision from page 2.'),
    ('2. Product approval reference', 'The actual record ID issued by the recognized approval register.'),
    ('3. Product approval timestamp', 'The actual verified signing timestamp, including timezone.'),
    ('4. Product signature reference', 'The issued signature-verification, event, or certificate reference.'),
], [3250, 6110])
p('For each SOURCE RECORD URI, use the permanent link above if that record exposes the corresponding value and verification evidence. Use a separate audit/certificate link where needed. This draft form reference is not a substitute for the issued approval ID.', 'Form note')
doc.add_heading('Exact material submitted for review', 2)
p(BUNDLE.name, 'Form metadata')
p('SHA-256: ' + BUNDLE_HASH, 'Technical hash')
p('This bundle freezes the contracts, current accessibility fixes, capture definitions, investigation, and September 5 local evidence manifest. Register this form and bundle together. The September 3 form is preserved as historical input; its incomplete decision does not apply to this revised package.', 'Form note')
p('After verified product approval, the separate credential / tenant / seeded-data authorization, fresh authenticated evidence, remaining performance measurements, and independent review must still pass. Missing or failed evidence keeps X1 blocked.', 'Form note')

doc.core_properties.title = 'Stoquify X1 Product Performance Approval'
doc.core_properties.subject = 'Unsigned decision form for product approval; no approval or signature is inferred'
doc.core_properties.author = 'Stoquify'
doc.core_properties.comments = 'standard_business_brief; memo_masthead; named Form/Budget/Technical styles are explicit form overrides.'
with (OUT / (STEM + '.docx')).open('xb') as output:
    doc.save(output)

# Structural preset and integrity checks before rendering.
assert sec.page_width.twips == 12240 and sec.page_height.twips == 15840
assert sec.left_margin.twips == sec.right_margin.twips == 1440
for t in doc.tables:
    widths = [int(c.get(qn('w:w'))) for c in t._tbl.tblGrid]
    assert sum(widths) == 9360
    assert t._tbl.tblPr.find(qn('w:tblInd')).get(qn('w:w')) == '120'
    for row in t.rows:
        assert [int(c._tc.get_or_add_tcPr().find(qn('w:tcW')).get(qn('w:w'))) for c in row.cells] == widths
receipt = {'documentReference': DOC_REF, 'status': 'UNSIGNED_NO_DECISION', 'reviewBundle': BUNDLE.name, 'reviewBundleSha256': BUNDLE_HASH, 'unsignedDocxSha256': hashlib.sha256((OUT / (STEM + '.docx')).read_bytes()).hexdigest(), 'sourceFileCount': len(SOURCE_BYTES), 'renderQa': 'pending'}
receipt['contractManifest'] = CONTRACT_MANIFEST.name
receipt['contractManifestSha256'] = hashlib.sha256(CONTRACT_MANIFEST.read_bytes()).hexdigest()
receipt['localEvidenceRunId'] = LOCAL_RUN_ID
with RECEIPT.open('x', encoding='utf-8') as output:
    output.write(json.dumps(receipt, indent=2) + '\n')
coverage_rows = '\n'.join(
    f"| {c['viewport']} | {c['locale']} | {c['routeCount']} | {c['usefulResultMsMin']} to {c['usefulResultMsMax']} | {c['maxRenderedRows']} |"
    for c in coverage
)
handoff = f'''# X1 local evidence and approval handoff

Prepared 2026-09-05. Product approval and authorization for release evidence remain pending.

The completed local authenticated run `{LOCAL_RUN_ID}` passed 43 Playwright checks, including aggregation of 42 screenshots. It covers seven finance/accounting routes in English and French on desktop, tablet, and mobile. The regression test run for route observation passed 13 tests. The local application process was stopped after capture.

## Captured scope

Database: `stoquify_replay_next_20260830`. Tenant: `rds_org_001`. Foreign tenant probe: `rds_org_002`. Persona: Carine Manga, `external_accountant`, user `rds_org_001_compliance_user_external_accountant`. These identifiers describe the fixture scope; the approval record must authorize any release evidence use.

Source: `{LOCAL_MANIFEST_PATH}`. Aggregation timestamp: `{local_manifest['capturedAt']}`. Classification: `{local_manifest['classification']}`. All artifacts retain `approvedForReleaseEvidence: false`.

| Viewport | Locale | Screenshots | Observed useful result time in ms | Maximum rendered rows |
| --- | --- | ---: | --- | ---: |
{coverage_rows}

These are observed development-server route times. Each cell contains seven different routes, not repeated samples for one performance signal. The table does not establish p75 budget acceptance, server-filter latency, cached page navigation, export duration, or scale-tier performance.

## Corrections represented in this package

- The cash-flow view, location, and period controls have localized accessible names.
- Scroll areas can receive keyboard focus.
- Mobile history cards carry row markers recognized by the evidence observer.
- The regression test covers mobile cards when desktop rows are hidden.

The first full run failed on missing control names, scroll-region focus, and mobile-row observation. The corrected run above covers all 42 combinations. The automated checks are limited to what the harness tests; independent review and any remaining performance work still apply.

## Material presented for the next product decision

Form: `docs/X1/{STEM}.docx`.

Review bundle: `docs/X1/{BUNDLE.name}`. SHA-256: `{BUNDLE_HASH}`.

Contract URI: `{contract_manifest['artifactUri']}`. Manifest SHA-256: `{receipt['contractManifestSha256']}`.

The September 3 form and bundle are retained. That form has an incomplete decision without a registered approval ID or verifiable signature; it does not authorize this revised code. Its proposed 900-day validity is carried forward for the approver to confirm or revise. The new form leaves decision, signature, appointment reference, and real signing timestamp for the authorized person.

## Inputs and next actions

| Item | Current status | Next action |
| --- | --- | --- |
| Local EN/FR capture | Complete for 42 combinations | Retain the local manifest and screenshots under their original timestamps. |
| Current contract package | Refreshed | Review and sign the September 5 package. |
| Product decision | Pending | Supply the decision, registered approval ID, actual signed timestamp, and signature reference. |
| Evidence authorization | Pending | Supply the security decision, credential approval ID, seeded-data approval ID, actual signed timestamp, and signature reference. |
| Approval automation | Requires configured authority and provider | Use the existing local profile for development; supply recognized approval records before release binding. |
| Release evidence | Pending authorization | Capture a new run after both effective approval timestamps; never relabel or backdate the local screenshots. |
| Performance acceptance | Not demonstrated by the route matrix | Evaluate the agreed signals and dataset scope with the approved measurement method. |

The nine human fields remain in `docs/X1/X1_FINANCE_ACCOUNTING_APPROVAL_INTAKE_PREFILLED_2026-09-01.docx`. IDs and timestamps must come from the recognized approval process. Automation can transfer issued values after they exist.
'''
with SUMMARY_PATH.open('x', encoding='utf-8') as output:
    output.write(handoff)
print(json.dumps(receipt, indent=2))
