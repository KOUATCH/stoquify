from __future__ import annotations

import argparse
import hashlib
import json
import os
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "docs" / "blockers-and-gates"
TARGET_DOCX = DOC_DIR / "Compliance authorization validation.docx"
CREDENTIAL_DOCX = DOC_DIR / "Development database roles and passwords.docx"
SNAPSHOT_JSON = ROOT / "what-next" / "hris-payroll-compliance-reseed-2026-08-20.json"
CREDENTIAL_JSON = ROOT / ".seed-artifacts" / "seed-login-credentials.json"
SOURCE_REGISTER_JSON = DOC_DIR / "COMPLIANCE_AUTHORIZATION_SUPPLIED_SOURCE_2026-08-20.json"
EVIDENCE_REGISTER_JSON = DOC_DIR / "COMPLIANCE_AUTHORIZATION_EVIDENCE_REGISTER_2026-08-20.json"
EXTERNAL_HANDOFF_MD = DOC_DIR / "COMPLIANCE_AUTHORIZATION_EXTERNAL_EVIDENCE_REQUEST_2026-08-20.md"
ARTIFACT_MANIFEST = DOC_DIR / "COMPLIANCE_AUTHORIZATION_REMEDIATION_ARTIFACTS_2026-08-20.sha256"

GATE_REPORTS = {
    "countryPackDevelopment": ROOT / "what-next" / "statutory-country-pack-development-readiness.json",
    "countryPackIntegration": ROOT / "what-next" / "statutory-country-pack-integration-readiness.json",
    "countryPackProduction": ROOT / "what-next" / "statutory-country-pack-production-readiness.json",
    "qualifiedReviewPreflight": ROOT / "what-next" / "AQSTOQFLOW_COMPLIANCE_AUTHORIZATION_REMEDIATION_REVIEW_PREFLIGHT_2026-08-20.json",
    "countryAdapterPilot": ROOT / "what-next" / "country-adapter-pilot-readiness.json",
}

ALLOWED_CLASSIFICATIONS = {
    "REPOSITORY_VERIFIED",
    "DATABASE_DEVELOPMENT_EVIDENCE",
    "SUPPLIED_UNVERIFIED_STATEMENT",
    "SYNTHETIC_DEVELOPMENT_PERSONA",
    "REQUIRES_HUMAN_REVIEW",
    "REQUIRES_QUALIFIED_REVIEW",
    "BLOCKED_EXTERNAL",
    "NOT_APPLICABLE",
    "REJECTED_AS_UNSUPPORTED",
}

REQUIRED_EVIDENCE_REQUEST_FIELDS = {
    "accountableOwner",
    "requiredEvidence",
    "authoritativeSource",
    "acceptanceTest",
    "developmentImpact",
    "integrationImpact",
    "productionImpact",
    "engineeringCompletable",
}

NAVY = "15324B"
BLUE = "DCEAF7"
GREEN = "DDF4E4"
AMBER = "FFF0C2"
RED = "F9D7D7"
GRAY = "E7EBEF"


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def repository_path(path: Path) -> str:
    try:
        value = path.relative_to(ROOT)
    except ValueError:
        value = path
    return str(value).replace("\\", "/")


def write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    os.replace(temporary, path)


def write_text_atomic(path: Path, content: str) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(content, encoding="utf-8", newline="\n")
    os.replace(temporary, path)


def validate_source_register(register: dict[str, Any]) -> None:
    if register.get("schemaVersion") != "1.0.0":
        raise ValueError("Compliance authorization source register must use schemaVersion 1.0.0.")

    claims = register.get("claims")
    evidence_requests = register.get("evidenceRequests")
    source_artifacts = register.get("sourceArtifacts")
    if not isinstance(claims, list) or not claims:
        raise ValueError("Compliance authorization source register must contain claims.")
    if not isinstance(evidence_requests, dict) or not evidence_requests:
        raise ValueError("Compliance authorization source register must contain evidenceRequests.")
    if not isinstance(source_artifacts, list) or not source_artifacts:
        raise ValueError("Compliance authorization source register must contain sourceArtifacts.")

    claim_ids: set[str] = set()
    for claim in claims:
        claim_id = claim.get("id")
        if not isinstance(claim_id, str) or not claim_id:
            raise ValueError("Every compliance authorization claim must have a stable id.")
        if claim_id in claim_ids:
            raise ValueError(f'Duplicate compliance authorization claim id "{claim_id}".')
        claim_ids.add(claim_id)

        classification = claim.get("sourceClassification")
        disposition = claim.get("disposition")
        if classification not in ALLOWED_CLASSIFICATIONS:
            raise ValueError(f'Claim "{claim_id}" has invalid sourceClassification "{classification}".')
        if disposition not in ALLOWED_CLASSIFICATIONS:
            raise ValueError(f'Claim "{claim_id}" has invalid disposition "{disposition}".')
        request_id = claim.get("evidenceRequestId")
        if request_id not in evidence_requests:
            raise ValueError(f'Claim "{claim_id}" references missing evidence request "{request_id}".')

    for request_id, request in evidence_requests.items():
        if not isinstance(request, dict):
            raise ValueError(f'Evidence request "{request_id}" must be an object.')
        missing = sorted(REQUIRED_EVIDENCE_REQUEST_FIELDS - set(request))
        if missing:
            raise ValueError(f'Evidence request "{request_id}" is missing: {", ".join(missing)}.')

    for artifact in source_artifacts:
        path_value = artifact.get("path")
        expected_hash = artifact.get("sha256")
        if not isinstance(path_value, str) or not path_value:
            raise ValueError("Every source artifact must have a repository-relative path.")
        if path_value.replace("\\", "/") == repository_path(TARGET_DOCX):
            raise ValueError("The generated compliance DOCX cannot be its own source artifact.")
        artifact_path = ROOT / path_value
        if not artifact_path.exists():
            raise FileNotFoundError(f'Source artifact "{path_value}" does not exist.')
        if not isinstance(expected_hash, str) or sha256(artifact_path) != expected_hash.lower():
            raise ValueError(f'Source artifact hash mismatch for "{path_value}".')


def load_source_register(path: Path = SOURCE_REGISTER_JSON) -> dict[str, Any]:
    register = load_json(path)
    validate_source_register(register)
    return register


def enriched_claims(register: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for claim in register["claims"]:
        rows.append({**claim, **register["evidenceRequests"][claim["evidenceRequestId"]]})
    return rows


def load_gate_statuses() -> list[dict[str, Any]]:
    statuses: list[dict[str, Any]] = []
    for gate_id, path in GATE_REPORTS.items():
        if not path.exists():
            statuses.append({
                "id": gate_id,
                "status": "MISSING_EVIDENCE",
                "generatedAt": None,
                "blockers": [repository_path(path)],
                "source": repository_path(path),
            })
            continue
        report = load_json(path)
        summary = report.get("summary") if isinstance(report.get("summary"), dict) else {}
        statuses.append({
            "id": gate_id,
            "status": summary.get("status") or report.get("status") or "UNKNOWN",
            "generatedAt": summary.get("generatedAt") or report.get("generatedAt"),
            "blockers": report.get("blockers") or [],
            "source": repository_path(path),
        })
    return statuses


def build_evidence_register(
    snapshot: dict[str, Any],
    source_register: dict[str, Any],
    gate_statuses: list[dict[str, Any]],
) -> dict[str, Any]:
    gate_by_id = {gate["id"]: gate for gate in gate_statuses}
    development_ready = gate_by_id.get("countryPackDevelopment", {}).get("status") == "READY_FOR_DEVELOPMENT_TESTING"
    integration_ready = gate_by_id.get("countryPackIntegration", {}).get("status") == "READY_FOR_CORE_INTEGRATION"
    production_status = gate_by_id.get("countryPackProduction", {}).get("status")
    return {
        "schemaVersion": "1.0.0",
        "generatedAt": snapshot["generatedAt"],
        "classification": "DEVELOPMENT_EVIDENCE_NOT_PRODUCTION_AUTHORIZATION",
        "sourceRegister": {
            "path": repository_path(SOURCE_REGISTER_JSON),
            "sha256": sha256(SOURCE_REGISTER_JSON),
        },
        "sourceArtifacts": source_register["sourceArtifacts"],
        "classifications": source_register["classificationDefinitions"],
        "claims": enriched_claims(source_register),
        "identityConflicts": source_register["identityConflicts"],
        "gateStatuses": gate_statuses,
        "assertions": {
            "developmentTestingPermitted": development_ready,
            "coreIntegrationPermitted": integration_ready,
            "productionAuthorized": False,
            "statutoryOrFiscalCertified": False,
            "liveAuthoritySubmissionPermitted": False,
            "observedProductionGateStatus": production_status,
        },
        "nonClaims": source_register["nonClaims"],
    }


def build_external_handoff(evidence_register: dict[str, Any]) -> str:
    lines = [
        "# Compliance authorization external and qualified-review evidence request — 2026-08-20",
        "",
        "Classification: **UNAPPROVED HANDOFF TEMPLATE — NOT AUTHORIZATION EVIDENCE**",
        "",
        "Complete these requests only through the stated authoritative source. Engineering must not prefill decisions, signatures, qualifications, appointments, regulator attestations, or legal conclusions.",
        "",
    ]
    emitted: set[str] = set()
    for claim in evidence_register["claims"]:
        request_id = claim["evidenceRequestId"]
        if request_id in emitted or claim["disposition"] in {"REPOSITORY_VERIFIED", "DATABASE_DEVELOPMENT_EVIDENCE", "NOT_APPLICABLE"}:
            continue
        emitted.add(request_id)
        lines.extend([
            f'## {request_id.replace("_", " ").title()}',
            "",
            f'- Related claims: {", ".join(item["field"] for item in evidence_register["claims"] if item["evidenceRequestId"] == request_id)}',
            f'- Accountable owner: {claim["accountableOwner"]}',
            f'- Required evidence: {claim["requiredEvidence"]}',
            f'- Authoritative source: {claim["authoritativeSource"]}',
            f'- Acceptance test: {claim["acceptanceTest"]}',
            f'- Development effect: {claim["developmentImpact"]}',
            f'- Integration effect: {claim["integrationImpact"]}',
            f'- Production effect: {claim["productionImpact"]}',
            f'- Engineering-completable: {"yes" if claim["engineeringCompletable"] else "no"}',
            "- Submitted evidence path: ____________________",
            "- Submitted evidence SHA-256: ____________________",
            "- Signer/reviewer stable subject ID: ____________________",
            "- Signature/approval method and timestamp: ____________________",
            "- Independent checker stable subject ID and later timestamp: ____________________",
            "- Independent checker result: ____________________",
            "",
        ])
    return "\n".join(lines).rstrip() + "\n"


def text(value: Any, fallback: str = "UNRESOLVED — no authoritative value available") -> str:
    if value is None:
        return fallback
    if isinstance(value, bool):
        return "YES" if value else "NO"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    rendered = str(value).strip()
    return rendered or fallback


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def set_cell(cell, value: Any, *, bold: bool = False, color: RGBColor | None = None, size: float = 7.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text(value))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: Iterable[Iterable[Any]], *, font_size: float = 7.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True, color=RGBColor(255, 255, 255), size=font_size)
        set_cell_shading(table.rows[0].cells[index], NAVY)
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            set_cell(row.cells[index], value, size=font_size)
            if len(table.rows) % 2 == 0:
                set_cell_shading(row.cells[index], "F6F8FA")
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(document: Document, title: str, level: int = 1) -> None:
    heading = document.add_heading(title, level=level)
    heading.style.font.name = "Aptos Display"
    heading.style.font.color.rgb = RGBColor(21, 50, 75)


def add_notice(document: Document, message: str, fill: str = AMBER) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), fill)
    set_cell(table.cell(0, 0), message, bold=True, size=9)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def configure_document(document: Document, *, landscape: bool = True) -> None:
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    for current_section in document.sections:
        header = current_section.header.paragraphs[0]
        header.text = "STOQUIFY — DEVELOPMENT COMPLIANCE EVIDENCE"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(7)
        header.runs[0].font.color.rgb = RGBColor(90, 100, 110)
        footer = current_section.footer.paragraphs[0]
        footer.text = "Synthetic local-development data. Database records do not constitute external authorization or certification."
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(7)
        footer.runs[0].font.color.rgb = RGBColor(120, 40, 40)


def configure_core_properties(document: Document, generated_at: str) -> None:
    properties = document.core_properties
    properties.title = "Compliance authorization validation"
    properties.subject = "Development compliance evidence and unresolved authorization handoff"
    properties.author = "Stoquify Compliance Evidence Generator"
    properties.last_modified_by = "Stoquify Compliance Evidence Generator"
    properties.comments = "Machine-generated development evidence. Not legal, statutory, fiscal, accounting, regulator, or production approval."
    properties.category = "Development compliance evidence"
    properties.keywords = "Stoquify, compliance, development evidence, fail-closed, Cameroon country pack"
    properties.revision = 1
    try:
        timestamp = datetime.fromisoformat(generated_at.replace("Z", "+00:00"))
        properties.created = timestamp
        properties.modified = timestamp
    except ValueError:
        properties.created = datetime.now(timezone.utc)
        properties.modified = properties.created


def word_lock_path(path: Path) -> Path:
    return path.with_name(f"~${path.name[2:]}")


def safe_output_path(path: Path) -> tuple[Path, bool]:
    if word_lock_path(path).exists():
        return path.with_name(f"{path.stem}.candidate{path.suffix}"), True
    return path, False


def normalize_docx_package(path: Path, generated_at: str) -> None:
    try:
        generated = datetime.fromisoformat(generated_at.replace("Z", "+00:00"))
    except ValueError:
        generated = datetime(1980, 1, 1, tzinfo=timezone.utc)
    zip_timestamp = (
        max(1980, generated.year),
        generated.month,
        generated.day,
        generated.hour,
        generated.minute,
        generated.second - (generated.second % 2),
    )
    normalized = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for entry in sorted(source.infolist(), key=lambda item: item.filename):
            info = zipfile.ZipInfo(entry.filename, date_time=zip_timestamp)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = entry.external_attr
            info.create_system = entry.create_system
            target.writestr(info, source.read(entry.filename))
    os.replace(normalized, path)


def build_compliance_document(
    snapshot: dict[str, Any],
    source_register: dict[str, Any],
    evidence_register: dict[str, Any],
    target_path: Path,
) -> None:
    document = Document()
    configure_document(document)
    configure_core_properties(document, snapshot["generatedAt"])
    title = document.add_heading("Compliance authorization validation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Reproducible development evidence, authorization boundary, and external-review handoff")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.color.rgb = RGBColor(21, 50, 75)
    add_notice(
        document,
        "CONTROL NOTICE: Every populated database field below is sourced from the local development database. "
        "Unsigned signatures, professional qualifications, production authorization and regulator attestations remain explicitly unresolved. "
        "No seed row is represented as external approval.",
    )

    add_heading(document, "1. Executive decision summary")
    gate_rows = [
        [gate["id"], gate["status"], gate.get("generatedAt"), ", ".join(gate.get("blockers") or []) or "NONE", gate["source"]]
        for gate in evidence_register["gateStatuses"]
    ]
    add_table(document, ["Boundary", "Observed status", "Generated at", "Blockers", "Evidence source"], gate_rows, font_size=6.8)
    assertions = evidence_register["assertions"]
    add_table(document, ["Decision", "Current result", "Meaning"], [
        ["Development testing", "PERMITTED" if assertions["developmentTestingPermitted"] else "NOT PROVEN", "Controlled development and sandbox testing only."],
        ["Core integration", "PERMITTED" if assertions["coreIntegrationPermitted"] else "NOT PROVEN", "Non-production integration with legal-effect paths disabled."],
        ["Production authorization", "BLOCKED", "No verified production approval; production gate remains fail-closed."],
        ["Statutory/fiscal certification", "BLOCKED", "Qualified review and applicable external evidence remain incomplete."],
        ["Live authority submission", "BLOCKED", "No live DGI/MINFI/CNPS submission authority or conformance is claimed."],
    ])

    add_heading(document, "2. Scope and non-claims")
    add_notice(document, "Development readiness is not production, statutory, fiscal, payroll, accounting, regulator, privacy, security, accessibility, or release certification.", RED)
    add_table(document, ["Non-claim"], [[item] for item in evidence_register["nonClaims"]], font_size=7.5)

    add_heading(document, "3. Document control")
    add_table(document, ["Field", "Value", "Authoritative source"], [
        ["Generated at", snapshot["generatedAt"], "Generation timestamp"],
        ["Generated document route", repository_path(target_path), "Repository artifact route"],
        ["Immutable source register", repository_path(SOURCE_REGISTER_JSON), "Hash-verified structured transcription"],
        ["Immutable source register SHA-256", sha256(SOURCE_REGISTER_JSON), "Source register bytes"],
        ["Machine-readable evidence register", repository_path(EVIDENCE_REGISTER_JSON), "Generated evidence/status contract"],
        ["External evidence handoff", repository_path(EXTERNAL_HANDOFF_MD), "Unapproved fillable handoff"],
        ["Artifact hash manifest", repository_path(ARTIFACT_MANIFEST), "Generated after document creation"],
        ["Database host", snapshot["dataSource"]["host"], "Resolved local DATABASE_URL (credentials excluded)"],
        ["Database port", snapshot["dataSource"]["port"], "Resolved local DATABASE_URL"],
        ["Database name", snapshot["dataSource"]["database"], "Resolved local DATABASE_URL"],
        ["Database schema", snapshot["dataSource"]["schema"], "Resolved local DATABASE_URL"],
        ["Environment", snapshot["dataSource"]["environment"], "NODE_ENV"],
        ["Database target fingerprint", snapshot["dataSource"]["fingerprint"], "SHA-256 of non-secret connection coordinates"],
    ])

    add_heading(document, "4. Source artifacts and evidence classifications")
    add_table(document, ["Source artifact", "SHA-256", "Classification", "Authority limitation"], [
        [artifact["path"], artifact["sha256"], artifact["classification"], artifact["authorityLimitation"]]
        for artifact in source_register["sourceArtifacts"]
    ], font_size=6.6)
    add_table(document, ["Classification", "Meaning"], [
        [item["code"], item["meaning"]] for item in source_register["classificationDefinitions"]
    ], font_size=7)

    add_heading(document, "5. Supplied claims and current disposition")
    add_table(document, ["Claim", "Supplied value", "Source class", "Current disposition", "Owner", "Development effect", "Production effect"], [
        [claim["field"], claim["value"], claim["sourceClassification"], claim["disposition"], claim["accountableOwner"], claim["developmentImpact"], claim["productionImpact"]]
        for claim in evidence_register["claims"]
    ], font_size=5.9)

    add_heading(document, "6. Identity and authority conflicts")
    add_notice(document, "Conflicting names are not automatically merged. Stable subject IDs and authoritative appointments are required before approval credit is possible.", AMBER)
    add_table(document, ["Conflict", "Observed values", "Status", "Reason", "Required resolution"], [
        [conflict["id"], " | ".join(conflict["values"]), conflict["status"], conflict["reason"], conflict["requiredResolution"]]
        for conflict in evidence_register["identityConflicts"]
    ], font_size=6.4)

    add_heading(document, "7. Development organizations and module context")
    add_table(document, ["Organization ID", "Name", "Trade name", "Country", "Currency", "Timezone", "Locale", "Requested modules"], [
        [organization["id"], organization["name"], organization.get("tradeName"), organization.get("country") or organization.get("countryCode"), organization["currency"], organization["timezone"], organization["defaultLocale"], ", ".join(organization["requestedModules"])]
        for organization in snapshot["organizations"]
    ], font_size=7)

    add_heading(document, "8. Supplied identity-to-database reconciliation")
    add_table(document, ["Supplied name", "Application role", "Stable user ID", "Payroll employee ID", "Database name", "Classification", "Signature", "Qualification"], [
        [row["suppliedName"], row["roleCode"], row["userId"], row["employeeId"], row["databaseName"], row["sourceClassification"], "NOT CLAIMED", "NOT CLAIMED"]
        for row in snapshot["suppliedIdentityMapping"]
    ], font_size=6.8)

    add_heading(document, "9. Complete database role, user and HRIS coverage")
    add_notice(
        document,
        "The rows below prove application RBAC and HRIS context only. Governance appointments, qualifications, signatures and external authorization require separate authentic evidence.",
        BLUE,
    )
    role_rows = []
    for row in snapshot["roleCoverage"]:
        user = row.get("user") or {}
        employee = row.get("employee") or {}
        assignment = employee.get("assignment") or {}
        position = assignment.get("position") or {}
        contract = employee.get("contract") or {}
        role_rows.append([
            row["organizationId"],
            row["roleCode"],
            row["roleName"],
            user.get("name"),
            user.get("email"),
            user.get("id"),
            employee.get("id"),
            position.get("id"),
            assignment.get("id"),
            contract.get("id"),
            row.get("sourceClassification"),
            row.get("authorityStatus"),
        ])
    add_table(document, ["Org", "Role code", "Role", "Person", "Email", "User ID", "Employee ID", "Position ID", "Assignment ID", "Contract ID", "Source class", "Authority status"], role_rows, font_size=5.8)

    add_heading(document, "10. Pilot location, terminal and drawer")
    pilot = snapshot["pilot"]
    add_table(document, ["Object", "Stable ID", "Name", "Database details", "Source"], [
        ["Location", pilot["location"]["id"], pilot["location"]["name"], pilot["location"], "Database locations table"],
        ["Terminal", pilot["terminal"]["id"], pilot["terminal"]["name"], pilot["terminal"], "Database pos_terminals table"],
        ["Cash drawer", pilot["drawer"]["id"], pilot["drawer"]["name"], pilot["drawer"], "Database cash_drawers table"],
    ], font_size=6.8)

    add_heading(document, "11. HRIS, payroll, accounting and compliance record counts")
    add_table(document, ["Database model/domain", "Row count", "Source"], [
        [key, value, "Local development database snapshot"] for key, value in snapshot["counts"].items()
    ])

    add_heading(document, "12. Referential integrity and segregation of duties")
    integrity = snapshot["integrity"]
    integrity_rows = [
        ["Expected compliance role rows", integrity["expectedComplianceRoleRows"], "PASS" if integrity["expectedComplianceRoleRows"] == integrity["actualComplianceRoleRows"] else "FAIL"],
        ["Actual compliance role rows", integrity["actualComplianceRoleRows"], "DATABASE"],
        ["Roles without users", len(integrity["rolesWithoutUsers"]), "PASS" if not integrity["rolesWithoutUsers"] else "FAIL"],
        ["Roles without HRIS employees", len(integrity["rolesWithoutEmployees"]), "PASS" if not integrity["rolesWithoutEmployees"] else "FAIL"],
        ["Employees without active contracts", len(integrity["employeesWithoutActiveContracts"]), "PASS" if not integrity["employeesWithoutActiveContracts"] else "FAIL"],
        ["Employees without primary assignments", len(integrity["employeesWithoutPrimaryAssignments"]), "PASS" if not integrity["employeesWithoutPrimaryAssignments"] else "FAIL"],
        ["Employees without active schedules", len(integrity["employeesWithoutActiveSchedules"]), "PASS" if not integrity["employeesWithoutActiveSchedules"] else "FAIL"],
        ["Cross-tenant employment assignments", integrity["crossTenantEmploymentAssignmentCount"], "PASS" if integrity["crossTenantEmploymentAssignmentCount"] == 0 else "FAIL"],
        ["Self-reporting relationships", integrity["selfReportingRelationshipCount"], "PASS" if integrity["selfReportingRelationshipCount"] == 0 else "FAIL"],
    ]
    add_table(document, ["Control", "Observed value", "Result"], integrity_rows)
    sod_rows = []
    for pair in integrity["segregationOfDutiesPairs"]:
        for organization in pair["organizations"]:
            sod_rows.append([organization["organizationId"], pair["makerRole"], organization["makerUserId"], pair["checkerRole"], organization["checkerUserId"], organization["result"]])
    add_table(document, ["Organization", "Maker/preparer role", "Maker user ID", "Checker/approver role", "Checker user ID", "Result"], sod_rows, font_size=6.8)

    add_heading(document, "13. Database and governance gate disposition")
    gate_rows = []
    for row in snapshot["gateMatrix"]:
        gate_rows.append([row["gate"], row["status"], row["source"]])
    add_table(document, ["Gate / blocker", "Status", "Right source / reason"], gate_rows)

    add_heading(document, "14. Qualified-review and country-pack production boundary")
    add_table(document, ["Boundary", "Observed status", "Blockers", "Result interpretation"], [
        [gate["id"], gate["status"], ", ".join(gate.get("blockers") or []) or "NONE", "Development/integration evidence only; no production or legal approval inferred."]
        for gate in evidence_register["gateStatuses"]
    ], font_size=6.8)

    add_heading(document, "15. Unresolved external or human-review evidence")
    add_table(document, ["Unresolved claim", "Disposition", "Owner", "Required evidence", "Acceptance test", "Engineering-completable"], [
        [claim["field"], claim["disposition"], claim["accountableOwner"], claim["requiredEvidence"], claim["acceptanceTest"], "YES" if claim["engineeringCompletable"] else "NO"]
        for claim in evidence_register["claims"]
        if claim["disposition"] not in {"REPOSITORY_VERIFIED", "DATABASE_DEVELOPMENT_EVIDENCE", "NOT_APPLICABLE"}
    ], font_size=5.8)

    add_heading(document, "16. Credential and password handling")
    reconciliation = snapshot["credentialReconciliation"]
    add_table(document, ["Field", "Value", "Source"], [
        ["Credential artifact", reconciliation["passwordSource"], "Local Git-ignored seed artifact"],
        ["Total login credentials", reconciliation["artifactCredentialCount"], "Credential artifact"],
        ["Required role credentials", reconciliation["requiredRoleCredentialCount"], "Credential artifact reconciled to database users"],
        ["Credentials missing from database", reconciliation["missingFromDatabase"] or "NONE", "Database reconciliation"],
        ["Passwords included in this compliance document", "NO", "Security boundary"],
        ["Password document", "docs/blockers-and-gates/Development database roles and passwords.docx", "Local Git-ignored handoff"],
    ])

    add_heading(document, "17. Controlled promotion sequence")
    add_table(document, ["Step", "Owner", "Required result"], [
        ["1. Engineering remediation", "Engineering and QA", "Development and integration gates pass; production paths remain fail-closed."],
        ["2. Qualified review", "Appointed qualified reviewers", "Identity, capacity, sources, fixture decisions, tie-outs, limitations, signature, and review window are complete."],
        ["3. Independent check", "Checker independent of implementer/reviewer", "Every retained hash, signature, appointment, decision, and timestamp is independently verified."],
        ["4. External conformance", "Authority liaison, security, and integration owners", "Applicable DGI/MINFI/CNPS contract, credentials, sandbox conformance, and operational controls pass."],
        ["5. Production gate", "Release authority", "All mandatory production gates pass for one frozen candidate without override."],
        ["6. Controlled activation", "Maker-checker release owners", "Bounded rollout, monitoring, kill switch, and rollback are approved and evidenced."],
    ], font_size=6.8)

    add_heading(document, "Appendix A — Supplied statements retained for traceability")
    add_notice(document, "Supplied text is retained as evidence of what was stated. Its current disposition controls; no supplied YES, typed name, or empty signature field is silently upgraded to approval.", GRAY)
    add_table(document, ["Field", "Supplied value", "Source", "Current disposition", "Notes"], [
        [claim["field"], claim["value"], claim["sourceArtifact"], claim["disposition"], claim["notes"]]
        for claim in source_register["claims"]
    ], font_size=6.2)

    DOC_DIR.mkdir(parents=True, exist_ok=True)
    temporary = target_path.with_suffix(".tmp.docx")
    document.save(temporary)
    normalize_docx_package(temporary, snapshot["generatedAt"])
    os.replace(temporary, target_path)


def build_credential_document(snapshot: dict[str, Any], credential_artifact: dict[str, Any]) -> None:
    document = Document()
    configure_document(document)
    title = document.add_heading("Development database roles and passwords", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_notice(
        document,
        "HIGHLY SENSITIVE DEVELOPMENT-ONLY CREDENTIALS. This DOCX is explicitly Git-ignored. "
        "Never commit, email, publish, reuse in production, or treat these personas as real external authorities.",
        RED,
    )
    add_heading(document, "1. Credential control")
    add_table(document, ["Field", "Value", "Source"], [
        ["Generated at", datetime.now(timezone.utc).isoformat(), "Document generator"],
        ["Credential artifact generated at", credential_artifact.get("generatedAt"), ".seed-artifacts/seed-login-credentials.json"],
        ["Database", snapshot["dataSource"]["database"], "Resolved local development target"],
        ["Database schema", snapshot["dataSource"]["schema"], "Resolved local development target"],
        ["Credential count", credential_artifact.get("credentialCount"), "Credential artifact"],
        ["Password source", ".seed-artifacts/seed-login-credentials.json", "Canonical seed artifact; passwords are not recoverable from database hashes"],
        ["Database verification", "Every required role email reconciled to an active database user", "Compliance seed snapshot"],
        ["Repository protection", "/docs/blockers-and-gates/Development database roles and passwords.docx", ".gitignore"],
    ])

    add_heading(document, "2. All development login personas")
    rows = []
    for credential in sorted(credential_artifact["credentials"], key=lambda item: (item["organizationId"], item["role"], item["email"])):
        rows.append([
            credential["organizationId"],
            credential["organizationName"],
            credential["name"],
            credential["role"],
            ", ".join(credential.get("roleCodes", [])),
            credential["email"],
            credential["password"],
            credential["loginRoute"],
            credential["locale"],
            credential.get("intendedScenario"),
            credential.get("note") or "SYNTHETIC DEVELOPMENT PERSONA",
        ])
    add_table(document, ["Org ID", "Organization", "Name", "Primary role", "Role codes", "Email", "Password", "Login route", "Locale", "Scenario", "Control note"], rows, font_size=5.7)

    add_heading(document, "3. Required role inventory")
    add_table(document, ["Organization", "Role code", "Role name", "User", "Email", "Active", "Verified", "Authority limitation"], [
        [row["organizationId"], row["roleCode"], row["roleName"], (row.get("user") or {}).get("name"), (row.get("user") or {}).get("email"), (row.get("user") or {}).get("isActive"), (row.get("user") or {}).get("isVerified"), row["authorityStatus"]]
        for row in snapshot["roleCoverage"]
    ], font_size=6.3)

    temporary = CREDENTIAL_DOCX.with_suffix(".tmp.docx")
    document.save(temporary)
    os.replace(temporary, CREDENTIAL_DOCX)


def write_artifact_manifest(paths: Iterable[Path]) -> None:
    rows = [f"{sha256(path)}  {repository_path(path)}" for path in paths]
    write_text_atomic(ARTIFACT_MANIFEST, "\n".join(rows) + "\n")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the Stoquify compliance authorization evidence packet.")
    parser.add_argument(
        "--include-credential-document",
        action="store_true",
        help="Also read the Git-ignored credential artifact and regenerate the separate password document.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    snapshot = load_json(SNAPSHOT_JSON)
    source_register = load_source_register()
    gate_statuses = load_gate_statuses()
    evidence_register = build_evidence_register(snapshot, source_register, gate_statuses)
    write_json_atomic(EVIDENCE_REGISTER_JSON, evidence_register)
    write_text_atomic(EXTERNAL_HANDOFF_MD, build_external_handoff(evidence_register))

    output_docx, lock_detected = safe_output_path(TARGET_DOCX)
    build_compliance_document(snapshot, source_register, evidence_register, output_docx)
    manifest_paths = [
        output_docx,
        SOURCE_REGISTER_JSON,
        EVIDENCE_REGISTER_JSON,
        EXTERNAL_HANDOFF_MD,
        *(ROOT / artifact["path"] for artifact in source_register["sourceArtifacts"]),
    ]
    write_artifact_manifest(manifest_paths)

    result: dict[str, Any] = {
        "complianceDocument": repository_path(output_docx),
        "complianceDocumentSha256": sha256(output_docx),
        "canonicalDocumentLockDetected": lock_detected,
        "canonicalDocumentUpdated": output_docx == TARGET_DOCX,
        "sourceRegister": repository_path(SOURCE_REGISTER_JSON),
        "sourceRegisterSha256": sha256(SOURCE_REGISTER_JSON),
        "evidenceRegister": repository_path(EVIDENCE_REGISTER_JSON),
        "evidenceRegisterSha256": sha256(EVIDENCE_REGISTER_JSON),
        "externalEvidenceHandoff": repository_path(EXTERNAL_HANDOFF_MD),
        "artifactManifest": repository_path(ARTIFACT_MANIFEST),
        "credentialDocumentGenerated": False,
        "passwordsPrinted": False,
    }

    if args.include_credential_document:
        credentials = load_json(CREDENTIAL_JSON)
        build_credential_document(snapshot, credentials)
        result.update({
            "credentialDocumentGenerated": True,
            "credentialDocument": repository_path(CREDENTIAL_DOCX),
            "credentialDocumentSha256": sha256(CREDENTIAL_DOCX),
            "credentialDocumentGitIgnored": True,
        })

    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
