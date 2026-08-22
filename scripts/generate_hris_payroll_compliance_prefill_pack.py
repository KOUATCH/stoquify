"""Generate the verified-only HRIS/payroll compliance prefill evidence pack.

The generator reads repository documents plus the redacted read-only database
snapshot. It never changes the original compliance document, frozen G1 contract,
live approval register, or database records.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts import generate_landing_innovation_pdfs as pdf_base


DATE = "2026-08-19"
GENERATED_AT = "2026-08-19T18:30:00Z"
OUT = ROOT / "docs" / "blockers-and-gates" / "hris-payroll-compliance-prefill"

ORIGINAL_DOCX = ROOT / "docs" / "Compliance" / "Complaince authorization validation.docx"
BASE_WORKING_DOCX = ROOT / "docs" / "blockers-and-gates" / "COMPLIANCE_AUTHORIZATION_G1_PREFILLED_WORKING_COPY_2026-08-19.docx"
OUTPUT_DOCX = OUT / f"COMPLIANCE_AUTHORIZATION_G1_EVIDENCE_RECONCILED_WORKING_COPY_{DATE}.docx"
OUTPUT_PDF = OUTPUT_DOCX.with_suffix(".pdf")
SNAPSHOT_PATH = OUT / f"HRIS_PAYROLL_DATABASE_READ_ONLY_SNAPSHOT_{DATE}.json"
PREFILL_STATUS_PATH = ROOT / "docs" / "blockers-and-gates" / "COMPLIANCE_AUTHORIZATION_G1_PREFILL_STATUS_2026-08-19.json"
FINALIZATION_PATH = ROOT / "docs" / "blockers-and-gates" / "G1_33_OBLIGATION_FINALIZATION_REGISTER_2026-08-19.json"
MATRIX_PATH = ROOT / "docs" / "blockers-and-gates" / "cameroon-hris-payroll-authority" / f"CAMEROON_EMPLOYMENT_DOCUMENT_REQUIREMENTS_MATRIX_{DATE}.json"
CONTRACT_PATH = ROOT / "docs" / "pos-enterprise-grade-audit" / "EXECUTION_06_G1_CONTRACT_FREEZE_V0_2_0.json"
LIVE_REGISTER_PATH = ROOT / "docs" / "pos-enterprise-grade-audit" / "EXECUTION_06_G1_CONTRACT_APPROVAL_REGISTER.json"

CONTRACT_SHA256 = "11434eb3e1af1826516426e90d2d53a2faa47ade361d91191b5f3e0c950a36db"
UNRESOLVED = "UNRESOLVED — GOVERNANCE EVIDENCE REQUIRED"
CONFLICTED = "CONFLICTED — INDEPENDENT RESOLUTION REQUIRED"
WARNING = "WORKING COPY — NOT AUTHORITY, SIGNATURE, APPROVAL OR RELEASE EVIDENCE"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    return path.resolve().relative_to(ROOT).as_posix()


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.rstrip() + "\n", encoding="utf-8")


def write_json(path: Path, value: Any) -> None:
    write_text(path, json.dumps(value, indent=2, ensure_ascii=False))


def esc(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", "<br>")


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    lines = [
        "| " + " | ".join(esc(header) for header in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(esc(cell) for cell in row) + " |" for row in rows)
    return "\n".join(lines)


def source_paths() -> list[Path]:
    paths: set[Path] = {
        ORIGINAL_DOCX,
        BASE_WORKING_DOCX,
        SNAPSHOT_PATH,
        PREFILL_STATUS_PATH,
        FINALIZATION_PATH,
        MATRIX_PATH,
        CONTRACT_PATH,
        LIVE_REGISTER_PATH,
        ROOT / "docs" / "blockers-and-gates" / "G1_NAMES_ROLES_AUTHORITY_EVIDENCE_MATRIX_2026-08-19.md",
        ROOT / "docs" / "blockers-and-gates" / "G1_APPROVAL_FIELD_DICTIONARY_AND_VALIDATION_RULES_2026-08-19.md",
        ROOT / "docs" / "blockers-and-gates" / "G1_UNRESOLVED_AUTHORITY_AND_APPROVAL_BLOCKERS_2026-08-19.json",
        ROOT / "docs" / "blockers-and-gates" / "G1_DECISION_BY_DECISION_COMPLETION_RUNBOOK_2026-08-19.md",
        ROOT / "docs" / "blockers-and-gates" / "COMPLIANCE_AUTHORIZATION_G1_UNRESOLVED_FIELDS_2026-08-19.json",
        ROOT / "prisma" / "schema.prisma",
        ROOT / "prisma" / "migrations" / "20260719190000_hris_org_manager_scope_foundation" / "migration.sql",
        ROOT / "services" / "hris" / "org.service.ts",
        ROOT / "lib" / "security" / "auth-session.ts",
        ROOT / "services" / "security" / "step-up-auth.service.ts",
        ROOT / "scripts" / "pos-g1-contract-gate.js",
        ROOT / "graphify-out" / "GRAPH_REPORT_actions.md",
        ROOT / "graphify-out" / "GRAPH_REPORT_types.md",
    }
    authority_pack = ROOT / "docs" / "blockers-and-gates" / "cameroon-hris-payroll-authority"
    if authority_pack.exists():
        paths.update(path for path in authority_pack.rglob("*") if path.is_file())
    blocker_root = ROOT / "docs" / "blockers-and-gates"
    name_pattern = re.compile(r"(?:COMPLIANCE_AUTHORIZATION|^G1_).*(?:\.md|\.json|\.docx|\.pdf|\.sha256)$", re.I)
    paths.update(path for path in blocker_root.iterdir() if path.is_file() and name_pattern.search(path.name))
    return sorted(path for path in paths if path.is_file())


def classify_source(path: Path) -> tuple[str, str, str]:
    value = rel(path)
    name = path.name.lower()
    if "/sources/" in f"/{value}" or name.startswith(("cameroon-labour", "cnps-", "dgi-", "ohada-")):
        return (
            "REGULATOR_OR_LEGAL_SOURCE_LEGAL_REVIEW_REQUIRED",
            "Statutory requirement and regulator provenance",
            "Does not prove a Stoquify identity, role appointment or approval",
        )
    if path == CONTRACT_PATH:
        return ("VERIFIED_REPOSITORY_CONTRACT", "Exact G1 decision contract binding", "Does not prove human authority or approval")
    if path == LIVE_REGISTER_PATH:
        return ("LIVE_REGISTER_REPOSITORY_TRUTH", "Current approval-register state", "Empty/pending fields are not approval evidence")
    if path == SNAPSHOT_PATH:
        return ("SYSTEM_DERIVED_READ_ONLY_SNAPSHOT", "Current tenant/identity/HRIS data state", "Application records alone do not prove legal employer or governance authority")
    if path.suffix.lower() in {".prisma", ".sql", ".ts", ".js"}:
        return ("REPOSITORY_TECHNICAL_TRUTH", "Schema, validation and runtime behavior", "Does not prove a real person or human approval")
    if "working_copy" in name or "completion_form" in name:
        return ("CANDIDATE_WORKING_DOCUMENT", "Form structure and candidate assertions", "Not independent identity, employment, authority or approval evidence")
    if path == ORIGINAL_DOCX:
        return ("IMMUTABLE_SOURCE_DOCUMENT", "Original compliance-document content", "Signature images and typed names are not authenticated evidence")
    return ("DERIVED_ASSESSMENT_OR_TEMPLATE", "Cross-reference, field definition or unresolved-state evidence", "Must not be treated as independent primary evidence")


def build_source_inventory() -> tuple[list[dict[str, Any]], str]:
    records = []
    for path in source_paths():
        classification, use, boundary = classify_source(path)
        records.append({
            "path": rel(path),
            "extension": path.suffix.lower(),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
            "classification": classification,
            "usedFor": use,
            "authorityBoundary": boundary,
        })
    counts = Counter(record["classification"] for record in records)
    markdown = f"""# HRIS/payroll compliance source inventory

Prepared {DATE}  
Classification: **EVIDENCE INVENTORY — NOT AUTHORITY OR APPROVAL EVIDENCE**

## Outcome

{len(records)} repository and documentary sources were inventoried and content-hashed. Repeated reports derived from the same source are not treated as independent corroboration.

{md_table(['Classification', 'Files'], [[key, value] for key, value in sorted(counts.items())])}

## Source register

{md_table(['Path', 'Class', 'SHA-256', 'Bytes', 'Can support', 'Cannot prove alone'], [[record['path'], record['classification'], record['sha256'], record['bytes'], record['usedFor'], record['authorityBoundary']] for record in records])}

## Handling rules

- Recompute hashes before relying on any source.
- Regulator materials require dated qualified Cameroon review.
- Candidate working documents never become authority merely because values are repeated.
- Database records are current application state, not corporate appointment instruments.
- The original compliance document remains unchanged.
"""
    return records, markdown


def build_identity_reconciliation(snapshot: dict[str, Any]) -> tuple[dict[str, Any], str]:
    rows = []
    for item in snapshot["identityReconciliation"]:
        top = item.get("topSimilarityCandidates", [])
        nearest = top[0] if top else None
        rows.append({
            "canonicalRole": item["canonicalRole"],
            "page6Candidate": item["candidateName"],
            "normalizedComparisonOnly": item["normalizedCandidateName"],
            "exactMatchCount": item["exactMatchCount"],
            "nearestDatabaseNameClue": nearest.get("databaseName") if nearest else None,
            "nearestSimilarity": nearest.get("similarity") if nearest else None,
            "stableSubjectId": None,
            "payrollEmployeeId": None,
            "classification": "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED",
            "reason": "No exact normalized database match; low-similarity names are rejected as identity evidence.",
            "requiredAction": "HR/security must verify the legal person and bind User.id to PayrollEmployee.id using controlled evidence.",
        })
    result = {
        "schemaVersion": "1.0.0",
        "generatedAt": snapshot["generatedAt"],
        "databaseTransaction": snapshot["transaction"],
        "summary": {
            "canonicalRoles": len(rows),
            "exactMatches": sum(row["exactMatchCount"] for row in rows),
            "verifiedIdentities": 0,
            "candidateNamesPropagated": len(rows),
        },
        "rows": rows,
    }
    markdown = f"""# HRIS/payroll name and identity reconciliation

Prepared {DATE}  
Classification: **CANDIDATE REVIEW — NOT IDENTITY OR AUTHORITY EVIDENCE**

## Result

- Page-6 candidates assessed: **{len(rows)}**
- Exact normalized database matches: **0**
- Verified stable subject IDs: **0**
- Database transaction: **{snapshot['transaction']}**

No database identity was inserted into the compliance form. The highest string-similarity clues are shown only to demonstrate that no plausible match exists; they are expressly rejected as identity evidence.

{md_table(['Canonical role', 'Page-6 candidate', 'Exact matches', 'Nearest DB clue', 'Similarity', 'Disposition'], [[row['canonicalRole'], row['page6Candidate'], row['exactMatchCount'], row['nearestDatabaseNameClue'] or 'None', row['nearestSimilarity'] if row['nearestSimilarity'] is not None else '', 'UNRESOLVED — HR/security verification required'] for row in rows])}

## Verification rule

An exact name would still be only a candidate linkage. Final identity requires tenant consistency, controlled identity evidence, a stable `User.id`, a linked `PayrollEmployee.id`, employment evidence and an independent checker.
"""
    return result, markdown


def logical_location(table_index: int) -> str:
    if table_index == 0:
        return "Control banner"
    if table_index == 1:
        return "Frozen-contract summary"
    if table_index == 2:
        return "Page 6 — canonical role roster"
    if table_index == 3:
        return "Page 7 — 33-obligation matrix"
    if 4 <= table_index <= 47:
        return "G1 decision detail sections"
    return "Existing evidence/completion appendix"


def canonical_field(header: str) -> str:
    text = re.sub(r"[^a-z0-9]+", "_", header.lower()).strip("_")
    mappings = {
        "accountable_name": "candidateAccountableName",
        "name": "candidateAccountableName",
        "stable_subject_id": "stableSubjectId",
        "authority_ref_appointment": "authorityReference",
        "authority_reference": "authorityReference",
        "effective_dates": "authorityEffectiveDates",
        "scope_effective_dates": "authorityScopeAndDates",
        "tenant_org": "tenantOrOrganizationId",
        "sod_coi": "sodCoiResult",
        "sod": "sodResult",
        "ready_to_sign": "signingEligibility",
        "auth_method_mfa": "authenticationMethodAndMfa",
        "fresh_auth_at": "freshAuthenticatedAt",
        "approved_at": "approvedAt",
        "evidence_sha_256": "signatureEvidenceSha256",
        "independent_verifier": "independentVerifier",
        "complete_sha_256": "contractSha256",
        "repository_path": "contractPath",
        "version": "contractVersion",
    }
    return mappings.get(text, text or "unlabeledField")


def classify_cell(header: str, value: str) -> tuple[str, str, str]:
    lower = value.lower()
    canonical = canonical_field(header)
    if canonical in {"contractSha256", "contractPath", "contractVersion"}:
        return ("AUTHORITATIVE_VERIFIED", rel(CONTRACT_PATH), "Recompute and compare frozen contract")
    if "candidate" in lower or canonical == "candidateAccountableName":
        return ("CANDIDATE_NOT_VERIFIED", rel(BASE_WORKING_DOCX), "Verify identity through HR/security")
    if any(marker in lower for marker in ("unresolved", "to be filled", "pass / fail", "yes / no", "not performed", "not available")):
        return ("UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "", "Obtain controlled primary evidence and independent review")
    if canonical in {"stableSubjectId", "authorityReference", "sodCoiResult", "sodResult", "freshAuthenticatedAt", "approvedAt", "signatureEvidenceSha256", "independentVerifier"}:
        return ("UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "", "No authentic evidence available")
    return ("CORROBORATED_FORM_CONTENT", rel(BASE_WORKING_DOCX), "Confirm against the controlling source")


def build_crosswalk() -> tuple[dict[str, Any], str]:
    document = Document(BASE_WORKING_DOCX)
    tables = []
    cells = []
    for table_index, table in enumerate(document.tables):
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        tables.append({
            "tableIndex": table_index,
            "logicalLocation": logical_location(table_index),
            "rows": len(table.rows),
            "columns": len(table.columns),
            "headings": headers,
        })
        for row_index, row in enumerate(table.rows[1:], start=1):
            for column_index, cell in enumerate(row.cells):
                header = headers[column_index] if column_index < len(headers) else f"Column {column_index + 1}"
                value = cell.text.strip()
                classification, source, rule = classify_cell(header, value)
                cells.append({
                    "tableIndex": table_index,
                    "logicalLocation": logical_location(table_index),
                    "rowIndex": row_index,
                    "columnIndex": column_index,
                    "heading": header,
                    "canonicalField": canonical_field(header),
                    "observedValue": value,
                    "classification": classification,
                    "sourcePath": source or None,
                    "sourceSha256": sha256(ROOT / source) if source and (ROOT / source).is_file() else None,
                    "verificationRule": rule,
                    "resolutionOwner": "HR/security" if "subject" in canonical_field(header).lower() else "Governance/evidence owner",
                })
    result = {
        "schemaVersion": "1.0.0",
        "generatedAt": GENERATED_AT,
        "sourceDocument": {"path": rel(BASE_WORKING_DOCX), "sha256": sha256(BASE_WORKING_DOCX)},
        "tableCount": len(tables),
        "cellCount": len(cells),
        "tables": tables,
        "cells": cells,
    }
    class_counts = Counter(cell["classification"] for cell in cells)
    markdown = f"""# HRIS/payroll compliance field crosswalk

Prepared {DATE}  
Classification: **FIELD PROVENANCE — NOT AUTHORITY OR APPROVAL EVIDENCE**

## Coverage

- Tables accounted for: **{len(tables)}**
- Data cells classified: **{len(cells)}**
- Full cell-level provenance is in the companion JSON artifact.

{md_table(['Classification', 'Cells'], [[key, value] for key, value in sorted(class_counts.items())])}

## Table and heading inventory

{md_table(['Table', 'Logical location', 'Rows', 'Columns', 'Headings'], [[table['tableIndex'], table['logicalLocation'], table['rows'], table['columns'], '; '.join(table['headings'])] for table in tables])}

## Population rule

Only fields classified `AUTHORITATIVE_VERIFIED` may be final-filled. Candidate names remain labelled candidates. Identity, employment, appointment, SoD, authentication, signature and verification fields remain unresolved until primary evidence exists.
"""
    return result, markdown


def build_field_provenance(snapshot: dict[str, Any]) -> dict[str, Any]:
    organizations = snapshot["organizationEvidence"]
    fields = [
        ("legalEmployerName", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "Two application tenants exist; neither is verified as the Cameroon legal employer.", "Corporate governance/legal"),
        ("tenantOrOrganizationId", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", f"Candidate application tenant IDs: {', '.join(item['id'] for item in organizations)}", "Governance/platform owner"),
        ("cnpsEmployerReference", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "No CNPS employer registration record found.", "Payroll/compliance"),
        ("dgiTaxpayerReference", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "Two masked application tax identifiers exist; no selected legal entity or DGI-issued evidence.", "Tax/legal"),
        ("hrEvidenceCustodian", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "No application role assignments or governance appointment evidence found.", "HR/governance"),
        ("governanceCustodian", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "No governance authority registry or appointment found.", "Corporate governance"),
        ("stableSubjectId", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "0/17 exact database name matches and 0 linked payroll employees.", "HR/security"),
        ("payrollEmployeeId", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "24 draft records exist, but none maps to a page-6 identity.", "HR"),
        ("employmentAssignmentId", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "No org units, positions or employment assignments exist.", "HR"),
        ("authorityReference", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "No exact G1 appointment or bounded delegation exists.", "Governance"),
        ("sodCoiResult", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "Stable identities and independent checker are absent.", "Controls reviewer"),
        ("freshAuthenticatedAt", None, "UNRESOLVED_GOVERNANCE_EVIDENCE_REQUIRED", "No fresh-assurance sessions exist in the current snapshot.", "IAM/security"),
        ("contractArtifactId", "STOQUIFY-POS-G1-CONTRACT-FREEZE-0.2.0-20260817", "AUTHORITATIVE_VERIFIED", "Frozen contract artifact.", "G1 contract owner"),
        ("contractVersion", "0.2.0", "AUTHORITATIVE_VERIFIED", "Frozen contract version.", "G1 contract owner"),
        ("contractPath", rel(CONTRACT_PATH), "AUTHORITATIVE_VERIFIED", "Repository path resolved.", "G1 contract owner"),
        ("contractSha256", sha256(CONTRACT_PATH), "AUTHORITATIVE_VERIFIED", "Recomputed content hash.", "Independent verifier"),
    ]
    return {
        "schemaVersion": "1.0.0",
        "generatedAt": snapshot["generatedAt"],
        "fields": [
            {
                "field": field,
                "populatedValue": value,
                "classification": classification,
                "evidence": evidence,
                "owner": owner,
                "sourcePath": rel(CONTRACT_PATH) if classification == "AUTHORITATIVE_VERIFIED" else rel(SNAPSHOT_PATH),
                "sourceSha256": sha256(CONTRACT_PATH) if classification == "AUTHORITATIVE_VERIFIED" else sha256(SNAPSHOT_PATH),
                "resolutionAction": "Retain verified value and rehash on use." if classification == "AUTHORITATIVE_VERIFIED" else "Obtain primary evidence and independent verification; do not guess.",
            }
            for field, value, classification, evidence, owner in fields
        ],
    }


def build_readiness(snapshot: dict[str, Any], prefill: dict[str, Any], finalization: dict[str, Any]) -> dict[str, Any]:
    candidates = {role["role"]: role["candidateName"] for role in prefill["roles"]}
    obligations = []
    for obligation in finalization["obligations"]:
        role = obligation["authority"]["approverRole"]
        obligations.append({
            "obligationId": obligation["obligationId"],
            "decisionId": obligation["contractBinding"]["decisionId"],
            "decisionTitle": obligation["contractBinding"]["decisionTitle"],
            "selectedOption": obligation["contractBinding"]["selectedOption"],
            "canonicalRole": role,
            "page6Candidate": candidates.get(role),
            "stableSubjectId": None,
            "payrollEmployeeId": None,
            "employmentAssignmentId": None,
            "authorityReference": None,
            "sodResult": None,
            "coiResult": None,
            "freshAuthenticationReference": None,
            "signatureEvidenceReference": None,
            "independentVerifier": None,
            "status": "BLOCKED_MISSING_VERIFIED_IDENTITY_EMPLOYMENT_ASSIGNMENT_AUTHORITY_SOD_AND_APPROVAL",
            "owner": obligation["operations"]["fieldOwner"],
            "resolutionAction": obligation["operations"]["resolutionAction"],
        })
    return {
        "schemaVersion": "1.0.0",
        "generatedAt": snapshot["generatedAt"],
        "classification": "READINESS_REGISTER_NOT_APPROVAL_EVIDENCE",
        "g1Passed": False,
        "productionAuthorized": False,
        "contractBinding": {
            "path": rel(CONTRACT_PATH),
            "sha256": sha256(CONTRACT_PATH),
            "expectedSha256": CONTRACT_SHA256,
            "hashMatch": sha256(CONTRACT_PATH) == CONTRACT_SHA256,
        },
        "summary": {
            "decisions": len({item["decisionId"] for item in obligations}),
            "obligations": len(obligations),
            "canonicalRoles": len(candidates),
            "verifiedIdentities": 0,
            "verifiedAssignments": 0,
            "verifiedAppointments": 0,
            "formalSodCoiResults": 0,
            "freshAuthenticatedApprovals": 0,
            "completedObligations": 0,
        },
        "databaseSnapshot": snapshot["recordCounts"],
        "obligations": obligations,
    }


def build_unresolved(snapshot: dict[str, Any], readiness: dict[str, Any]) -> tuple[dict[str, Any], str]:
    items = [
        ("EMPLOYER-01", "Confirm the exact Cameroon legal employer and establishment", "Corporate governance/legal", "Signed corporate decision plus registry evidence"),
        ("TENANT-01", "Select the application tenant that represents that legal employer", "Governance/platform owner", "Tenant-to-legal-entity binding"),
        ("CNPS-01", "Obtain the CNPS employer registration reference and current-status evidence", "Payroll/compliance", "CNPS-issued evidence, masked/hash recorded"),
        ("DGI-01", "Confirm the DGI taxpayer reference for the selected legal employer", "Tax/legal", "DGI-issued evidence reconciled to legal entity"),
        ("OWNER-01", "Appoint accountable HR and governance evidence custodians", "Corporate governance", "Stable IDs, authority references, scope and dates"),
        ("IDENTITY-01", "Verify all 17 page-6 people", "HR/security", "Controlled identity evidence and exact User.id↔PayrollEmployee.id binding"),
        ("EMPLOYMENT-01", "Activate evidence-backed employee masters", "HR", "Employment source evidence, active status and signed-document hash"),
        ("ASSIGN-01", "Create effective-dated org units, positions and assignments", "HR", "Approved assignment records; termination/suspension fail-closed"),
        ("AUTH-01", "Issue 17 exact G1 role appointments or bounded delegations", "Governance", "Issuer authority, tenant, scope, dates and evidence hash"),
        ("SOD-01", "Run independent SoD/COI and qualification checks", "Controls reviewer", "Stable-ID results and approved exceptions if any"),
        ("APPROVAL-01", "Collect 33 deliberate fresh-authenticated approvals", "Human signers/IAM", "Ten-minute window, explicit intent and exact contract hash"),
        ("VERIFY-01", "Independently resolve and rehash every evidence envelope", "Independent verifier", "Pass/fail verification record"),
    ]
    result = {
        "schemaVersion": "1.0.0",
        "generatedAt": snapshot["generatedAt"],
        "g1Passed": False,
        "items": [
            {"id": item_id, "unresolved": issue, "owner": owner, "acceptanceEvidence": evidence, "status": "OPEN"}
            for item_id, issue, owner, evidence in items
        ],
    }
    markdown = f"""# HRIS/payroll unresolved evidence register

Prepared {DATE}  
Status: **OPEN — G1 REMAINS 0/11 AND 0/33**

## Current database facts

{md_table(['Record', 'Count'], [[key, value] for key, value in snapshot['recordCounts'].items()])}

Additional facts:

- All 24 payroll employees are `DRAFT`.
- Linked `User.id` values: 0.
- Payroll contracts: 24, active contracts: 0, signed-document hashes: 0.
- Application role links: {snapshot.get('applicationRoleLinkCount', 0)}.
- Application tenant records: 2; legal-employer selection: unresolved.

## Required resolutions

{md_table(['ID', 'Unresolved requirement', 'Owner', 'Acceptance evidence'], [[item_id, issue, owner, evidence] for item_id, issue, owner, evidence in items])}

## Gate boundary

The database can provide stable identifiers only after an authentic person/employee binding exists. It cannot issue corporate authority, SoD clearance or human approval. `policy:gates` and `verify:release` remain out of scope until the actual G1 validator is eligible.
"""
    return result, markdown


def set_cell_text(cell, value: str, *, bold: bool = False, color: str | None = None, size: float = 7.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def style_table(table, header_fill: str = "15324B") -> None:
    table.style = "Table Grid"
    if table.rows:
        for cell in table.rows[0].cells:
            shade_cell(cell, header_fill)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(7)
    for row in table.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)


def append_status_table(document: Document, title: str, headers: list[str], rows: list[list[str]]) -> None:
    document.add_heading(title, level=2)
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    style_table(table)


def build_working_copy(snapshot: dict[str, Any], identity: dict[str, Any], readiness: dict[str, Any], provenance: dict[str, Any]) -> None:
    shutil.copy2(BASE_WORKING_DOCX, OUTPUT_DOCX)
    document = Document(OUTPUT_DOCX)

    first_source_paragraph = document.paragraphs[0]
    cover_title = first_source_paragraph.insert_paragraph_before()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = cover_title.add_run("STOQUIFY HRIS/PAYROLL EVIDENCE-RECONCILED WORKING COPY")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(21, 50, 75)
    cover_warning = first_source_paragraph.insert_paragraph_before()
    cover_warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning_run = cover_warning.add_run(WARNING)
    warning_run.bold = True
    warning_run.font.size = Pt(12)
    warning_run.font.color.rgb = RGBColor(156, 0, 6)
    cover_context = first_source_paragraph.insert_paragraph_before()
    cover_context.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_context.add_run(
        "Legacy source pages following this cover contain supplied names, timestamps, "
        "signature images and authorization claims. They are preserved for traceability "
        "only and receive zero identity, authority, approval, production or certification credit."
    ).font.size = Pt(11)
    cover_state = first_source_paragraph.insert_paragraph_before()
    cover_state.alignment = WD_ALIGN_PARAGRAPH.CENTER
    state_run = cover_state.add_run(
        "Current verified state: 0/17 identities; 0 HR assignments; 0 G1 appointments; "
        "0 formal SoD/COI results; 0/33 approvals."
    )
    state_run.bold = True
    state_run.font.size = Pt(12)
    state_run.font.color.rgb = RGBColor(156, 0, 6)
    cover_hash = first_source_paragraph.insert_paragraph_before()
    cover_hash.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_hash.add_run(f"Frozen G1 contract SHA-256: {sha256(CONTRACT_PATH)}").font.size = Pt(9)
    cover_break = first_source_paragraph.insert_paragraph_before()
    cover_break.add_run().add_break(WD_BREAK.PAGE)

    legacy_exact_replacements = {
        "Production authorization: YES": "UNVERIFIED LEGACY ASSERTION — NO PRODUCTION AUTHORIZATION",
        "Statutory/fiscal certification authorization: YES": "UNVERIFIED LEGACY ASSERTION — NO STATUTORY/FISCAL CERTIFICATION",
    }
    legacy_prefixes = (
        "Operator approval timestamp:",
        "Checker approval timestamp:",
        "Product approval timestamp:",
        "Controller approval timestamp:",
    )
    legacy_signature_labels = (
        "Product approver signature:",
        "DGI Signature:",
        "Minfi Signature:",
        "CNPS Signature:",
    )
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        replacement = legacy_exact_replacements.get(text)
        if replacement:
            paragraph.text = replacement
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.color.rgb = RGBColor(156, 0, 6)
        elif text.startswith(legacy_prefixes):
            paragraph.text = f"UNVERIFIED LEGACY ASSERTION — {text} — ZERO APPROVAL CREDIT"
            paragraph.runs[0].font.color.rgb = RGBColor(156, 0, 6)
        elif text in legacy_signature_labels:
            paragraph.text = f"UNVERIFIED LEGACY SIGNATURE LABEL — {text} — ZERO APPROVAL CREDIT"
            paragraph.runs[0].font.color.rgb = RGBColor(156, 0, 6)

    set_cell_text(document.tables[0].cell(0, 0), f"{WARNING}. Fresh database review: 0/17 exact identity matches; 0 employment assignments; 0 delegations; 0/33 obligations complete.", bold=True, color="9C0006", size=8)
    shade_cell(document.tables[0].cell(0, 0), "FFC7CE")
    set_cell_text(
        document.tables[1].cell(4, 1),
        "Authority evidence: 17 candidate names; 0 verified identities; 0 HR assignments; 0 G1 appointments; 0 formal SoD/COI results; 0 approvals.",
        bold=True,
        color="9C0006",
        size=7,
    )

    role_table = document.tables[2]
    for row in role_table.rows[1:]:
        set_cell_text(row.cells[3], "UNRESOLVED — NO EXACT DB MATCH", color="9C0006")
        set_cell_text(row.cells[4], "UNRESOLVED — G1 APPOINTMENT REQUIRED", color="9C0006")
        set_cell_text(row.cells[5], "UNRESOLVED — HR ASSIGNMENT ABSENT", color="9C0006")
        set_cell_text(row.cells[6], "UNRESOLVED — LEGAL EMPLOYER/TENANT BINDING REQUIRED", color="9C0006")
        set_cell_text(row.cells[7], "UNRESOLVED — STABLE ID REQUIRED", color="9C0006")
        set_cell_text(row.cells[8], "UNRESOLVED — INDEPENDENT CHECKER REQUIRED", color="9C0006")

    obligation_table = document.tables[3]
    for row in obligation_table.rows[1:]:
        set_cell_text(row.cells[4], "UNRESOLVED", color="9C0006")
        set_cell_text(row.cells[5], "UNRESOLVED", color="9C0006")
        set_cell_text(row.cells[6], "NOT ASSESSED", color="9C0006")
        set_cell_text(row.cells[7], "NO — BLOCKED", bold=True, color="9C0006")

    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    heading = document.add_heading("HRIS/payroll evidence reconciliation appendix", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph(WARNING)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].bold = True
    paragraph.runs[0].font.color.rgb = RGBColor(156, 0, 6)

    organizations = snapshot["organizationEvidence"]
    append_status_table(document, "A. Legal employer and evidence owners", ["Field", "Database/document clue", "Verified value", "Status / required action"], [
        ["Legal employer", "; ".join(f"{item['name']} ({item['id']})" for item in organizations), "", "UNRESOLVED — application tenants are not corporate authority evidence"],
        ["Tenant binding", "; ".join(item["id"] for item in organizations), "", "UNRESOLVED — governance must bind one tenant to the legal employer"],
        ["DGI taxpayer reference", "; ".join(item["taxIdentifierMasked"] or "absent" for item in organizations), "", "UNRESOLVED — obtain DGI-issued evidence for selected entity"],
        ["CNPS employer reference", "Not found", "", "UNRESOLVED — obtain CNPS-issued employer evidence"],
        ["HR evidence custodian", "No assigned application roles or authority record", "", "UNRESOLVED — governance appointment required"],
        ["Governance custodian", "No authority registry or appointment", "", "UNRESOLVED — governance appointment required"],
    ])

    append_status_table(document, "B. Read-only database snapshot", ["Record", "Count", "Meaning"], [
        [key, str(value), {
            "organizations": "Application tenants; not legal-employer proof",
            "users": "Application identities; no page-6 exact matches",
            "payrollEmployees": "All 24 are DRAFT and unlinked",
            "orgUnits": "Required HR structure absent",
            "positions": "Required HR structure absent",
            "employmentAssignments": "Effective-dated eligibility unavailable",
            "reportingRelationships": "Manager scope unavailable",
            "managerDelegations": "No bounded delegation evidence",
            "freshAssuranceSessions": "No reusable approval authentication evidence",
        }.get(key, "Current application record count")]
        for key, value in snapshot["recordCounts"].items()
    ])

    append_status_table(document, "C. Page-6 identity resolution", ["Canonical G1 role", "Page-6 candidate", "Exact database match", "Stable IDs", "Required action"], [
        [row["canonicalRole"], row["page6Candidate"], "0", "UNRESOLVED", "HR/security identity and employment verification"]
        for row in identity["rows"]
    ])

    append_status_table(document, "D. G1 readiness", ["Measure", "Result", "Disposition"], [
        ["Frozen contract SHA-256", readiness["contractBinding"]["sha256"], "VERIFIED REPOSITORY TRUTH"],
        ["Decisions represented", str(readiness["summary"]["decisions"]), "Complete structural coverage"],
        ["Role obligations represented", str(readiness["summary"]["obligations"]), "Complete structural coverage"],
        ["Verified identities", "0", "BLOCKED"],
        ["Verified appointments", "0", "BLOCKED"],
        ["Formal SoD/COI results", "0", "BLOCKED"],
        ["Fresh-authenticated approvals", "0", "BLOCKED"],
    ])

    append_status_table(document, "E. Key provenance", ["Field", "Value", "Classification", "Source"], [
        [field["field"], str(field["populatedValue"] or ""), field["classification"], field["sourcePath"]]
        for field in provenance["fields"]
    ])

    document.add_heading("F. How to complete this document", level=2)
    instructions = [
        "Corporate governance identifies the real Cameroon legal employer and binds it to one tenant using controlled evidence.",
        "HR/security verifies each person and records the exact User.id and PayrollEmployee.id; do not select a close name from the database.",
        "HR activates evidence-backed employee records, organizational units, positions and effective-dated assignments.",
        "Governance issues exact G1 role appointments or bounded delegations with issuer authority, scope and dates.",
        "An independent controls reviewer records SoD/COI and required qualifications using stable IDs.",
        "Only then may each human signer freshly authenticate and approve the exact decision bound to the frozen contract SHA-256.",
        "A different verifier resolves the exported evidence, recomputes its SHA-256 and records pass/fail before any live-register import.",
    ]
    for index, instruction in enumerate(instructions, start=1):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{index}. ").bold = True
        paragraph.add_run(instruction)

    document.add_paragraph(
        f"Detailed cell-level provenance: {rel(OUT / f'HRIS_PAYROLL_COMPLIANCE_FIELD_CROSSWALK_{DATE}.json')}"
    )
    document.add_paragraph(
        f"Read-only database snapshot: {rel(SNAPSHOT_PATH)}; SHA-256 {sha256(SNAPSHOT_PATH)}"
    )
    document.add_paragraph(
        f"Original source remained immutable: {rel(ORIGINAL_DOCX)}; SHA-256 {sha256(ORIGINAL_DOCX)}"
    )
    document.save(OUTPUT_DOCX)


def render_pdfs(paths: list[Path]) -> None:
    def decor(canvas, document) -> None:
        canvas.saveState()
        width, height = pdf_base.A4
        if document.page > 1:
            canvas.setStrokeColor(pdf_base.RULE)
            canvas.line(18 * pdf_base.mm, height - 13 * pdf_base.mm, width - 18 * pdf_base.mm, height - 13 * pdf_base.mm)
            canvas.setFont("Helvetica-Bold", 7.5)
            canvas.setFillColor(pdf_base.NAVY)
            canvas.drawString(18 * pdf_base.mm, height - 10 * pdf_base.mm, "STOQUIFY")
            canvas.setFont("Helvetica", 7.5)
            canvas.setFillColor(pdf_base.MUTED)
            canvas.drawRightString(width - 18 * pdf_base.mm, height - 10 * pdf_base.mm, "HRIS/payroll compliance evidence")
        canvas.setStrokeColor(pdf_base.RULE)
        canvas.line(18 * pdf_base.mm, 13 * pdf_base.mm, width - 18 * pdf_base.mm, 13 * pdf_base.mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(pdf_base.MUTED)
        canvas.drawString(18 * pdf_base.mm, 9 * pdf_base.mm, DATE)
        canvas.drawRightString(width - 18 * pdf_base.mm, 9 * pdf_base.mm, f"{document.page:02d}")
        canvas.restoreState()

    pdf_base.DATE = DATE
    pdf_base.page_decor = decor
    for path in paths:
        pdf_base.build_document(
            path,
            path.with_suffix(".pdf"),
            "HRIS / PAYROLL COMPLIANCE EVIDENCE",
            "Verified-only identity, employment, authority and G1 readiness reconciliation",
        )


def changelog(snapshot: dict[str, Any], inventory_count: int, crosswalk: dict[str, Any]) -> str:
    return f"""# HRIS/payroll compliance prefill changelog

Prepared {DATE}

## Changes made

- Created a separate working copy from `{rel(BASE_WORKING_DOCX)}`; the original and prior working copies were not overwritten.
- Recomputed the frozen contract hash and preserved its exact binding.
- Executed a PostgreSQL transaction with `SET TRANSACTION READ ONLY`; rolled it back after retrieval.
- Redacted emails, compensation and raw tax/social identifiers from the snapshot.
- Reconciled 17 page-6 candidates; no exact database matches were found.
- Marked identity, employment assignment, authority, SoD/COI, authentication, signature and verifier fields unresolved.
- Added legal-employer, tenant, database, identity, readiness, provenance and completion-instruction appendices.
- Inventoried and hashed {inventory_count} sources.
- Classified {crosswalk['cellCount']} cells across {crosswalk['tableCount']} tables.

## Database results

```json
{json.dumps(snapshot['recordCounts'], indent=2)}
```

All 24 payroll employees are `DRAFT`; 0 are linked to `User.id`; 0 HR assignments, delegations or fresh-assurance sessions exist.

## Safety assertions

- Database writes performed: `false`
- Live approval register modified: `false`
- Frozen G1 contract modified: `false`
- Human approvals performed: `false`
- G1 passed: `false`
- `policy:gates` run: `false`
- `verify:release` run: `false`
"""


def manifest() -> str:
    entries = []
    for path in sorted(OUT.rglob("*")):
        if path.is_file() and not path.name.endswith(".sha256"):
            entries.append(f"{sha256(path)}  {rel(path)}")
    return "\n".join(entries) + "\n"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    if "--manifest-only" in sys.argv:
        write_text(OUT / f"HRIS_PAYROLL_COMPLIANCE_PREFILL_ARTIFACTS_{DATE}.sha256", manifest())
        return

    snapshot = load_json(SNAPSHOT_PATH)
    prefill = load_json(PREFILL_STATUS_PATH)
    finalization = load_json(FINALIZATION_PATH)

    if sha256(CONTRACT_PATH) != CONTRACT_SHA256:
        raise RuntimeError("Frozen G1 contract hash drift; stop evidence preparation")

    inventory, inventory_md = build_source_inventory()
    identity, identity_md = build_identity_reconciliation(snapshot)
    crosswalk, crosswalk_md = build_crosswalk()
    provenance = build_field_provenance(snapshot)
    readiness = build_readiness(snapshot, prefill, finalization)
    unresolved, unresolved_md = build_unresolved(snapshot, readiness)

    inventory_base = OUT / f"HRIS_PAYROLL_SOURCE_INVENTORY_{DATE}"
    identity_base = OUT / f"HRIS_PAYROLL_NAME_IDENTITY_RECONCILIATION_{DATE}"
    crosswalk_base = OUT / f"HRIS_PAYROLL_COMPLIANCE_FIELD_CROSSWALK_{DATE}"
    unresolved_base = OUT / f"HRIS_PAYROLL_UNRESOLVED_EVIDENCE_REGISTER_{DATE}"

    write_json(inventory_base.with_suffix(".json"), {
        "schemaVersion": "1.0.0",
        "generatedAt": GENERATED_AT,
        "classification": "EVIDENCE_INVENTORY_NOT_AUTHORITY_OR_APPROVAL",
        "sources": inventory,
    })
    write_text(inventory_base.with_suffix(".md"), inventory_md)
    write_json(OUT / f"HRIS_PAYROLL_FIELD_PROVENANCE_REGISTER_{DATE}.json", provenance)
    write_json(identity_base.with_suffix(".json"), identity)
    write_text(identity_base.with_suffix(".md"), identity_md)
    write_json(crosswalk_base.with_suffix(".json"), crosswalk)
    write_text(crosswalk_base.with_suffix(".md"), crosswalk_md)
    write_json(OUT / f"HRIS_PAYROLL_G1_33_OBLIGATION_READINESS_{DATE}.json", readiness)
    write_json(unresolved_base.with_suffix(".json"), unresolved)
    write_text(unresolved_base.with_suffix(".md"), unresolved_md)

    build_working_copy(snapshot, identity, readiness, provenance)
    write_text(OUT / f"HRIS_PAYROLL_COMPLIANCE_PREFILL_CHANGELOG_{DATE}.md", changelog(snapshot, len(inventory), crosswalk))
    render_pdfs([
        inventory_base.with_suffix(".md"),
        identity_base.with_suffix(".md"),
        crosswalk_base.with_suffix(".md"),
        unresolved_base.with_suffix(".md"),
    ])
    write_text(OUT / f"HRIS_PAYROLL_COMPLIANCE_PREFILL_ARTIFACTS_{DATE}.sha256", manifest())
    print(json.dumps({
        "output": rel(OUT),
        "sources": len(inventory),
        "tablesCrosswalked": crosswalk["tableCount"],
        "cellsClassified": crosswalk["cellCount"],
        "roles": identity["summary"]["canonicalRoles"],
        "exactIdentityMatches": identity["summary"]["exactMatches"],
        "decisions": readiness["summary"]["decisions"],
        "obligations": readiness["summary"]["obligations"],
        "workingCopy": rel(OUTPUT_DOCX),
        "workingCopyPdfPendingExternalConversion": not OUTPUT_PDF.exists(),
    }, indent=2))


if __name__ == "__main__":
    main()
