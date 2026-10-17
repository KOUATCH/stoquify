from pathlib import Path
import json, hashlib, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[2]
VERSION = '1.0'
DATE = '7 October 2026'
sources = {
 'S1': 'what-next/platform-correctness-roadmap/2026-09-30-run-01/build-roadmap.cjs',
 'S2': 'what-next/platform-remediation/2026-09-30-run-01/execution-register.json',
 'S3': 'what-next/platform-remediation/2026-09-30-run-01/CAMEROON_FIRST_CONTINUATION_2026-10-07.md',
 'S4': 'what-next/platform-remediation/2026-09-30-run-01/CAMEROON_SOURCE_REFRESH_2026-10-07.md',
 'S5': 'docs/approval-register/STOQUIFY_APPROVAL_REGISTER_START_HERE_2026-09-03.md',
 'S6': 'docs/approval-register/STOQUIFY_APPROVAL_REGISTER_OPERATING_GUIDE_2026-09-03.md',
 'S7': 'docs/blockers-and-gates/G1_NAMES_ROLES_AUTHORITY_EVIDENCE_MATRIX_2026-08-19.md',
 'S8': 'what-next/platform-correctness-roadmap/2026-09-30-run-01/roadmap-register.json',
 'S9': 'docs/system-audit/04_IMPLEMENTATION_ROADMAP.md',
 'S10': 'docs/blockers-and-gates/cameroon-hris-payroll-authority/templates/TEMPLATE_SOD_COI_DECLARATION_2026-08-19.md',
}

decisions = [
dict(id='D01', title='Product scope and first customers', exact='Initial segment, locations, currencies, enabled modules and capability promises',
 meaning='Choose who the first supported product is for and what you will promise them. A country choice alone does not select the customer type or authorize a live launch.',
 confirmed='Cameroon first, later CEMAC expansion, pre-production/demo data and phased implementation are confirmed. Customer sector, entity type, site count, currencies and first-release modules are not yet selected. [S3]',
 recommendation='Proposed: select one defined Cameroon customer cohort and qualify online workflows first. Treat accounting and HR/payroll as separately gated capabilities. Use synthetic demonstrations until the relevant acceptance evidence is available.',
 alternatives='A narrow cohort limits test combinations and support cost. Multiple sectors/sites at launch increase configuration, training and regression coverage. Do not assume every implemented module is promised.',
 owner='Product owner and delivery lead', review='QA lead; domain owners for every promised capability', gate='R00 before the release-scope baseline; R33 still controls actual release.',
 evidence='Customer needs or pilot brief; entity and location inventory; capability matrix; current package acceptance and exclusions. [S1, S2, S8]',
 deferral='Individual capabilities may be excluded with a reason and owner. The target cohort and explicit release boundary cannot remain ambiguous when R00 is accepted.',
 questions=[('Customer and entity','Which sector, business/entity type and intended customer cohort?'),('Sites and users','Which locations, number of sites and user roles?'),('Currency scope','Which transaction, functional and reporting currencies? Refer the policy to D02.'),('Online modules','Which modules are required for the first cohort? List each.'),('Regulated capabilities','Which accounting/payroll/tax outputs are promised, conditional or excluded?'),('Delivery boundary','What is demo only, controlled pilot, later release or out of scope?')]),
dict(id='D02', title='Accounting policy choices', exact='Functional currency, rate/rounding, partial-year, close and year-end policies',
 meaning='Agree the accounting rules against which the implementation and reports will be assessed. A green test is not an accountant approving the underlying policy.',
 confirmed='Cameroon is the first jurisdiction. No functional-currency policy, rate source, rounding method, fiscal calendar or close policy has been approved in this pack. [S3, S4]',
 recommendation='Proposed: have the accounting owner and qualified accountant document the policy for the selected entity and reporting framework, using explicit examples. Keep unsupported currency/reporting cases unavailable until mappings and source amounts are verified.',
 alternatives='Single-currency scope reduces conversion cases. Multi-currency scope requires approved rate authority, dates, missing-rate handling and reproducible conversion/rounding. Close-only locking and professionally qualified close have different evidence requirements.',
 owner='Accounting owner', review='Qualified accountant; engineering verifies the implementation against approved examples', gate='R02, R06, R07 and R20 before financial implementation acceptance.',
 evidence='Entity/framework applicability; rate-source evidence; calculation examples; fiscal-year setup; journal/report tie-outs; close, reopening and correction controls. [S1, S2, S4]',
 deferral='The detailed policy may remain pending during development, with a named reviewer and due gate. Affected financial outputs cannot be represented as qualified.',
 questions=[('Entity and framework','What entity/sector and accounting/reporting framework applies?'),('Currencies and rates','Functional/reporting currencies; FX source, rate date and missing-rate rule?'),('Rounding','Precision and rounding method at line, tax, journal and report levels?'),('Financial calendar','Year start/end; first partial year and period boundaries?'),('Close and correction','What blocks close; who locks/reopens; how are corrections and reversals handled?'),('Year end','Required adjustments, opening balances and statements; attach reviewed examples.')]),
dict(id='D03', title='Offline operating limits', exact='Offline devices, storage/eviction, supported tenders, fiscal limits and recovery',
 meaning='Decide whether offline selling is in scope and the exact devices, payments and recovery limits that must be proven before admitting it.',
 confirmed='The current continuation does not qualify a complete offline pilot. No supported device/tender matrix is approved here. [S2, S3]',
 recommendation='Proposed: defer live offline admission until R04/R05 data-conservation evidence and R32 device/pilot criteria pass. Continue controlled engineering tests with synthetic data.',
 alternatives='Online-only admission has a narrower recovery surface. A controlled offline pilot adds device enrollment, storage loss, retries, uncertain payments, fiscal and support obligations. Browser storage alone is not a durability guarantee.',
 owner='POS/product owner', review='Security owner and country owner; QA and support for device/recovery evidence', gate='R04/R05 design; finalize before R32 offline admission.',
 evidence='Browser/OS/device matrix; capacity and eviction tests; replay/conservation tests; tender/provider/fiscal review; operator recovery drill. [S1, S2]',
 deferral='If offline is excluded, record that decision and prohibit offline capability promises/admission. Reopen D03 before enabling it; do not mark the capability verified.',
 questions=[('Admission choice','Defer offline, controlled pilot, or request qualified release?'),('Devices','Exact device models, OS and browser versions; enrollment/revocation owner?'),('Storage','Queue/capacity limits, eviction detection, older-client shutdown and export policy?'),('Payments','Which tenders are allowed offline; how are uncertain outcomes resolved?'),('Fiscal limits','Which local fiscal/numbering restrictions need country review?'),('Recovery','Lost device, interrupted sync, duplicate/lost-response recovery and support steps?')]),
dict(id='D04', title='Access privacy and retention', exact='Role matrix, legal holds, retention and HR/workforce disclosure',
 meaning='Specify who can see and change each class of data, what must be retained, and how sensitive workforce information is protected.',
 confirmed='Current work uses test/demo data. No production retention schedule, complete access matrix or employee disclosure policy is approved in this pack. [S3]',
 recommendation='Proposed: use minimum role-based access, explicit employee/manager/accountant scopes and controlled export permissions. Keep payroll details restricted and define legal-hold precedence before adopting deletion schedules.',
 alternatives='Broad convenience access increases disclosure risk and review burden. More granular access requires role testing and support procedures. Retention periods must follow entity/data applicability review, not a guessed universal duration.',
 owner='Privacy/records owner and HR owner', review='Application security; qualified legal/privacy advice for applicable retention and disclosure duties', gate='R08, R10, R16 and R28 before enforcement acceptance and sensitive launch.',
 evidence='Data inventory and classification; role/tenant/action matrix; employee disclosure rules; retention and hold sources; export and access-revocation tests. [S1, S2]',
 deferral='Missing policy blocks affected sensitive use or deletion. Record the exact restricted data/actions and reviewer due date; pre-production status alone does not prove every record is synthetic.',
 questions=[('Access matrix','For each role: view/change/approve/export rights and organization/employee scope?'),('Sensitive data','Who may see salary, identity, health/leave documents and bank details?'),('Self service','What can employees, managers, HR and delegated accountants access?'),('Retention','Schedule by record class, authority source, start trigger and archive owner?'),('Holds and deletion','Who imposes/releases holds; how do holds override deletion and backups?'),('Disclosure','Permitted recipients, purpose, masking, audit trail and revocation process?')]),
dict(id='D05', title='Country review and external authority', exact='Named country experts, rule versions, providers and submission authority',
 meaning='Name qualified people to review each regulated capability and identify who can authorize real provider operations or statutory submissions.',
 confirmed='Cameroon-first source research exists, with unresolved current-law interpretation and reviewer requirements. Source discovery is not statutory approval. No expert appointments or provider/submission authority are confirmed here. [S4]',
 recommendation='Proposed: review accounting, tax, payroll/labour, social insurance and fiscal/declaration capabilities separately. Bind each approval to a rule version, applicable entity/employee cases and independent expected examples.',
 alternatives='Prepare exports for authorized human review, or later qualify direct submission through a named provider/authority. These are distinct capabilities; approving one does not authorize the other.',
 owner='Country/integration leads and business approvers', review='Qualified Cameroon accounting, tax, payroll/labour and social-insurance expertise as applicable; independent technical reviewer', gate='Begin appointments at R00; complete per capability at R18/R19 before activation.',
 evidence='Qualifications and appointments; current source snapshots; rule/version/effective-date map; reviewed fixtures; provider acceptance and submission mandates. [S1, S4, S5, S6]',
 deferral='Record missing experts/sources/providers as unresolved. Keep that regulated capability or real submission unavailable; independent engineering may continue.',
 questions=[('Capability coverage','Which accounting, payroll, tax, CNPS, fiscal or declaration capabilities?'),('Experts','For each capability: real reviewer, qualification, Cameroon scope and independence?'),('Rule versions','Exact pack/rule version, sources, effective dates and customer applicability?'),('Expected results','Who independently checks normal, boundary and exception examples?'),('Providers','Provider/account/capability and sandbox/live environment; credential reference only?'),('Submission mandate','Who may prepare, approve and execute each real payment or filing?')]),
dict(id='D06', title='Optional feature priorities', exact='Which optional lot/landed-cost/assets/budget/cohort features are required',
 meaning='Choose which optional extensions are justified by an actual customer need. A promised optional feature becomes required for that cohort.',
 confirmed='No acceptance of R34–R37 optional capabilities is established by the general modernization request. [S1, S8]',
 recommendation='Proposed: assess each feature independently and defer those without a named customer decision, usable source data and measurable acceptance scenario.',
 alternatives='A bounded extension can satisfy a real sector requirement but adds implementation and assurance work. Deferral protects delivery focus; it also forbids promising that feature for the selected cohort.',
 owner='Product owner with actual customer representative', review='Relevant inventory/accounting/finance/data domain owner', gate='R34–R37 before committing implementation capacity; earlier if promised to the first cohort.',
 evidence='Customer requirement, source-data assessment, bounded feature definition, cost/effort review and acceptance example. [S1, S8]',
 deferral='Required / deferred / not applicable must be recorded for each feature, with reason and reassessment trigger. Do not treat all optional features as one yes/no decision.',
 questions=[('Traceability','Lot/serial/expiry and unit conversion: required, deferred or not applicable; why?'),('Inventory economics','Landed cost and replenishment: scope, customer need and source quality?'),('Fixed assets','Asset subledger/automation: material customer requirement and accounting owner?'),('Planning','Budgets and scenario cash forecasting: decisions supported and data required?'),('Customer analysis','Cohort/retention analytics: business purpose and identity/consent constraints?'),('Commitment','For each accepted feature: cohort, owner, acceptance example, effort and due gate?')]),
dict(id='D07', title='Operational service objectives', exact='Deployment workload, latency, error, backlog, cost, RPO/RTO and alert objectives',
 meaning='Agree the operating conditions and measurable service limits before judging operational readiness. RPO means acceptable data-loss window; RTO means acceptable recovery time.',
 confirmed='Current evidence does not qualify a fresh integrated build/package, complete browser journeys or production recovery. No measured production service objectives are approved here. [S2, S3]',
 recommendation='Proposed: first measure representative workloads in a controlled environment, then agree explicit thresholds and measurement windows with product, finance and SRE. Do not publish unmeasured service promises.',
 alternatives='Tighter latency/recovery targets generally require more capacity, recovery automation and support expenditure. A smaller bounded cohort may reduce the initial operating envelope but still needs measured acceptance.',
 owner='SRE/operations owner', review='Product owner and finance owner; independent QA/technical evidence review', gate='R23/R30 before measuring operational acceptance against agreed limits.',
 evidence='Deployment topology; workload model; reproducible baseline measurements; restore drill; worker/backlog monitoring; alert delivery and cost model. [S1, S2]',
 deferral='Leave unknown thresholds blank with a measurement owner/date. That service promise remains unqualified. Record zero silent event loss, duplicate economic effects and unauthorized access as proposed invariants, not measured achievements.',
 questions=[('Environment and workload','Hosting/topology; tenant/user concurrency, data volume and peak transactions?'),('Latency and errors','Critical journeys, percentile/window, acceptable latency and error-rate thresholds?'),('Workers','Queue age/backlog/throughput targets and escalation conditions?'),('Recovery','RPO and RTO per service/data class; how will restore be measured?'),('Cost and alerts','Budget, measurement unit/window, alert recipient and escalation response?'),('Acceptance method','Test dataset, duration, observer, evidence location and remeasurement trigger?')]),
dict(id='D08', title='Pilot admission and observation', exact='Pilot cohorts, minimum volume, business-cycle coverage and observation window',
 meaning='Define a fair pilot test in advance: participants, duration, business cycles, acceptance evidence and stop conditions.',
 confirmed='Cameroon is prioritized. No named pilot customer, transaction-volume threshold, observation window or pilot authorization has been confirmed. [S3]',
 recommendation='Proposed: admit a bounded cohort only after its applicable security, data, country and provider gates pass. Agree the pilot criteria before seeing the results; keep online and offline admission distinct.',
 alternatives='A shorter/narrower pilot costs less but cannot demonstrate unobserved business cycles. Broader admission needs more support and evidence. A synthetic rehearsal is useful but does not establish real-customer pilot performance.',
 owner='Release owner and pilot operations owner', review='Product, QA and relevant domain owners; country/provider reviewers for regulated or external capabilities', gate='Before R32/R33; criteria must not be invented after results.',
 evidence='Named cohort and permitted capabilities; signed criteria; support/rollback plan; representative business-cycle evidence; gate matrix and independent review. [S1, S2]',
 deferral='Pilot may be deferred while engineering continues. No admission or release approval follows from a draft plan or successful rehearsal.',
 questions=[('Cohort','Customers/sites/users/devices, contacts and online/offline capabilities?'),('Entry criteria','Which R packages and country/provider approvals must pass first?'),('Coverage','Which sales, returns, close, payroll and reporting cycles must be observed?'),('Volume and time','Minimum volume per workflow, duration/window and who sets the threshold?'),('Stop and recovery','Stop criteria, incident owner, data reconciliation and rollback/support plan?'),('Exit authority','Who reviews results, accepts residual limits and separately authorizes release?')]),
]

all_docs = {}
canonical_roles = [
 'Product owner + delivery lead', 'Accounting owner + qualified accountant',
 'POS/product + security + country owner', 'Privacy/records + HR + AppSec',
 'Country/integration leads + business approvers', 'Product + actual customer + domain owner',
 'SRE + product + finance', 'Release owner + pilot operations',
]
for item, roles in zip(decisions, canonical_roles):
 item['owner'] = roles
def make(stem, title, docid):
 d=Document(); sec=d.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
 sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
 for name in ['Normal','Body Text']:
  s=d.styles[name]; s.font.name='Calibri'; s.font.size=Pt(11); s.paragraph_format.space_after=Pt(6)
  s.paragraph_format.line_spacing=1.05
 for name,size in [('Title',23),('Heading 1',16),('Heading 2',12)]:
  s=d.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '17324D')
  s.paragraph_format.space_after=Pt(7)
 d.core_properties.title=title; d.core_properties.author=''; d.core_properties.last_modified_by=''
 hp=sec.header.paragraphs[0]; hp.text='STOQUIFY  |  R00  |  DRAFT - AWAITING HUMAN VALIDATION'; hp.runs[0].font.size=Pt(8)
 fp=sec.footer.paragraphs[0]; fp.text=f'{docid}  |  v{VERSION}  |  {DATE}     |     '
 fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
 for run in fp.runs: run.font.size=Pt(8)
 d.add_paragraph(title,'Title'); d.add_paragraph(f'{docid} | Revision {VERSION} | {DATE}')
 d.add_paragraph('Organization or project legal owner: ______________________________')
 all_docs[stem]={'doc':d,'md':[f'# {title}',f'{docid} | Revision {VERSION} | {DATE}', 'DRAFT — AWAITING HUMAN VALIDATION', 'Organization or project legal owner: [To complete]']}
 return stem
def p(k,text,style=None):
 all_docs[k]['doc'].add_paragraph(text,style); all_docs[k]['md'].append(text)
def h(k,text): p(k,text,'Heading 2')
def page(k,title):
 all_docs[k]['doc'].add_page_break(); p(k,title,'Heading 1'); all_docs[k]['md'].append('<!-- page break -->')
def table(k,headers,rows,widths=None):
 d=all_docs[k]['doc']; t=d.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
 if widths:
  for col,width in zip(t.columns,widths): col.width=Inches(width)
 for c,s in zip(t.rows[0].cells,headers): c.text=s
 repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for c,s in zip(cells,row): c.text=str(s)
 for n,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
  for i,c in enumerate(row.cells):
   if widths: c.width=Inches(widths[i])
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(5); par.paragraph_format.space_before=Pt(4)
    for run in par.runs: run.font.size=Pt(10.5); run.bold=(n==0)
   if n==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EAF0F5'); c._tc.get_or_add_tcPr().append(sh)
 all_docs[k]['md'] += ['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str,row)).replace('\n','<br>')+' |' for row in rows]

g=make('01_START_HERE_v1.0','R00 validation guide','R00 GUIDE')
p(g,'Start with the decision workbook. Confirm what the first supported product will include, appoint real reviewers, and record each decision against an exact version. These are internal governance templates, not statutory forms. No signature or approval has been supplied.')
table(g,['Document','Use and who completes it'],[
('02 Decision workbook','Owner drafts answers to D01–D08; each relevant specialist reviews their scope. One brief and one response page per decision.'),
('03 Reviewer appointments','Owner/delegator proposes named appointments; appointees accept and a verifier checks authority and independence.'),
('04 Approval record','Duplicate for each decision and signing authority. Reviewers record conclusions; authorized decision maker records the outcome; verification follows.'),
('05 Evidence and open items','Delivery lead tracks evidence, missing people/answers, restrictions and R00 acceptance. This document itself needs no signature.')],[2.15,4.95])
h(g,'What you can decide yourself')
p(g,'If you hold the relevant project/business authority, you can choose the target customer and product scope (D01), optional priorities (D06), propose appointments and accept commercial constraints. Final D01/D06 decisions still require joint concurrence from the roles in document 03. Your authority is not established by this pack; complete the authority fields. D03 and D08 require operational/security or release input. D07 requires SRE and finance input.')
h(g,'What needs another reviewer')
p(g,'D02 needs qualified accounting review. D04 needs privacy/records, HR and security review. D05 needs qualified Cameroon expertise and provider/submission authority for each capability. Independent technical review cannot be supplied by a second account for the implementer. Compatible roles may be combined only where the adopted policy and stricter domain controls permit it. [S5, S6]')
h(g,'What you sign')
p(g,'First confirm appointments in document 03 after authority checks; each appointee accepts. Then use document 04 for the exact completed decision sheet and its evidence. Do not sign a blank proposal or approve all eight decisions through an unexplained blanket signature. The guide and checklist are working aids.')
page(g,'How to validate and return the pack')
steps=[
('1 Read the brief','Open document 02. Read the confirmed facts, proposal, alternatives and due gate for each decision. A proposed recommendation is not your selected decision.'),
('2 Complete the response','Answer every numbered question. Enter the exact selected option and rationale. Use “not yet available” for missing facts; do not invent thresholds, legal periods or reviewer names.'),
('3 Check supporting evidence','Open each cited item. Check scope, date, version and whether it actually supports the choice. Ask for demonstrations or worked examples where a policy affects calculations or workflow.'),
('4 Appoint reviewers','Use document 03. Establish names, relevant qualifications, authority, scope, validity and conflicts. Obtain acceptance. If a role is missing, assign someone to obtain it and record the capability restriction.'),
('5 Obtain scoped conclusions','Give the relevant specialist the sheet and evidence. Ask what is correct, uncertain and missing, what cases were checked and what would invalidate their conclusion. Do not turn a technical review into legal certification.'),
('6 Record the outcome','Use document 04: approve, reject or request revision. Conditional approval requires named conditions, owners and due gates and is not final acceptance. A deferral/not-applicable decision needs a reason and capability restriction.'),
('7 Freeze before signing','Resolve edits, identify final revision, filenames and pages, and freeze the evidence bundle. A custodian can calculate a SHA-256 file digest to detect later changes. A digest does not prove identity or approval.'),
('8 Sign and verify','Use the adopted internal method. Paper: sign the exact final pages, retain the original, scan and verify identity/authority through an established channel. Native confirmation may be used only if the deployed register and verification process are available and qualified; its availability is not established here.'),
('9 Return privately','Place completed documents and authority evidence in an access-controlled approved repository. Send only a secure reference to the custodian. Do not commit identity scans, signatures, private employee data or credentials to the public code repository.'),
('10 Reconcile and maintain','Custodian checks exact version, scope, signer authority and verification, then records the decision. Keep signing, receipt and verification dates separate. Changes or revoked authority require a new review; retain superseded/revoked history.')]
for title,text in steps: h(g,title); p(g,text)
page(g,'Acceptance limits and a worked example')
p(g,'R00 can advance when every in-scope capability has a named accountable owner, release boundary, acceptance scenario and evidence location; decisions and valid deferrals are recorded and reviewer appointments are verified. Later-stage decisions need owners, due gates and operating restrictions. The delivery lead prepares the evidence and product owner/QA lead review it. R00 is not complete merely because forms exist. [S2, S8]')
p(g,'After R00, the actual software and operating evidence must still pass the relevant gates: security/tenant isolation, data and financial integrity, HRIS/payroll workflows, country/provider qualification, build and authenticated browser checks, recovery and pilot/release acceptance. Scope approval is not production, payment or filing authorization. No legal conclusion about signature enforceability is made by this internal guide.')
h(g,'Fictional example only not an approval')
p(g,'Example D06: Fictional owner A, supported by fictional customer B, selects “defer customer-cohort analytics for the first pilot.” Reason: the fictional pilot has no reviewed customer-identity/consent dataset or defined retention-analysis need. Restriction: no cohort-analytics capability is promised or enabled for that pilot. Reopen when a customer requirement and suitable data/consent review exist. Fictional domain reviewer C agrees only with this scope decision. Evidence: a fictional customer interview and data assessment. These names and references are illustrative, do not appear in live appointment fields, and confer no approval.')
h(g,'Sources and reuse')
p(g,'S1 defines these eight R00 decisions; S8 defines package gates. S2 is the current execution record. S3 confirms your operating choices; S4 records research limits. S5/S6 provide the existing appointment and verification process. S7 is used only to prevent mistaken reuse of separate POS G1 decisions or unverified people/signature images. R00 D01–D08 is not POS G1 D-01–D-11.')
p(g,'This pack adapts the existing process to R00 because the existing X1/G1 forms approve different objects. It does not replace those controls or create a new approval platform. Source paths and hashes are listed in SOURCE_INDEX.md and decision-data.json beside these documents.')

w=make('02_DECISION_WORKBOOK_v1.0','R00 decision workbook','R00 WORKBOOK')
for i,item in enumerate(decisions):
 if i: page(w,item['id']+' '+item['title'])
 else: p(w,item['id']+' '+item['title'],'Heading 1')
 p(w,'Canonical definition: '+item['exact']+'. [S1]')
 p(w,item['meaning'])
 for label,key in [('Already confirmed','confirmed'),('Proposal for your review','recommendation'),('Alternatives and consequences','alternatives'),('Evidence to inspect','evidence'),('Decision timing and deferral','deferral')]: h(w,label); p(w,item[key])
 p(w,'Accountable roles: '+item['owner']+'. Review: '+item['review']+'.')
 p(w,'Due gate: '+item['gate'])
 page(w,item['id']+' response and validation')
 p(w,'DRAFT. Complete this page, attach supporting evidence, and use document 04 for the actual decision and signature. A recommendation is not selected until you record it here.')
 table(w,['Question to answer','Your answer or evidence reference'],[(f'{item["id"]}.{n+1} {a}\n{b}','________________________________________________\n________________________________________________') for n,(a,b) in enumerate(item['questions'])],[3.35,3.75])
 p(w,'Exact selected proposal or changes: _______________________________________\n______________________________________________________________________')
 p(w,'Reason and material consequences: ________________________________________\n______________________________________________________________________')
 p(w,'Conditions or unresolved items and restrictions: _____________________________\n______________________________________________________________________')
 p(w,'Named owner: __________________  Required reviewers: _____________________')
 p(w,'Evidence versions/locations: __________________  Due gate/date: ______________')
 p(w,'Decision outcome is recorded separately in document 04. Reopen on a change to scope, policy, applicable law/rules, evidence or authority relevant to this decision.')

a=make('03_REVIEWER_APPOINTMENTS_v1.0','R00 reviewer appointments','R00 APPOINTMENTS')
p(a,'Complete the role matrix, then copy the appointment page for each person/scope. Joint authorities below retain their decision responsibility; additional technical review does not replace their concurrence. All names and authority checks are pending.')
table(a,['ID','Accountable roles','Required review coverage'],[(x['id'],x['owner'],x['review']) for x in decisions],[.6,2.75,3.75])
h(a,'Authority and independence')
p(a,'An owner may also act as product/delivery lead if their authority covers both. Product and security separation is the existing proposed default; independent technical review remains required. Combine other roles only after documenting competence, authority, conflicts and applicable domain restrictions. A creator cannot satisfy an independent review requirement by self-signing under a different title. [S5, S6]')
p(a,'Authority may be supported by ownership records, an executive/board mandate, a contract or a scoped appointment. A personally owned project may need an owner-establishment declaration supported by ownership evidence and independent verification. Qualification and delegated authority are separate checks. Use secure references, not attached private identity documents.')
p(a,'Assess maker/checker, signer/verifier, evidence producer/reviewer and personal/financial conflicts. The subject must not approve their own conflict assessment. Delegated powers cannot exceed the delegator’s scope or validity. [S6, S10]')
page(a,'Appointment and acceptance form')
p(a,'DRAFT TEMPLATE. Duplicate for each person/scope. This form becomes effective only through the adopted authority and verification process.')
table(a,['Field','Complete before confirming'],[
('Organization and project scope','________________________________________________'),('Real person and role','________________________________________________'),('Account or secure identity reference','________________________________________________'),('Decisions and powers covered','D___  Scope/limits: __________________________________'),('Delegator and authority source','________________________________________________'),('Qualification evidence if required','Private reference: __________________________________'),('Starts and expires or review trigger','________________________________________________'),('Conflicts and independence assessment','________________________________________________'),('Missing evidence and restriction','________________________________________________')],[2.5,4.6])
p(a,'Appointing authority: I appoint this person only for the scope and limits stated above, subject to verified authority. This does not grant unrelated production, payment or filing powers.')
p(a,'Name/role: _________________________  Signature/confirmation: ______________\nDate/time/time zone: __________________  Method: _________________________')
p(a,'Appointee: I accept the stated responsibility and disclose relevant conflicts.')
p(a,'Name: _____________________________  Signature/confirmation: ______________\nDate/time/time zone: ___________________________________________________')
p(a,'Verifier name/role: __________________  Verifier authority reference: ___________\nEvidence checked for issuer/appointee: ______________________________________\nIndependence result: ____________________________________________________\nVerifier signature or confirmation reference: _________________________________\nVerification date/time/time zone: __________________________________________\nOutcome: [ ] Verified  [ ] Unverified  [ ] Revision needed')

s=make('04_APPROVAL_RECORD_v1.0','R00 decision approval record','R00 DECISION RECORD')
p(s,'DRAFT TEMPLATE. Copy for each decision and required authority. This document ID identifies a template, not an executed approval. Leave the receipt/reference blank until the real decision is recorded.')
p(s,'Specify which exact artifact the signer confirms. Word and PDF are different files with different hashes; a completed or signed copy also differs from this blank template. Preserve and identify each representation separately.')
p(s,'Acceptance requires affirmative scoped decisions or concurrence from every applicable joint authority listed in document 03. Unresolved required dissent or a missing authority prevents acceptance; one owner signature cannot override it.')
table(s,['Binding field','Complete with the final reviewed version'],[
('Decision ID and title','D___  __________________________________________________'),('Organization and approved scope','________________________________________________________'),('Workbook revision and pages','________________________________________________________'),('Exact selected decision','________________________________________________________\n________________________________________________________'),('Evidence files and revisions','________________________________________________________\n________________________________________________________'),('Frozen bundle or secure location','________________________________________________________'),('Digest if recorded by custodian','________________________________________________________'),('Limits and exclusions','________________________________________________________'),('Conditions and responsible people','________________________________________________________\n________________________________________________________'),('Due dates and review triggers','________________________________________________________')],[2.1,5.0])
h(s,'Meaning of the approval')
p(s,'The decision approves only the identified scope/policy proposal. It does not assert that the implementation meets it. It does not authorize production deployment, real payments or statutory filing, and does not certify Cameroon or OHADA compliance. Expert conclusions remain limited to their explicit reviewed scope.')
page(s,'Review decision and verification')
h(s,'Specialist review repeat for each required reviewer')
p(s,'Reviewer name/role: _____________________  Appointment reference: __________\nQualification/scope: _____________________  Review date: ___________________\nEvidence and cases reviewed: ____________________________________________\nConclusion, exclusions and uncertainties: __________________________________\n______________________________________________________________________\n[ ] Supports proposal  [ ] Supports subject to stated conditions  [ ] Does not support\nSignature/confirmation and method: _______________________________________')
h(s,'Authorized decision maker')
p(s,'Name and role: _________________________  Authority reference: ______________\nOutcome: [ ] Approve  [ ] Reject  [ ] Request revision\nIf approving: [ ] Unconditional for stated scope  [ ] Conditional as listed\nReason and exact conditions: _____________________________________________\n______________________________________________________________________\nSignature/confirmation: __________________  Method: ________________________\nSigning date/time/time zone: ______________________________________________')
h(s,'Custodian verification after signing')
p(s,'Verifier name/role: ______________________  Verifier authority reference: __________\nVerification method: _____________________  Signer evidence checked: __________\nExact signed version matches reviewed bundle: [ ] Yes  [ ] No  [ ] Unverified\nAll required authorities concur and independence checked: [ ] Yes  [ ] No  [ ] Unverified\nReceived date/time/time zone: ____________________________________________\nVerifier signature or confirmation reference: _________________________________\nVerification date/time/time zone: __________________________________________\nRecord outcome: [ ] Verified  [ ] Unverified  [ ] Revision required\nSecure signed-file location: ______________________________________________\nActual receipt/reference if issued: _________________________________________')
p(s,'Do not backdate a confirmation or treat an image of a signature as verified identity. For paper signing, retain the original and verify the exact scan through the adopted process. A changed file, scope, authority or material condition requires a new linked decision; retain the old history. Conditional approval does not close an unresolved release gate. [S6]')

e=make('05_EVIDENCE_AND_OPEN_ITEMS_v1.0','R00 evidence and open items','R00 EVIDENCE')
p(e,'The delivery lead maintains this checklist and the associated structured decision data. It records progress and missing evidence; it cannot supply a human approval. Store private supporting records outside the public repository.')
table(e,['Check','Evidence needed before acceptance'],[
('Scope and baseline','Exact candidate/baseline, migration/toolchain references, cohort, currencies, modules, exclusions and existing findings. [S2, S3]'),('Decisions','Each D01–D08 has a selected decision or valid bounded deferral, rationale, owner, due gate and evidence location.'),('People','Verified names, appointments, authority scope, qualification where needed, acceptance and independence assessment.'),('Decision recording','Exact revision and reviewed evidence; authentic decision outcome; signing and verification records; conditions and restrictions.'),('Acceptance scenarios','For every in-scope capability, identify the workflow/test/example that would demonstrate success and the responsible reviewer.'),('Records integrity','All links resolve; signed records are private and preserved; supersession/revocation does not overwrite original history.'),('R00 review','Delivery lead prepares; product owner and QA lead assess the complete scope/owner/decision evidence. R00 remains PARTIAL until acceptance is actually met.')],[1.65,5.45])
h(e,'Current evidence limits')
p(e,'October 7 engineering evidence records 59 focused tests, four isolated PostgreSQL tests, targeted lint and full typecheck passing. That does not qualify the full product. Shared leave-balance concurrency, authenticated browser workflows, integrated build/package, recovery and statutory validation remain open. [S3]')
h(e,'Statuses used in this pack')
p(e,'Templates: DRAFT / AWAITING HUMAN VALIDATION. Unmade decisions and identities are null in decision-data.json. Future outcomes must be supported by real records; conditional or deferred items retain their restrictions. Canonical R00 stays PARTIAL and humanAssignment stays UNCONFIRMED until verified updates are justified.')
page(e,'Unresolved decisions and completion review')
table(e,['Item','Still needed','Owner and due gate'],[(x['id'],x['exact'],x['owner']+'\n'+x['gate']+'\nNamed person/date: __________') for x in decisions],[.5,3.25,3.35])
p(e,'For each open item add: secure evidence reference, missing action, named responsible person, target date, operating restriction, resolution evidence and reviewer. Reuse this row structure in the structured register rather than marking an empty field approved.')
p(e,'R00 review conclusion: [ ] Continue preparation  [ ] Request revision  [ ] Recommend acceptance for stated scope\nPrepared by/date: __________________  Product/QA review references: ___________\nUnresolved restrictions: __________________________________________________')
p(e,'No R00 conclusion bypasses R18/R19 country/provider gates, R31 release evidence, R32 offline pilot requirements or R33 release authority. A later final release decision must identify the exact candidate and permitted cohort/capabilities.')
page(e,'Capability and open item worksheet')
p(e,'Copy this page for each in-scope capability or unresolved condition. Use separate entries for online, offline, accounting and HR/payroll cohorts. This links your choices to the evidence required by R00.')
table(e,['Field','Your entry'],[
('Decision and package IDs','D___ / R___  __________________________________________'),
('Capability and module','________________________________________________________'),
('Country provider and customer cohort','________________________________________________________'),
('Operating boundary','[ ] Demo  [ ] Proposed pilot  [ ] Future release  [ ] Excluded'),
('Named accountable owner','________________________________________________________'),
('Required authority and reviewer','________________________________________________________'),
('Acceptance scenario and expected result','________________________________________________________\n________________________________________________________'),
('Evidence location version and date','________________________________________________________\n________________________________________________________'),
('Missing decision or evidence','________________________________________________________'),
('Remediation owner and due date or gate','________________________________________________________'),
('Restriction until resolved','________________________________________________________'),
('Resolution evidence and reviewed date','________________________________________________________'),
('Status and linked decision record','________________________________________________________')],[2.75,4.35])
p(e,'Baseline binding for this scope: commit/revision ______________________________\nRelevant uncommitted files/manifest ________________________________________\nMigration/toolchain evidence reference ______________________________________\nKnown findings and changed-evidence triage __________________________________')

OUT.mkdir(parents=True,exist_ok=True)
for stem,item in all_docs.items():
 for element in (item['doc'].element, item['doc'].styles.element):
  for border in element.xpath('.//w:pBdr'):
   border.getparent().remove(border)
 item['doc'].save(OUT/(stem+'.docx'))
 (OUT/(stem+'.md')).write_text('\n\n'.join(item['md'])+'\n',encoding='utf-8')
for item in decisions:
 item.update({'status':'AWAITING_HUMAN_VALIDATION','selectedDecision':None,'decisionMaker':None,'reviewers':[], 'outcome':None,'signature':None,'approvalReference':None,'validationTracking':{'namedOwner':None,'dueDate':None,'evidenceReferences':[],'unresolvedItems':[],'operatingRestrictions':[],'resolutionEvidence':None,'verifiedBy':None,'verifiedAt':None}})
data={'schemaVersion':'stoquify-r00-preparation.v1','revision':VERSION,'date':'2026-10-07','status':'DRAFT','programme':'platform-correctness-roadmap/2026-09-30-run-01','package':'R00','productionAuthorized':False,'confirmedContext':['Cameroon first','Later CEMAC expansion','Pre-production test/demo data','Full phased modernization'], 'decisions':decisions, 'sources':[{'id':key,'path':value,'sha256':hashlib.sha256((ROOT/value).read_bytes()).hexdigest()} for key,value in sources.items()]}
(OUT/'decision-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
index=['# R00 source index','Revision 1.0 | 7 October 2026','These are preparation sources, not signed approval receipts. References are relative to the repository root.']
for key,value in sources.items(): index += [f'- {key}: [{Path(value).name}](../../../{value})']
(OUT/'SOURCE_INDEX.md').write_text('\n\n'.join(index)+'\n',encoding='utf-8')
readme=['# R00 decision and reviewer pack','DRAFT — AWAITING HUMAN VALIDATION | Revision 1.0 | 7 October 2026','Start with the guide, then complete the workbook. Read the guide before signing anything. R00 remains PARTIAL.']
for stem,item in all_docs.items(): readme += [f'- {item["doc"].core_properties.title}: [Word]({stem}.docx) · [PDF]({stem}.pdf) · [Markdown]({stem}.md)']
readme += ['Editable Word documents contain ordinary text/table fields, not electronic-signature widgets. PDFs are printable review copies. Completed private records must be stored through the adopted controlled process, not committed publicly.', '[Source index](SOURCE_INDEX.md) · [Structured decisions](decision-data.json)', 'Documents 03 and 04 contain the signature/confirmation fields. Copy document 03 for each appointee and document 04 for each decision/required authority. The remaining documents support preparation and review.']
(OUT/'README.md').write_text('\n\n'.join(readme)+'\n',encoding='utf-8')
print(json.dumps({'documents':len(all_docs),'decisions':len(decisions),'directory':str(OUT)}))
